from pathlib import Path
import re
from datetime import date, datetime, timedelta

from django.conf import settings
from django.http import FileResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_GET
from django.views.decorators.http import require_POST, require_http_methods

from app.financial_register import ledger_service
from app.financial_register import service as financial_register_service
from app.financial_register.disbursements import service as disbursement_service
from app.financial_register.contribution_coverage import (
    service as coverage_service,
    workbook as coverage_workbook,
)
from app.financial_register.reconciliation import service as reconciliation_service
from app.financial_register.statements import service as statement_service
from app.financial_register.verification import service as verification_service
from app.financial_register.fines import service as fines_service
from app.financial_register.contribution_monitoring import (
    service as contribution_monitoring_service,
)
from app.financial_register.audit_trail import service as audit_trail_service
from app.security import audit_log


ADMIN_MEMBER_ID = "TLC/001/2026"
EXCLUDED_PAYMENT_STATUSES = {"CANCELLED", "REJECTED"}


def _data_file(name):
    return Path(settings.TLC_DATA_DIR) / "spreadsheets" / name


def _configure_data_paths():
    members_file = _data_file("01_MEMBERS.xlsx")
    contributions_file = _data_file("02_CONTRIBUTIONS.xlsx")
    coverage_file = _data_file("03_CONTRIBUTION_COVERAGE.xlsx")

    ledger_service.MEMBERS_FILE = members_file
    ledger_service.CONTRIBUTIONS_FILE = contributions_file
    disbursement_service.DISBURSEMENTS_FILE = _data_file(
        "11_TREASURY_DISBURSEMENTS.xlsx"
    )
    financial_register_service.MEMBERS_FILE = members_file
    financial_register_service.CONTRIBUTIONS_FILE = contributions_file
    verification_service.CONTRIBUTIONS_FILE = contributions_file
    coverage_workbook.COVERAGE_FILE = coverage_file
    coverage_service.COVERAGE_FILE = coverage_file
    coverage_service.CONTRIBUTIONS_FILE = contributions_file
    reconciliation_service.MEMBERS_FILE = members_file
    reconciliation_service.CONTRIBUTIONS_FILE = contributions_file
    audit_log.ACTIVITY_DIR = Path(settings.TLC_DATA_DIR) / "activity"
    audit_log.ACTIVITY_FILE = audit_log.ACTIVITY_DIR / "recent_activity.json"


def _clean(value):
    return "" if value is None else str(value).strip()


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _format_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return _clean(value)


def _format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    return _clean(value)


def _read_member_contributions(member_id):
    from openpyxl import load_workbook

    contributions_file = _data_file("02_CONTRIBUTIONS.xlsx")
    if not contributions_file.exists():
        raise FileNotFoundError

    workbook = load_workbook(
        contributions_file,
        read_only=True,
        data_only=True,
    )
    try:
        sheet = workbook["Contributions"]
        headers = [_clean(cell.value) for cell in sheet[1]]
        rows = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in row):
                continue
            record = dict(zip(headers, row))
            row_member_id = _clean(record.get("Member ID"))
            if row_member_id != _clean(member_id):
                continue
            rows.append(
                {
                    "contribution_id": _clean(record.get("Contribution ID")),
                    "receipt_number": _clean(record.get("Receipt Number")),
                    "member_id": row_member_id,
                    "full_name": _clean(record.get("Full Name")),
                    "phone_number": _clean(record.get("Phone Number")),
                    "email": _clean(record.get("Email")),
                    "payment_date": _format_date(record.get("Payment Date")),
                    "payment_time": _format_time(record.get("Payment Time")),
                    "payment_type": _clean(record.get("Payment Type")),
                    "description": _clean(record.get("Description")),
                    "amount": _amount(record.get("Amount")),
                    "payment_method": _clean(record.get("Payment Method")),
                    "transaction_reference": _clean(
                        record.get("Transaction Reference")
                    ),
                    "status": _clean(record.get("Status")),
                    "recorded_by": _clean(record.get("Recorded By")),
                    "created_at": _format_date(record.get("Created At")),
                }
            )
        return sorted(
            rows,
            key=lambda item: (
                item["payment_date"],
                item["payment_time"],
            ),
            reverse=True,
        )
    finally:
        workbook.close()


@require_GET
def member_contributions(request):
    member_id = request.session.get("member_id")
    if not member_id:
        return JsonResponse(
            {"success": False, "message": "You must be logged in."},
            status=401,
        )

    try:
        contributions = _read_member_contributions(member_id)
        valid_contributions = [
            item
            for item in contributions
            if item["status"].upper() not in EXCLUDED_PAYMENT_STATUSES
        ]
        current_month = datetime.now().strftime("%Y-%m")
        summary = {
            "total": sum(item["amount"] for item in valid_contributions),
            "this_month": sum(
                item["amount"]
                for item in valid_contributions
                if item["payment_date"].startswith(current_month)
            ),
            "count": len(valid_contributions),
            "latest": contributions[0] if contributions else None,
        }
        return JsonResponse(
            {
                "success": True,
                "member_id": member_id,
                "summary": summary,
                "contributions": contributions,
            }
        )
    except FileNotFoundError:
        return JsonResponse(
            {
                "success": False,
                "message": "Contribution records are unavailable.",
            },
            status=500,
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to load your contributions.",
            },
            status=500,
        )


def _is_authorized(request):
    return (
        str(request.session.get("member_id") or "").strip().upper()
        == ADMIN_MEMBER_ID
    )


@require_GET
def financial_register(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )

    year = _clean(request.GET.get("year"))
    month = _clean(request.GET.get("month"))
    if year and (not year.isdigit() or len(year) != 4):
        return JsonResponse(
            {"success": False, "message": "Invalid year."},
            status=400,
        )
    year = year or None

    if month:
        if not month.isdigit() or not 1 <= int(month) <= 12:
            return JsonResponse(
                {"success": False, "message": "Invalid month."},
                status=400,
            )
        if not year:
            return JsonResponse(
                {"success": False, "message": "A month requires a year."},
                status=400,
            )
        month = str(int(month)).zfill(2)
    else:
        month = None

    try:
        _configure_data_paths()
        report = ledger_service.build_financial_ledger(year=year, month=month)
        return JsonResponse({"success": True, "report": report})
    except FileNotFoundError as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to build the TLC financial register.",
            },
            status=500,
        )


def _attach_statement_coverage(statement, member_id, year, month):
    if not year or not month:
        return
    try:
        coverage = coverage_service.get_member_monthly_coverage(
            member_id.upper(),
            int(year),
            int(month),
        )
    except Exception:
        return
    if not isinstance(coverage, dict):
        return

    weeks = coverage.get("weeks") or []
    statement["coverage"] = coverage
    statement["coverageWeeks"] = weeks
    statement["coverageAmount"] = coverage.get("amount", 0)
    statement["coverageWeekCount"] = coverage.get("week_count", len(weeks))

    receipts = {
        str(row["Receipt Number"]).strip().upper()
        for row in weeks
        if row.get("Receipt Number")
    }
    contribution_ids = {
        str(row["Contribution ID"]).strip().upper()
        for row in weeks
        if row.get("Contribution ID")
    }
    if not receipts and not contribution_ids:
        return

    try:
        yearly = statement_service.build_member_statement(
            member_id=member_id.upper(),
            year=str(int(year)),
            month=None,
        )
    except Exception:
        return
    if not isinstance(yearly, dict):
        return

    history = (yearly.get("member") or {}).get("transactions")
    if not isinstance(history, list):
        history = yearly.get("transactions")
    if not isinstance(history, list):
        return

    identity_keys = (
        "receipt",
        "receipt_number",
        "receiptNumber",
        "Receipt Number",
        "reference",
        "payment_reference",
        "paymentReference",
        "contribution_id",
        "contributionId",
        "Contribution ID",
    )

    def identities(transaction):
        return {
            str(transaction[key]).strip().upper()
            for key in identity_keys
            if transaction.get(key) not in (None, "")
        }

    source_transactions = [
        transaction
        for transaction in history
        if isinstance(transaction, dict)
        and (
            identities(transaction).intersection(receipts)
            or identities(transaction).intersection(contribution_ids)
        )
    ]
    member = statement.setdefault("member", {})
    current = member.get("transactions")
    if not isinstance(current, list):
        current = []
    existing = set()
    for transaction in current:
        if isinstance(transaction, dict):
            existing.update(identities(transaction))
    for transaction in source_transactions:
        transaction_ids = identities(transaction)
        if not transaction_ids.intersection(existing):
            current.append(transaction)
            existing.update(transaction_ids)
    member["transactions"] = current
    statement["transactions"] = current
    statement["coverage_source_transactions"] = source_transactions


@require_GET
def member_statement(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    member_id = _clean(request.GET.get("member_id"))
    year = _clean(request.GET.get("year"))
    month = _clean(request.GET.get("month"))
    if not member_id:
        return JsonResponse(
            {"success": False, "message": "Member ID is required."},
            status=400,
        )
    if year and (not year.isdigit() or len(year) != 4):
        return JsonResponse(
            {"success": False, "message": "Invalid year."},
            status=400,
        )
    if month:
        if not year:
            return JsonResponse(
                {"success": False, "message": "A month requires a year."},
                status=400,
            )
        if not month.isdigit() or not 1 <= int(month) <= 12:
            return JsonResponse(
                {"success": False, "message": "Invalid month."},
                status=400,
            )
        month = str(int(month)).zfill(2)

    try:
        _configure_data_paths()
        statement = statement_service.build_member_statement(
            member_id=member_id,
            year=year or None,
            month=month or None,
        )
        if statement is None:
            return JsonResponse(
                {"success": False, "message": "Member was not found."},
                status=404,
            )
        _attach_statement_coverage(statement, member_id, year, month)
        return JsonResponse({"success": True, "statement": statement})
    except ValueError as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=400,
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to build the member statement.",
            },
            status=500,
        )


@require_GET
def verify_receipt(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    receipt_number = _clean(request.GET.get("receipt"))
    if not receipt_number:
        return JsonResponse(
            {"success": False, "message": "Receipt number is required."},
            status=400,
        )
    try:
        _configure_data_paths()
        receipt = verification_service.find_receipt(receipt_number)
        if receipt is None:
            return JsonResponse(
                {
                    "success": False,
                    "message": (
                        "No valid TLC receipt was found with that reference."
                    ),
                },
                status=404,
            )
        return JsonResponse({"success": True, "receipt": receipt})
    except FileNotFoundError as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )
    except Exception:
        return JsonResponse(
            {"success": False, "message": "Unable to verify the receipt."},
            status=500,
        )


@require_GET
def receipt_options(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        _configure_data_paths()
        return JsonResponse(
            {
                "success": True,
                "results": verification_service.list_receipt_options(),
            }
        )
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_GET
def financial_reconciliation(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        _configure_data_paths()
        result = reconciliation_service.build_reconciliation()
        return JsonResponse({"success": True, "reconciliation": result})
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_GET
def fines_register(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    year = request.GET.get("year") or None
    month = request.GET.get("month") or None
    try:
        if month and not year:
            return JsonResponse(
                {"success": False, "message": "A month requires a year."},
                status=400,
            )
        _configure_data_paths()
        result = fines_service.build_fines_register(year=year, month=month)
        return JsonResponse({"success": True, "fines": result})
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_GET
def contribution_monitoring(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        try:
            weeks = int(request.GET.get("weeks", 8))
        except (TypeError, ValueError):
            weeks = 8
        weeks = max(1, min(12, weeks))
        _configure_data_paths()
        result = contribution_monitoring_service.build_contribution_monitoring(
            history_weeks=weeks
        )
        return JsonResponse({"success": True, "monitoring": result})
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_GET
def financial_audit_trail(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        _configure_data_paths()
        result = audit_trail_service.build_financial_audit_trail(
            limit=request.GET.get("limit", 200)
        )
        return JsonResponse({"success": True, "audit": result})
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_POST
def reset_financial_audit_trail(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        _configure_data_paths()
        return JsonResponse(audit_trail_service.clear_financial_audit_trail())
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


def get_disbursements(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    year = request.GET.get("year") or None
    month = request.GET.get("month") or None
    if year:
        if not str(year).isdigit() or len(str(year)) != 4:
            return JsonResponse(
                {"success": False, "message": "Invalid year."},
                status=400,
            )
        year = str(year)
    if month:
        if not str(month).isdigit() or not 1 <= int(month) <= 12:
            return JsonResponse(
                {"success": False, "message": "Invalid month."},
                status=400,
            )
        if not year:
            return JsonResponse(
                {"success": False, "message": "A month requires a year."},
                status=400,
            )
        month = str(int(month)).zfill(2)
    try:
        _configure_data_paths()
        result = disbursement_service.build_disbursement_register(
            year=year,
            month=month,
        )
        return JsonResponse({"success": True, "register": result})
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


def create_disbursement(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    try:
        import json

        data = json.loads(request.body or b"{}")
        if not isinstance(data, dict):
            data = {}
    except (json.JSONDecodeError, UnicodeDecodeError):
        data = {}
    try:
        _configure_data_paths()
        result = disbursement_service.record_disbursement(
            amount=data.get("amount"),
            category=data.get("category"),
            purpose=data.get("purpose"),
            paid_to=data.get("paid_to"),
            payment_method=data.get("payment_method"),
            reference=data.get("reference"),
            authorized_by=data.get("authorized_by"),
            recorded_by=data.get("recorded_by"),
            notes=data.get("notes", ""),
            transaction_date=data.get("date"),
            transaction_time=data.get("time"),
        )
        audit_log.record_activity(
            ADMIN_MEMBER_ID,
            action="Treasury Disbursement Recorded",
            description=(
                f"Disbursement {result['disbursement_id']} recorded "
                f"for KSh {result['amount']:,.2f}."
            ),
            category="financial",
            details={
                "disbursement_id": result["disbursement_id"],
                "amount": result["amount"],
                "category": result["category"],
                "purpose": result["purpose"],
                "paid_to": result["paid_to"],
                "payment_method": result["payment_method"],
                "reference": result["reference"],
            },
        )
        return JsonResponse(
            {"success": True, "disbursement": result},
            status=201,
        )
    except ValueError as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=400,
        )
    except Exception as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=500,
        )


@require_http_methods(["GET", "POST"])
def disbursements(request):
    if request.method == "GET":
        return get_disbursements(request)
    return create_disbursement(request)


@require_GET
def member_financial_summary(request):
    member_id = request.session.get("member_id")
    if not member_id:
        return JsonResponse(
            {"success": False, "message": "You must be logged in."},
            status=401,
        )
    try:
        contributions = _read_member_contributions(member_id)
        valid = [
            item
            for item in contributions
            if item["status"].upper() not in EXCLUDED_PAYMENT_STATUSES
        ]
        now = datetime.now()
        amounts = [float(item.get("amount") or 0) for item in valid]
        return JsonResponse(
            {
                "success": True,
                "member_id": member_id,
                "summary": {
                    "total": sum(amounts),
                    "this_month": sum(
                        amount
                        for amount, item in zip(amounts, valid)
                        if item["payment_date"].startswith(now.strftime("%Y-%m"))
                    ),
                    "this_year": sum(
                        amount
                        for amount, item in zip(amounts, valid)
                        if item["payment_date"].startswith(str(now.year))
                    ),
                    "payment_count": len(valid),
                    "average": sum(amounts) / len(amounts) if amounts else 0,
                    "first_contribution": valid[-1] if valid else None,
                    "latest_contribution": valid[0] if valid else None,
                },
            }
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to load your financial summary.",
            },
            status=500,
        )


def _payment_transaction_date(transaction):
    for key in (
        "date",
        "payment_date",
        "transaction_date",
        "received_date",
        "created_at",
        "received_at",
    ):
        value = transaction.get(key)
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = _clean(value)
        for date_format in (
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%d %B %Y",
            "%d %b %Y",
            "%B %d, %Y",
            "%b %d, %Y",
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%dT%H:%M:%S.%f",
        ):
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                continue
        match = re.search(r"\d{1,2}\s+[A-Za-z]+\s+\d{4}", text)
        if match:
            for date_format in ("%d %B %Y", "%d %b %Y"):
                try:
                    return datetime.strptime(
                        match.group(0), date_format
                    ).date()
                except ValueError:
                    continue
    return None


def _payment_transaction_text(transaction):
    return " ".join(
        str(transaction[key])
        for key in (
            "description",
            "payment_type",
            "type",
            "transaction_type",
            "details",
            "notes",
            "remarks",
        )
        if transaction.get(key)
    )


def _payment_week_count(transaction):
    for key in ("weeks", "week_count", "coverage_weeks", "number_of_weeks"):
        value = transaction.get(key)
        if value is not None:
            try:
                count = int(float(str(value).strip()))
                if count > 0:
                    return count
            except (TypeError, ValueError):
                pass
    match = re.search(
        r"(\d+)\s*[- ]?\s*WEEKS?",
        _payment_transaction_text(transaction),
        flags=re.IGNORECASE,
    )
    return int(match.group(1)) if match else 0


def _payment_transaction_amount(transaction):
    value = (
        transaction.get("amount")
        or transaction.get("amount_paid")
        or transaction.get("paid_amount")
        or transaction.get("value")
    )
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group(0)) if match else 0.0


def _build_member_payment_coverage(history_statement, year, month):
    if not year or not month:
        weeks = []
        return {
            "year": None,
            "month": None,
            "month_label": "",
            "period": "",
            "amount": 0.0,
            "week_count": 0,
            "weeks_catered_for": 0,
            "weeks": weeks,
            "coverageWeeks": weeks,
            "coverageAmount": 0.0,
            "coverageWeekCount": 0,
        }

    target_year = int(year)
    target_month = int(month)
    member = history_statement.get("member") or {}
    transactions = member.get("transactions")
    if not isinstance(transactions, list):
        transactions = history_statement.get("transactions") or []
    coverage_rows = []
    for transaction in transactions:
        if not isinstance(transaction, dict):
            continue
        paid_date = _payment_transaction_date(transaction)
        week_count = _payment_week_count(transaction)
        total_amount = _payment_transaction_amount(transaction)
        if (
            paid_date is None
            or paid_date.year not in (target_year - 1, target_year)
            or week_count <= 0
            or total_amount <= 0
        ):
            continue
        weekly_amount = total_amount / week_count
        for index in range(week_count):
            covered_date = paid_date + timedelta(weeks=index)
            if covered_date.year == target_year and covered_date.month == target_month:
                received = paid_date.strftime("%d %B %Y")
                formatted_covered = covered_date.strftime("%d %B %Y")
                coverage_rows.append(
                    {
                        "week": index + 1,
                        "week_number": index + 1,
                        "label": f"WEEK {index + 1}",
                        "date": formatted_covered,
                        "coverage_date": formatted_covered,
                        "payment_received": received,
                        "payment_received_date": received,
                        "amount": round(weekly_amount, 2),
                        "description": f"Payment received {received}",
                        "source_amount": round(total_amount, 2),
                        "source_week_count": week_count,
                    }
                )
    coverage_rows.sort(
        key=lambda row: _payment_transaction_date({"date": row["coverage_date"]})
        or date.min
    )
    amount = round(sum(row["amount"] for row in coverage_rows), 2)
    week_count = len(coverage_rows)
    month_label = f"{target_year}-{target_month:02d}"
    return {
        "year": target_year,
        "month": target_month,
        "month_label": month_label,
        "period": month_label,
        "amount": amount,
        "coverageAmount": amount,
        "week_count": week_count,
        "coverageWeekCount": week_count,
        "weeks_catered_for": week_count,
        "weeks": coverage_rows,
        "coverageWeeks": coverage_rows,
    }


@require_GET
def member_payment_statement(request):
    member_id = _clean(request.session.get("member_id")).upper()
    if not member_id:
        return JsonResponse(
            {
                "success": False,
                "message": "Please log in to view your payment statement.",
            },
            status=401,
        )
    year = _clean(request.GET.get("year"))
    month = _clean(request.GET.get("month"))
    if year and (not year.isdigit() or len(year) != 4):
        return JsonResponse(
            {"success": False, "message": "Invalid year."},
            status=400,
        )
    if month:
        if not year:
            return JsonResponse(
                {"success": False, "message": "A month requires a year."},
                status=400,
            )
        if not month.isdigit() or not 1 <= int(month) <= 12:
            return JsonResponse(
                {"success": False, "message": "Invalid month."},
                status=400,
            )
        month = str(int(month)).zfill(2)
    try:
        _configure_data_paths()
        statement = statement_service.build_member_statement(
            member_id=member_id,
            year=year or None,
            month=month or None,
        )
        if statement is None:
            return JsonResponse(
                {
                    "success": False,
                    "message": "Your member statement was not found.",
                },
                status=404,
            )
        history_statement = (
            statement_service.build_member_statement(
                member_id=member_id,
                year=year,
                month=None,
            )
            if year and month
            else statement
        )
        coverage = _build_member_payment_coverage(
            history_statement,
            year,
            month,
        )
        if year and month and coverage["weeks"]:
            covered_dates = {
                _payment_transaction_date({"date": row["coverage_date"]})
                for row in coverage["weeks"]
            }
            covered_dates.discard(None)
            source_transactions = []
            history_member = (history_statement or {}).get("member") or {}
            history_transactions = history_member.get("transactions")
            if not isinstance(history_transactions, list):
                history_transactions = (history_statement or {}).get(
                    "transactions", []
                )
            for transaction in history_transactions:
                if not isinstance(transaction, dict):
                    continue
                paid_date = _payment_transaction_date(transaction)
                count = _payment_week_count(transaction)
                if not paid_date or count <= 0:
                    continue
                coverage_end = paid_date + timedelta(weeks=count) - timedelta(days=1)
                if any(paid_date <= covered <= coverage_end for covered in covered_dates):
                    source_transactions.append(transaction)

            member = statement.setdefault("member", {})
            current = member.get("transactions")
            if not isinstance(current, list):
                current = []

            def transaction_key(transaction):
                return (
                    _clean(
                        transaction.get("receipt")
                        or transaction.get("receipt_number")
                        or transaction.get("reference")
                        or transaction.get("payment_reference")
                    ),
                    str(_payment_transaction_date(transaction) or ""),
                    round(_payment_transaction_amount(transaction), 2),
                    _payment_transaction_text(transaction),
                )

            existing = {transaction_key(item) for item in current}
            for transaction in source_transactions:
                key = transaction_key(transaction)
                if key not in existing:
                    current.append(transaction)
                    existing.add(key)
            current.sort(
                key=lambda item: _payment_transaction_date(item) or date.min,
                reverse=True,
            )
            member["transactions"] = current
        statement["coverage"] = coverage
        statement["coverageWeeks"] = coverage["coverageWeeks"]
        statement["coverageAmount"] = coverage["coverageAmount"]
        statement["coverageWeekCount"] = coverage["coverageWeekCount"]
        return JsonResponse({"success": True, "statement": statement})
    except ValueError as error:
        return JsonResponse(
            {"success": False, "message": str(error)},
            status=400,
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to load your payment statement.",
            },
            status=500,
        )


@require_GET
def financial_register_page(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    return render(request, "portal/pages/tlc_financial_register.html")


def _admin_page(request, page_name):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )
    return render(request, f"portal/pages/{page_name}.html")


@require_GET
def member_statements_page(request):
    return _admin_page(request, "tlc_member_statements")


@require_GET
def receipt_verification_page(request):
    return _admin_page(request, "tlc_receipt_verification")


@require_GET
def financial_reconciliation_page(request):
    return _admin_page(request, "tlc_financial_reconciliation")


@require_GET
def fines_register_page(request):
    return _admin_page(request, "tlc_fines_register")


@require_GET
def contribution_monitoring_page(request):
    return _admin_page(request, "tlc_contribution_monitoring")


@require_GET
def financial_audit_trail_page(request):
    return _admin_page(request, "tlc_financial_audit_trail")


@require_GET
def disbursements_page(request):
    return render(request, "portal/pages/tlc_disbursements.html")


@require_GET
def download_excel(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )

    try:
        _configure_data_paths()
        from app.financial_register.exporter import create_excel_report

        return FileResponse(
            create_excel_report(),
            as_attachment=True,
            filename="TLC_Financial_Register.xlsx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to generate the Excel report.",
            },
            status=500,
        )


@require_GET
def download_pdf(request):
    if not _is_authorized(request):
        return JsonResponse(
            {"success": False, "message": "Access denied."},
            status=403,
        )

    try:
        _configure_data_paths()
        from app.financial_register.exporter import create_pdf_report

        return FileResponse(
            create_pdf_report(),
            as_attachment=True,
            filename="TLC_Financial_Register.pdf",
            content_type="application/pdf",
        )
    except Exception:
        return JsonResponse(
            {
                "success": False,
                "message": "Unable to generate the PDF report.",
            },
            status=500,
        )
