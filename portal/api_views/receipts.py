import json
import mimetypes
import threading
import uuid
from datetime import date, datetime, timezone
from pathlib import Path

from django.conf import settings
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import render
from django.views.decorators.http import require_GET, require_POST
from openpyxl import load_workbook

from app.receipts.receipt_generator import create_receipt
from app.receipts.receipt_number import format_receipt_number


ADMIN_MEMBER_ID = "TLC/001/2026"
EXCLUDED_STATUSES = {"CANCELLED", "REJECTED"}
_activity_lock = threading.Lock()


def _data_path(*parts):
    return Path(settings.TLC_DATA_DIR).joinpath(*parts)


def _read_json(path, fallback):
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return fallback
    return value


def _load_receipts():
    path = _data_path("receipts", "receipts.json")
    if not path.exists():
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text("[]", encoding="utf-8")
        except OSError:
            return []
    data = _read_json(path, [])
    return data if isinstance(data, list) else []


def _save_receipts(receipts):
    path = _data_path("receipts", "receipts.json")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(receipts, indent=4, ensure_ascii=False),
        encoding="utf-8",
    )


def _member_details(member_id):
    workbook_path = _data_path("spreadsheets", "01_MEMBERS.xlsx")
    if not member_id or not workbook_path.is_file():
        return None

    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        try:
            if "Members" not in workbook.sheetnames:
                return None
            rows = workbook["Members"].iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                return None
            headers = [str(value).strip() if value is not None else "" for value in headers]
            for row in rows:
                member = dict(zip(headers, row))
                current_id = str(member.get("Member ID") or "").strip().upper()
                if current_id != member_id.strip().upper():
                    continue
                return {
                    "member_id": current_id,
                    "full_name": str(member.get("Full Name") or "").strip(),
                    "phone_number": str(member.get("Phone Number") or "").strip(),
                    "email": str(member.get("Email Address") or "").strip(),
                    "membership_status": str(member.get("Membership Status") or "").strip(),
                    "registration_date": str(member.get("Registration Date") or "").strip(),
                    "account_status": str(member.get("Account Status") or "").strip(),
                }
        finally:
            workbook.close()
    except Exception:
        return None
    return None


def _enrich_receipt(receipt):
    if not isinstance(receipt, dict):
        return receipt
    member = _member_details(receipt.get("member_id"))
    if not member:
        return receipt
    enriched = dict(receipt)
    for key in ("member_id", "full_name", "phone_number", "email"):
        enriched[key] = member[key]
    return enriched


def _member_receipts(member_id):
    member_id = str(member_id or "").strip().upper()
    if not member_id:
        return []
    return [
        _enrich_receipt(receipt)
        for receipt in _load_receipts()
        if isinstance(receipt, dict)
        and str(receipt.get("member_id", "")).strip().upper() == member_id
    ]


def _record_activity(member_id, action, description, category, details):
    member_id = str(member_id or "").strip().upper()
    if not member_id:
        return
    path = _data_path("activity", "recent_activity.json")
    path.parent.mkdir(parents=True, exist_ok=True)
    record = {
        "id": uuid.uuid4().hex,
        "member_id": member_id,
        "action": action,
        "description": description,
        "category": category,
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "details": details if isinstance(details, dict) else {},
    }
    with _activity_lock:
        records = _read_json(path, [])
        if not isinstance(records, list):
            records = []
        records.append(record)
        path.write_text(
            json.dumps(records[-2000:], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def _clean(value):
    return "" if value is None else str(value).strip()


def _format_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = _clean(value)
    if not text:
        return ""
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).strftime("%Y-%m-%d")
    except ValueError:
        return text.split(" ", 1)[0]


def _format_time(value):
    if isinstance(value, datetime) or hasattr(value, "strftime"):
        try:
            return value.strftime("%H:%M:%S")
        except Exception:
            pass
    return _clean(value)


def _contribution_rows():
    path = _data_path("spreadsheets", "02_CONTRIBUTIONS.xlsx")
    if not path.is_file():
        raise FileNotFoundError("TLC contribution workbook was not found.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Contributions"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return
        headers = [_clean(value) for value in headers]
        for row in rows:
            yield dict(zip(headers, row))
    finally:
        workbook.close()


def _find_verified_receipt(receipt_number):
    for record in _contribution_rows():
        stored_number = _clean(record.get("Receipt Number"))
        if stored_number.upper() != receipt_number.upper():
            continue
        status = _clean(record.get("Status")).upper()
        if status in EXCLUDED_STATUSES:
            return None
        try:
            amount = float(record.get("Amount") or 0)
        except (TypeError, ValueError):
            amount = 0.0
        return {
            "verified": True,
            "receipt_number": stored_number,
            "member_id": _clean(record.get("Member ID")),
            "full_name": _clean(record.get("Full Name")),
            "payment_date": _format_date(record.get("Payment Date")),
            "payment_time": _format_time(record.get("Payment Time")),
            "payment_type": _clean(record.get("Payment Type")),
            "description": _clean(record.get("Description")),
            "amount": amount,
            "payment_method": _clean(record.get("Payment Method")),
            "transaction_reference": _clean(record.get("Transaction Reference")),
            "contribution_id": _clean(record.get("Contribution ID")),
            "status": status,
            "recorded_by": _clean(record.get("Recorded By")),
        }
    return None


def _search_verified_receipts(query, limit=10):
    query = _clean(query)
    if not query:
        return []
    query_upper = query.upper()
    results = []
    for record in _contribution_rows():
        receipt_number = _clean(record.get("Receipt Number"))
        status = _clean(record.get("Status")).upper()
        if not receipt_number or status in EXCLUDED_STATUSES:
            continue
        digits = "".join(character for character in receipt_number if character.isdigit())
        if query_upper not in receipt_number.upper() and not (query.isdigit() and query in digits):
            continue
        try:
            amount = float(record.get("Amount") or 0)
        except (TypeError, ValueError):
            amount = 0.0
        results.append({
            "receipt_number": receipt_number,
            "member_id": _clean(record.get("Member ID")),
            "full_name": _clean(record.get("Full Name")),
            "payment_date": _format_date(record.get("Payment Date")),
            "payment_type": _clean(record.get("Payment Type")),
            "amount": amount,
            "status": status,
        })
        if len(results) >= max(1, int(limit)):
            break
    return results


def _receipt_options():
    results = []
    for record in _contribution_rows():
        receipt_number = _clean(record.get("Receipt Number"))
        status = _clean(record.get("Status")).upper()
        if receipt_number and status not in EXCLUDED_STATUSES:
            results.append({
                "receipt_number": receipt_number,
                "full_name": _clean(record.get("Full Name")),
            })
    return results


def _documents():
    path = _data_path("spreadsheets", "08_DOCUMENTS.xlsx")
    if not path.is_file():
        return []
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            if "Documents" not in workbook.sheetnames:
                return []
            documents = []
            for row in workbook["Documents"].iter_rows(min_row=2, values_only=True):
                if not any(row):
                    continue
                document = {
                    "document_id": row[0] or "",
                    "document_name": row[1] or "",
                    "category": row[2] or "",
                    "description": row[3] or "",
                    "file_name": row[4] or "",
                    "date_added": row[5] or "",
                    "status": row[6] or "",
                }
                if str(document["status"]).strip().upper() == "ACTIVE":
                    documents.append(document)
            return documents
        finally:
            workbook.close()
    except Exception:
        return []


def _json_error(message, status):
    return JsonResponse({"success": False, "message": message}, status=status)


@require_POST
def record_receipt(request):
    try:
        data = json.loads(request.body or b"{}")
    except (json.JSONDecodeError, UnicodeDecodeError):
        data = {}
    if not isinstance(data, dict):
        data = {}

    member_id = _clean(data.get("member_id")).upper()
    amount = data.get("amount")
    payment_type = _clean(data.get("payment_type"))
    payment_reference = _clean(data.get("payment_reference"))
    description = _clean(data.get("description"))
    payment_method = _clean(data.get("payment_method", "M-PESA"))
    recorded_by = _clean(data.get("recorded_by"))
    if not member_id:
        return _json_error("Member ID is required.", 400)
    if amount is None:
        return _json_error("Payment amount is required.", 400)
    if not payment_type:
        return _json_error("Payment type is required.", 400)
    member = _member_details(member_id)
    if not member:
        return _json_error("Member ID was not found.", 404)

    try:
        now_year = datetime.now().year
        highest_sequence = 0
        for receipt in _load_receipts():
            parts = str(receipt.get("receipt_number", "")).strip().split("/")
            if len(parts) != 4 or parts[:3] != ["TLC", "RCPT", str(now_year)]:
                continue
            try:
                highest_sequence = max(highest_sequence, int(parts[3]))
            except ValueError:
                continue
        receipt_number = format_receipt_number(highest_sequence + 1, now_year)
        receipt = create_receipt(
            member_id=member["member_id"],
            member_name=member["full_name"],
            amount=amount,
            payment_type=payment_type,
            payment_reference=payment_reference,
            description=description,
            receipt_number=receipt_number,
            phone_number=member["phone_number"],
            email=member["email"],
            payment_method=payment_method,
            recorded_by=recorded_by,
        )
        receipts = _load_receipts()
        receipts.append(receipt)
        _save_receipts(receipts)
        _record_activity(
            member_id,
            "Receipt Generated",
            f"Receipt {receipt_number} was generated for KES {float(amount):,.2f}.",
            "receipt",
            {
                "receipt_number": receipt_number,
                "amount": float(amount),
                "payment_type": payment_type,
                "payment_method": payment_method,
                "payment_reference": payment_reference,
            },
        )

        email_sent = False
        email_message = "Member has no email address."
        if member.get("email"):
            try:
                from app.receipts.receipt_pdf_generator import generate_receipt_pdf
                from app.services.email_service import send_receipt_email

                pdf = generate_receipt_pdf(receipt)
                email_sent, email_message = send_receipt_email(
                    recipient_email=member["email"],
                    receipt_number=receipt_number,
                    member_name=member["full_name"],
                    amount=amount,
                    pdf_bytes=pdf,
                )
            except Exception:
                email_message = "Receipt created, but its PDF email could not be sent."
        return JsonResponse({
            "success": True,
            "message": "Receipt created successfully.",
            "receipt": receipt,
            "email_sent": email_sent,
            "email_message": email_message,
        }, status=201)
    except ValueError as error:
        return _json_error(str(error), 400)
    except Exception:
        return _json_error("Unable to create the receipt.", 500)


@require_GET
def get_member_receipts(request):
    member_id = request.session.get("member_id")
    if not member_id:
        return _json_error("Please log in to view your receipts.", 401)
    receipts = _member_receipts(member_id)
    _record_activity(
        member_id,
        "Viewed Receipts",
        "Viewed personal TLC payment receipts.",
        "receipts",
        {"receipt_count": len(receipts)},
    )
    return JsonResponse({"success": True, "receipts": receipts})


@require_GET
def get_member_receipt(request, receipt_number):
    member_id = str(request.session.get("member_id") or "").strip().upper()
    if not member_id:
        return _json_error("Please log in to view this receipt.", 401)
    receipt_number = str(receipt_number).strip()
    if member_id == ADMIN_MEMBER_ID:
        receipt = next(
            (
                _enrich_receipt(item)
                for item in _load_receipts()
                if isinstance(item, dict)
                and str(item.get("receipt_number", "")).strip() == receipt_number
            ),
            None,
        )
        if receipt is None:
            return _json_error("Receipt not found.", 404)
        _record_activity(
            member_id,
            "Viewed Receipt",
            f"Financial administrator viewed receipt {receipt_number}.",
            "receipts",
            {
                "receipt_number": receipt_number,
                "receipt_member_id": receipt.get("member_id", ""),
                "amount": receipt.get("amount", ""),
                "payment_type": receipt.get("payment_type", ""),
            },
        )
        return JsonResponse({"success": True, "receipt": receipt})

    for receipt in _member_receipts(member_id):
        if str(receipt.get("receipt_number", "")).strip() != receipt_number:
            continue
        _record_activity(
            member_id,
            "Viewed Receipt",
            f"Viewed receipt {receipt_number}.",
            "receipts",
            {"receipt_number": receipt_number, "amount": receipt.get("amount", "")},
        )
        return JsonResponse({"success": True, "receipt": receipt})
    return _json_error("Receipt not found.", 404)


@require_GET
def receipt_qr(request):
    receipt_number = request.GET.get("receipt", "").strip()
    if not receipt_number:
        return _json_error("Receipt number is required.", 400)
    try:
        import qrcode
        from io import BytesIO

        verification_url = (
            request.build_absolute_uri("/").rstrip("/")
            + "/member/receipt/"
            + receipt_number.replace("/", "%2F")
        )
        qr = qrcode.QRCode(
            version=None,
            error_correction=qrcode.constants.ERROR_CORRECT_H,
            box_size=8,
            border=4,
        )
        qr.add_data(verification_url)
        qr.make(fit=True)
        image = qr.make_image(fill_color="#173F6B", back_color="#FFFFFF")
        output = BytesIO()
        image.save(output, format="PNG")
        output.seek(0)
        response = HttpResponse(output.getvalue(), content_type="image/png")
        response["Content-Disposition"] = 'inline; filename="tlc-receipt-qr.png"'
        return response
    except Exception:
        return _json_error("Unable to generate the receipt QR code.", 500)


@require_GET
def documents_list(request):
    return JsonResponse({"success": True, "documents": _documents()})


@require_GET
def document_details(request, document_id):
    document = next(
        (doc for doc in _documents() if str(doc["document_id"]).strip() == str(document_id).strip()),
        None,
    )
    if not document:
        return _json_error("Document not found.", 404)
    return JsonResponse({"success": True, "document": document})


@require_GET
def document_file(request, file_name):
    file_name = str(file_name).strip()
    if not file_name or Path(file_name).name != file_name:
        return _json_error("Document file not found.", 404)
    if not any(str(doc.get("file_name", "")).strip() == file_name for doc in _documents()):
        return _json_error("Document file not found.", 404)
    file_path = _data_path("documents", file_name).resolve()
    documents_dir = _data_path("documents").resolve()
    if documents_dir not in file_path.parents:
        return _json_error("Document file not found.", 404)
    if not file_path.is_file():
        return _json_error("Document file is unavailable.", 404)
    member_id = request.session.get("member_id")
    if member_id:
        _record_activity(
            member_id,
            "Document Accessed",
            f"Opened the TLC document '{file_name}'.",
            "document",
            {"document_id": "", "file_name": file_name},
        )
    content_type = mimetypes.guess_type(file_name)[0] or "application/octet-stream"
    return FileResponse(file_path.open("rb"), content_type=content_type, as_attachment=False, filename=file_name)


def _authorized(request):
    return str(request.session.get("member_id") or "").strip().upper() == ADMIN_MEMBER_ID


@require_GET
def verify_receipt(request):
    if not _authorized(request):
        return _json_error("Access denied.", 403)
    receipt_number = request.GET.get("receipt", "").strip()
    if not receipt_number:
        return _json_error("Receipt number is required.", 400)
    try:
        receipt = _find_verified_receipt(receipt_number)
        if receipt is None:
            return _json_error("No valid TLC receipt was found with that reference.", 404)
        return JsonResponse({"success": True, "receipt": receipt})
    except FileNotFoundError as error:
        return _json_error(str(error), 500)
    except Exception:
        return _json_error("Unable to verify the receipt.", 500)


@require_GET
def search_receipts(request):
    if not _authorized(request):
        return _json_error("Access denied.", 403)
    try:
        results = _search_verified_receipts(request.GET.get("q", ""))
        return JsonResponse({"success": True, "results": results})
    except Exception:
        return _json_error("Unable to search receipts.", 500)


@require_GET
def receipt_options(request):
    if not _authorized(request):
        return _json_error("Access denied.", 403)
    try:
        return JsonResponse({"success": True, "results": _receipt_options()})
    except Exception as error:
        return _json_error(str(error), 500)


@require_GET
def constitution_info(request):
    if not _constitution_text():
        return _json_error("TLC Constitution is currently unavailable.", 404)
    return JsonResponse({
        "success": True,
        "name": "TLC Constitution",
        "title": "THE LEGACY CIRCLE",
        "subtitle": "OFFICIAL CONSTITUTION",
        "motto": "United in Purpose. Built to Last.",
        "founding_date": "1 September 2026",
    })


def _constitution_text():
    path = _data_path("documents", "TLC_Constitution.txt")
    try:
        return path.read_text(encoding="utf-8") if path.is_file() else ""
    except OSError:
        return ""


@require_GET
def constitution_text(request):
    text = _constitution_text()
    if not text:
        return _json_error("TLC Constitution is currently unavailable.", 404)
    return JsonResponse({"success": True, "text": text})


def _constitution_pdf(request, as_attachment):
    file_path = _data_path("documents", "TLC_Constitution.pdf")
    if not file_path.is_file():
        return _json_error("Official Constitution PDF not found.", 404)
    member_id = request.session.get("member_id")
    if member_id:
        action = "Constitution Downloaded" if as_attachment else "Constitution Opened"
        description = (
            "Downloaded the official TLC Constitution."
            if as_attachment
            else "Opened the official TLC Constitution for viewing or printing."
        )
        _record_activity(
            member_id,
            action,
            description,
            "document",
            {"document": "THE_LEGACY_CIRCLE_OFFICIAL_CONSTITUTION.pdf"},
        )
    return FileResponse(
        file_path.open("rb"),
        content_type="application/pdf",
        as_attachment=as_attachment,
        filename="THE_LEGACY_CIRCLE_OFFICIAL_CONSTITUTION.pdf",
    )


@require_GET
def download_constitution(request):
    return _constitution_pdf(request, as_attachment=True)


@require_GET
def print_constitution(request):
    return _constitution_pdf(request, as_attachment=False)


@require_GET
def receipts_page(request):
    return render(request, "portal/pages/receipts.html")


@require_GET
def receipt_details_page(request, receipt_number):
    return render(request, "portal/pages/member-receipt.html")


@require_GET
def receipt_verification_page(request):
    if not _authorized(request):
        return _json_error("Access denied.", 403)
    return render(request, "portal/pages/tlc_receipt_verification.html")
