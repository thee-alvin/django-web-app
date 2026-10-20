import base64
import hashlib
import hmac
import json
import logging
import os
import re
from datetime import datetime, timedelta
from pathlib import Path

from django.conf import settings
from django.http import JsonResponse
from django.views.decorators.http import require_GET, require_POST
from openpyxl import load_workbook


MEMBERS_FILE = Path(settings.TLC_DATA_DIR) / "spreadsheets" / "01_MEMBERS.xlsx"
ALLOWED_PROFILE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}
logger = logging.getLogger(__name__)


def _json_body(request):
    try:
        payload = json.loads(request.body or b"{}")
    except (TypeError, ValueError, UnicodeDecodeError):
        return {}
    return payload if isinstance(payload, dict) else {}


def _normalise_phone(value):
    digits = "".join(character for character in str(value or "") if character.isdigit())
    if digits.startswith("254"):
        return "0" + digits[3:]
    if len(digits) == 9:
        return "0" + digits
    return digits


def _member_rows(workbook):
    sheet = workbook["Members"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return [], []
    headers = [str(value).strip() if value is not None else "" for value in headers]
    return headers, [dict(zip(headers, row)) for row in rows]


def _member_profile_picture_path(member_id):
    safe_id = str(member_id or "").strip().upper().replace("/", "_").replace("\\", "_")
    if not safe_id:
        return ""

    directory = Path(settings.TLC_PROFILE_PICTURE_DIR)
    for extension in (".jpg", ".jpeg", ".png", ".webp"):
        if (directory / f"{safe_id}{extension}").exists():
            return f"{safe_id}{extension}"
    return ""


def _get_member_by_id(member_id):
    member_id = str(member_id or "").strip().upper()
    if not member_id or not MEMBERS_FILE.exists():
        return None

    workbook = load_workbook(MEMBERS_FILE, read_only=True, data_only=True)
    try:
        _, members = _member_rows(workbook)
        for member in members:
            stored_id = str(member.get("Member ID") or "").strip().upper()
            if stored_id != member_id:
                continue
            return {
                "member_id": stored_id,
                "full_name": str(member.get("Full Name") or "").strip(),
                "phone": str(member.get("Phone Number") or "").strip(),
                "email": str(member.get("Email Address") or "").strip(),
                "membership_status": str(member.get("Membership Status") or "").strip(),
                "registration_date": str(member.get("Registration Date") or "").strip(),
                "account_status": str(member.get("Account Status") or "").strip(),
                "last_login": str(member.get("Last Login") or "").strip(),
                "profile_picture": _member_profile_picture_path(stored_id),
            }
        return None
    finally:
        workbook.close()


def _werkzeug_password_hash(password):
    salt = base64.urlsafe_b64encode(os.urandom(16)).decode("ascii").rstrip("=")
    digest = hashlib.scrypt(
        str(password).encode("utf-8"),
        salt=salt.encode("utf-8"),
        n=32768,
        r=8,
        p=1,
        maxmem=132 * 32768 * 8,
    )
    return f"scrypt:32768:8:1${salt}${digest.hex()}"


def _check_password_hash(encoded, password):
    try:
        method, salt, expected = str(encoded).split("$", 2)
        if method.startswith("scrypt:"):
            _, n, r, p = method.split(":")
            expected_bytes = bytes.fromhex(expected)
            n, r, p = int(n), int(r), int(p)
            actual = hashlib.scrypt(
                str(password).encode("utf-8"),
                salt=salt.encode("utf-8"),
                n=n,
                r=r,
                p=p,
                dklen=len(expected_bytes),
                maxmem=132 * n * r * p,
            )
            return hmac.compare_digest(actual, expected_bytes)
        if method.startswith("pbkdf2:"):
            _, algorithm, iterations = method.split(":")
            expected_bytes = bytes.fromhex(expected)
            actual = hashlib.pbkdf2_hmac(
                algorithm,
                str(password).encode("utf-8"),
                salt.encode("utf-8"),
                int(iterations),
                dklen=len(expected_bytes),
            )
            return hmac.compare_digest(actual, expected_bytes)
    except (ValueError, TypeError, MemoryError):
        return False
    return False


def _record_activity(member_id, action, description):
    try:
        from app.security import audit_log

        activity_dir = Path(settings.TLC_DATA_DIR) / "activity"
        audit_log.ACTIVITY_DIR = activity_dir
        audit_log.ACTIVITY_FILE = activity_dir / "recent_activity.json"
        audit_log.record_activity(
            member_id,
            action=action,
            description=description,
            category="authentication" if action == "Member Login" else "profile",
        )
    except Exception:
        logger.exception("Unable to record member activity.")


def _failure(message, status):
    return JsonResponse({"success": False, "message": message}, status=status)


@require_POST
def forgot_password(request):
    data = _json_body(request)
    member_id = str(data.get("member_id", "")).strip().upper()
    phone = str(data.get("phone", "")).strip()
    email = str(data.get("email", "")).strip().lower()
    if not member_id or not phone or not email:
        return _failure("Please provide your Membership ID, phone number and email.", 400)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE, read_only=True, data_only=True)
        _, members = _member_rows(workbook)
        for member in members:
            if (
                str(member.get("Member ID") or "").strip().upper() == member_id
                and _normalise_phone(member.get("Phone Number")) == _normalise_phone(phone)
                and str(member.get("Email Address") or "").strip().lower() == email
            ):
                if not member.get("Password Hash"):
                    return _failure(
                        "This account has not been activated yet. Please use First Time Login.",
                        403,
                    )
                return JsonResponse({"success": True, "message": "Account verified successfully."})
        return _failure("The details provided do not match our TLC records.", 401)
    except Exception:
        return _failure("Unable to verify your TLC account.", 500)
    finally:
        if workbook is not None:
            workbook.close()


@require_POST
def verify_member(request):
    data = _json_body(request)
    member_id = str(data.get("member_id", "")).strip().upper()
    phone = str(data.get("phone", "")).strip()
    if not member_id or not phone:
        return _failure("Please enter your Member ID and registered phone number.", 400)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are currently unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE, read_only=True, data_only=True)
        headers, members = _member_rows(workbook)
        required = {"Member ID", "Full Name", "Phone Number", "Account Status"}
        if not required.issubset(headers):
            return _failure("TLC member spreadsheet is missing required columns.", 500)
        for member in members:
            stored_id = str(member.get("Member ID") or "").strip()
            if stored_id != member_id or str(member.get("Phone Number") or "").strip() != phone:
                continue
            if str(member.get("Account Status") or "").strip().lower() not in {"active", "pending"}:
                return _failure("This TLC membership is not eligible for account activation.", 403)
            if member.get("Password Hash"):
                return _failure(
                    "This TLC account has already been activated. Please use Member Login.",
                    409,
                )
            return JsonResponse(
                {
                    "success": True,
                    "message": "Member verified successfully.",
                    "member": {
                        "member_id": stored_id,
                        "name": str(member.get("Full Name") or ""),
                    },
                }
            )
        return _failure("The Member ID and phone number do not match our TLC records.", 401)
    except Exception:
        return _failure("Unable to verify the member record.", 500)
    finally:
        if workbook is not None:
            workbook.close()


@require_POST
def activate_account(request):
    data = _json_body(request)
    member_id = str(data.get("member_id", "")).strip().upper()
    password = str(data.get("password", ""))
    if not member_id or not password:
        return _failure("Member ID and password are required.", 400)
    if len(password) < 8:
        return _failure("Password must be at least 8 characters.", 400)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are currently unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE)
        sheet = workbook["Members"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        required = {"Member ID", "Password Hash", "Account Status"}
        if not required.issubset(headers):
            return _failure("TLC member spreadsheet is missing required columns.", 500)
        member_id_column = headers.index("Member ID") + 1
        password_column = headers.index("Password Hash") + 1
        status_column = headers.index("Account Status") + 1
        for row_number in range(2, sheet.max_row + 1):
            current_id = str(sheet.cell(row=row_number, column=member_id_column).value or "").strip()
            if current_id != member_id:
                continue
            if sheet.cell(row=row_number, column=password_column).value:
                return _failure("This TLC account has already been activated. Please log in.", 409)
            status = str(sheet.cell(row=row_number, column=status_column).value or "").strip().lower()
            if status not in {"active", "pending"}:
                return _failure("This member account is not eligible for activation.", 403)
            sheet.cell(row=row_number, column=password_column).value = _werkzeug_password_hash(password)
            sheet.cell(row=row_number, column=status_column).value = "active"
            workbook.save(MEMBERS_FILE)
            return JsonResponse({"success": True, "message": "TLC account activated successfully."})
        return _failure("Member ID was not found.", 404)
    except PermissionError:
        return _failure("The member spreadsheet is currently open. Please close it and try again.", 423)
    except Exception:
        return _failure("Unable to activate the TLC account.", 500)
    finally:
        if workbook is not None:
            workbook.close()


@require_POST
def member_login(request):
    data = _json_body(request)
    member_id = str(data.get("member_id", "")).strip().upper()
    password = str(data.get("password", ""))
    if not member_id or not password:
        return _failure("Please enter your Member ID and password.", 400)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are currently unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE, read_only=True, data_only=True)
        _, members = _member_rows(workbook)
        for member in members:
            current_id = str(member.get("Member ID") or "").strip()
            if current_id != member_id:
                continue
            account_status = str(member.get("Account Status") or "").strip().lower()
            password_hash = member.get("Password Hash")
            if not password_hash:
                return _failure(
                    "Your account has not been activated yet. Please use First Time Login.",
                    403,
                )
            if account_status != "active":
                return _failure("Your TLC account is not active.", 403)
            if _check_password_hash(password_hash, password):
                request.session.flush()
                request.session["member_id"] = current_id.strip().upper()
                request.session["role"] = "MEMBER"
                _record_activity(current_id, "Member Login", "Signed in to the TLC Member Portal.")
                return JsonResponse(
                    {
                        "success": True,
                        "message": "Login successful.",
                        "member": {
                            "member_id": current_id,
                            "name": str(member.get("Full Name") or ""),
                        },
                    }
                )
            break
        return _failure("Invalid Member ID or password.", 401)
    except Exception:
        return _failure("Unable to process the login.", 500)
    finally:
        if workbook is not None:
            workbook.close()


@require_POST
def reset_password(request):
    data = _json_body(request)
    member_id = str(data.get("member_id", "")).strip().upper()
    password = str(data.get("password", ""))
    if not member_id or not password:
        return _failure("Member ID and new password are required.", 400)
    if len(password) < 8:
        return _failure("Password must be at least 8 characters.", 400)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE)
        sheet = workbook["Members"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        if "Member ID" not in headers or "Password Hash" not in headers:
            return _failure("Required password columns are missing.", 500)
        id_column = headers.index("Member ID") + 1
        password_column = headers.index("Password Hash") + 1
        for row_number in range(2, sheet.max_row + 1):
            stored_id = str(sheet.cell(row=row_number, column=id_column).value or "").strip()
            if stored_id != member_id:
                continue
            sheet.cell(row=row_number, column=password_column).value = _werkzeug_password_hash(password)
            workbook.save(MEMBERS_FILE)
            return JsonResponse({"success": True, "message": "Password reset successfully."})
        return _failure("Member ID was not found.", 404)
    except PermissionError:
        return _failure("Please close the member spreadsheet and try again.", 423)
    except Exception:
        return _failure("Unable to reset the password.", 500)
    finally:
        if workbook is not None:
            workbook.close()


@require_POST
def member_logout(request):
    request.session.flush()
    return JsonResponse({"success": True, "message": "Logged out successfully."})


def _session_member_id(request):
    return str(request.session.get("member_id") or "").strip().upper()


@require_GET
def member_profile(request):
    member_id = _session_member_id(request)
    if not member_id:
        return _failure("No active member session.", 401)
    member = _get_member_by_id(member_id)
    if not member:
        return _failure("Member record could not be found.", 404)
    return JsonResponse({"success": True, "member": member})


@require_GET
def get_profile(request):
    member_id = _session_member_id(request)
    if not member_id:
        return _failure("Please log in to view your profile.", 401)
    member = _get_member_by_id(member_id)
    if not member:
        return _failure("Your TLC member record could not be found.", 404)
    return JsonResponse({"success": True, "member": member})


def _profile_picture_filename(member_id, extension):
    return str(member_id).strip().upper().replace("/", "_").replace("\\", "_") + extension


@require_POST
def update_profile(request):
    member_id = _session_member_id(request)
    if not member_id:
        return _failure("Please log in to update your profile.", 401)
    if not MEMBERS_FILE.exists():
        return _failure("TLC member records are unavailable.", 500)

    workbook = None
    try:
        workbook = load_workbook(MEMBERS_FILE)
        sheet = workbook["Members"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        required = {"Member ID", "Full Name", "Phone Number", "Email Address"}
        missing = sorted(required.difference(headers))
        if missing:
            return _failure("Required member columns are missing: " + ", ".join(missing), 500)
        columns = {name: headers.index(name) + 1 for name in headers}
        target_row = None
        for row_number in range(2, sheet.max_row + 1):
            stored_id = str(sheet.cell(row=row_number, column=columns["Member ID"]).value or "").strip().upper()
            if stored_id == member_id:
                target_row = row_number
                break
        if target_row is None:
            return _failure("Your TLC member record was not found.", 404)

        full_name = str(request.POST.get("full_name") or "").strip()
        phone = str(request.POST.get("phone") or "").strip()
        email = str(request.POST.get("email") or "").strip().lower()
        if not full_name:
            return _failure("Full name cannot be empty.", 400)
        if not phone:
            return _failure("Phone number cannot be empty.", 400)
        if not email or "@" not in email:
            return _failure("Please enter a valid email address.", 400)

        sheet.cell(row=target_row, column=columns["Full Name"]).value = full_name
        sheet.cell(row=target_row, column=columns["Phone Number"]).value = phone
        sheet.cell(row=target_row, column=columns["Email Address"]).value = email

        photo = request.FILES.get("profile_picture")
        if photo and photo.name:
            filename = Path(photo.name.replace("\\", "/")).name
            extension = Path(filename).suffix.lower()
            if not filename or extension not in ALLOWED_PROFILE_EXTENSIONS:
                return _failure("Use JPG, JPEG, PNG or WEBP for your profile picture.", 400)
            picture_dir = Path(settings.TLC_PROFILE_PICTURE_DIR)
            picture_dir.mkdir(parents=True, exist_ok=True)
            safe_id = member_id.replace("/", "_").replace("\\", "_")
            for old_extension in ALLOWED_PROFILE_EXTENSIONS:
                old_picture = picture_dir / f"{safe_id}{old_extension}"
                if old_picture.exists():
                    old_picture.unlink()
            destination = picture_dir / _profile_picture_filename(member_id, extension)
            with destination.open("wb") as output:
                for chunk in photo.chunks():
                    output.write(chunk)

        workbook.save(MEMBERS_FILE)
        _record_activity(
            member_id,
            "Profile Updated",
            "Your TLC member profile was updated successfully.",
        )
        return JsonResponse(
            {
                "success": True,
                "message": "Your TLC profile was updated successfully.",
                "member": _get_member_by_id(member_id),
            }
        )
    except PermissionError:
        return _failure(
            "The member spreadsheet is currently open. Please close it and try again.",
            423,
        )
    except Exception:
        return _failure("Unable to update your TLC profile.", 500)
    finally:
        if workbook is not None:
            workbook.close()


def _configure_financial_data_paths():
    data_dir = Path(settings.TLC_DATA_DIR)
    spreadsheets = data_dir / "spreadsheets"

    from app.financial_register import ledger_service
    from app.financial_register import service as financial_service
    from app.financial_register.contribution_coverage import service as coverage_service
    from app.financial_register.contribution_coverage import workbook as coverage_workbook
    from app.financial_register.disbursements import service as disbursement_service

    ledger_service.MEMBERS_FILE = spreadsheets / "01_MEMBERS.xlsx"
    ledger_service.CONTRIBUTIONS_FILE = spreadsheets / "02_CONTRIBUTIONS.xlsx"
    financial_service.MEMBERS_FILE = spreadsheets / "01_MEMBERS.xlsx"
    financial_service.CONTRIBUTIONS_FILE = spreadsheets / "02_CONTRIBUTIONS.xlsx"
    disbursement_service.DISBURSEMENTS_FILE = spreadsheets / "11_TREASURY_DISBURSEMENTS.xlsx"
    coverage_workbook.COVERAGE_FILE = spreadsheets / "03_CONTRIBUTION_COVERAGE.xlsx"
    coverage_service.COVERAGE_FILE = coverage_workbook.COVERAGE_FILE
    coverage_service.CONTRIBUTIONS_FILE = spreadsheets / "02_CONTRIBUTIONS.xlsx"


_DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%d %B %Y",
    "%d %b %Y",
    "%B %d, %Y",
    "%b %d, %Y",
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%dT%H:%M:%S.%f",
)


def _clean(value):
    return str(value or "").strip()


def _number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(match.group()) if match else 0.0


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        try:
            return value
        except Exception:
            pass
    text = _clean(value)
    if not text:
        return None
    for date_format in _DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    match = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    if match:
        for date_format in ("%d %B %Y", "%d %b %Y"):
            try:
                return datetime.strptime(match.group(0), date_format).date()
            except ValueError:
                continue
    return None


def _transaction_date(transaction):
    for key in ("date", "payment_date", "transaction_date", "received_date", "created_at", "received_at"):
        if key in transaction:
            parsed = _parse_date(transaction.get(key))
            if parsed:
                return parsed
    return None


def _transaction_text(transaction):
    return " ".join(
        str(transaction[key])
        for key in ("description", "payment_type", "type", "transaction_type", "details", "notes", "remarks")
        if transaction.get(key)
    )


def _extract_week_count(transaction):
    for key in ("weeks", "week_count", "coverage_weeks", "number_of_weeks"):
        value = transaction.get(key)
        if value is not None:
            try:
                count = int(float(str(value).strip()))
                if count > 0:
                    return count
            except (TypeError, ValueError):
                pass
    match = re.search(r"(\d+)\s*[- ]?\s*WEEKS?", _transaction_text(transaction), flags=re.IGNORECASE)
    return int(match.group(1)) if match else 0


def _extract_transactions(statement):
    member = statement.get("member") or {}
    transactions = member.get("transactions")
    if isinstance(transactions, list):
        return transactions
    transactions = statement.get("transactions")
    return transactions if isinstance(transactions, list) else []


def _build_coverage(all_year_statement, year, month):
    if not year or not month:
        return {
            "year": None,
            "month": None,
            "month_label": "",
            "period": "",
            "amount": 0.0,
            "week_count": 0,
            "weeks_catered_for": 0,
            "weeks": [],
            "coverageWeeks": [],
            "coverageAmount": 0.0,
            "coverageWeekCount": 0,
        }

    target_year, target_month = int(year), int(month)
    coverage_rows = []
    for transaction in _extract_transactions(all_year_statement):
        if not isinstance(transaction, dict):
            continue
        payment_date = _transaction_date(transaction)
        if not payment_date or payment_date.year not in {target_year, target_year - 1}:
            continue
        week_count = _extract_week_count(transaction)
        if week_count <= 0:
            continue
        total_amount = _number(
            transaction.get("amount")
            or transaction.get("amount_paid")
            or transaction.get("paid_amount")
            or transaction.get("value")
        )
        if total_amount <= 0:
            continue
        for index in range(week_count):
            covered_date = payment_date + timedelta(weeks=index)
            if covered_date.year == target_year and covered_date.month == target_month:
                formatted_date = covered_date.strftime("%d %B %Y")
                formatted_payment = payment_date.strftime("%d %B %Y")
                coverage_rows.append(
                    {
                        "week": index + 1,
                        "week_number": index + 1,
                        "label": f"WEEK {index + 1}",
                        "date": formatted_date,
                        "coverage_date": formatted_date,
                        "payment_received": formatted_payment,
                        "payment_received_date": formatted_payment,
                        "amount": round(total_amount / week_count, 2),
                        "description": f"Payment received {formatted_payment}",
                        "source_amount": round(total_amount, 2),
                        "source_week_count": week_count,
                    }
                )
    coverage_rows.sort(key=lambda row: _parse_date(row["coverage_date"]) or datetime.min.date())
    monthly_amount = round(sum(row["amount"] for row in coverage_rows), 2)
    week_count = len(coverage_rows)
    month_label = f"{target_year}-{target_month:02d}"
    return {
        "year": target_year,
        "month": target_month,
        "month_label": month_label,
        "period": month_label,
        "amount": monthly_amount,
        "coverageAmount": monthly_amount,
        "week_count": week_count,
        "coverageWeekCount": week_count,
        "weeks_catered_for": week_count,
        "weeks": coverage_rows,
        "coverageWeeks": coverage_rows,
    }


@require_GET
def member_payment_statement(request):
    member_id = _session_member_id(request)
    if not member_id:
        return _failure("Please log in to view your payment statement.", 401)
    year = _clean(request.GET.get("year"))
    month = _clean(request.GET.get("month"))
    if year and (not year.isdigit() or len(year) != 4):
        return _failure("Invalid year.", 400)
    if month:
        if not year:
            return _failure("A month requires a year.", 400)
        if not month.isdigit() or not 1 <= int(month) <= 12:
            return _failure("Invalid month.", 400)
        month = str(int(month)).zfill(2)

    try:
        _configure_financial_data_paths()
        from app.financial_register.statements.service import build_member_statement

        statement = build_member_statement(member_id=member_id, year=year or None, month=month or None)
        if statement is None:
            return _failure("Your member statement was not found.", 404)
        history_statement = (
            build_member_statement(member_id=member_id, year=year, month=None)
            if year and month
            else statement
        )
        coverage = _build_coverage(history_statement, year, month)
        if year and month and coverage.get("weeks"):
            covered_dates = {_parse_date(row.get("coverage_date")) for row in coverage["weeks"]}
            covered_dates.discard(None)
            source_transactions = []
            for transaction in _extract_transactions(history_statement):
                payment_date = _transaction_date(transaction)
                if not payment_date:
                    continue
                week_count = _extract_week_count(transaction)
                if week_count <= 0:
                    continue
                coverage_end = payment_date + timedelta(weeks=week_count) - timedelta(days=1)
                if any(payment_date <= date <= coverage_end for date in covered_dates):
                    source_transactions.append(transaction)

            member = statement.setdefault("member", {})
            current = member.get("transactions")
            if not isinstance(current, list):
                current = []

            def transaction_key(transaction):
                return (
                    _clean(
                        transaction.get("receipt")
                        or transaction.get("receipt_number")
                        or transaction.get("reference")
                        or transaction.get("payment_reference")
                    ),
                    str(_transaction_date(transaction) or ""),
                    round(
                        _number(
                            transaction.get("amount")
                            or transaction.get("amount_paid")
                            or transaction.get("paid_amount")
                            or transaction.get("value")
                        ),
                        2,
                    ),
                    _transaction_text(transaction),
                )

            existing = {transaction_key(transaction) for transaction in current}
            for transaction in source_transactions:
                key = transaction_key(transaction)
                if key not in existing:
                    current.append(transaction)
                    existing.add(key)
            current.sort(key=lambda transaction: _transaction_date(transaction) or datetime.min.date(), reverse=True)
            member["transactions"] = current

        statement["coverage"] = coverage
        statement["coverageWeeks"] = coverage["coverageWeeks"]
        statement["coverageAmount"] = coverage["coverageAmount"]
        statement["coverageWeekCount"] = coverage["coverageWeekCount"]
        return JsonResponse({"success": True, "statement": statement})
    except ValueError as error:
        return _failure(str(error), 400)
    except Exception:
        return _failure("Unable to load your payment statement.", 500)
