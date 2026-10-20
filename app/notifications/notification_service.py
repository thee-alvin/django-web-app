from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]
NOTIFICATIONS_WORKBOOK = BASE_DIR / "data" / "spreadsheets" / "09_NOTIFICATIONS.xlsx"


def get_notifications():
    if not NOTIFICATIONS_WORKBOOK.exists():
        return []

    try:
        wb = load_workbook(NOTIFICATIONS_WORKBOOK, data_only=True)

        if "Notifications" not in wb.sheetnames:
            return []

        ws = wb["Notifications"]
        notifications = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue

            item = {
                "notification_id": row[0] or "",
                "title": row[1] or "",
                "message": row[2] or "",
                "type": row[3] or "General",
                "priority": row[4] or "Normal",
                "target": row[5] or "All Members",
                "date": row[6] or "",
                "time": row[7] or "",
                "status": row[8] or "",
            }

            if str(item["status"]).strip().upper() == "ACTIVE":
                notifications.append(item)

        notifications.sort(
            key=lambda x: f"{x['date']} {x['time']}",
            reverse=True
        )

        return notifications

    except Exception:
        return []


def get_notification(notification_id):
    notification_id = str(notification_id).strip()

    for item in get_notifications():
        if str(item["notification_id"]).strip() == notification_id:
            return item

    return None
