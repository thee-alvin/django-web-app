from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

from app.elections.results_service import (
    calculate_and_store_results
)


BASE_DIR = Path(__file__).resolve().parents[2]

ELECTIONS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "06_ELECTIONS.xlsx"
)


def clean(value):
    return str(value or "").strip()


def get_election():
    wb = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Elections"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        for row in rows:
            if not any(x is not None for x in row):
                continue

            return {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

        return None

    finally:
        wb.close()


def close_election(election_id=None):

    election = get_election()

    if not election:
        raise ValueError("No election is configured.")

    election_id = (
        election_id
        or clean(election.get("Election ID"))
    )

    current_status = clean(
        election.get("Status")
    ).upper()

    if current_status == "ELECTION CLOSED":
        return {
            "success": True,
            "message": "Election is already closed.",
            "results": []
        }

    if current_status != "VOTING OPEN":
        raise ValueError(
            f"Election cannot be closed from status: {current_status}"
        )

    wb = load_workbook(ELECTIONS_FILE)

    try:
        ws = wb["Elections"]

        headers = [
            clean(cell.value)
            for cell in ws[1]
        ]

        election_id_col = (
            headers.index("Election ID") + 1
        )

        status_col = (
            headers.index("Status") + 1
        )

        results_status_col = (
            headers.index("Results Status") + 1
        )

        target_row = None

        for row in range(2, ws.max_row + 1):
            stored_id = clean(
                ws.cell(
                    row=row,
                    column=election_id_col
                ).value
            )

            if stored_id == election_id:
                target_row = row
                break

        if target_row is None:
            raise ValueError(
                "Election could not be found."
            )

        ws.cell(
            target_row,
            status_col
        ).value = "ELECTION CLOSED"

        ws.cell(
            target_row,
            results_status_col
        ).value = "CALCULATING"

        wb.save(ELECTIONS_FILE)

    finally:
        wb.close()

    results = calculate_and_store_results(
        election_id
    )

    wb = load_workbook(ELECTIONS_FILE)

    try:
        ws = wb["Elections"]

        ws.cell(
            target_row,
            results_status_col
        ).value = "READY"

        audit_ws = wb["Election Audit"]

        audit_ws.append([
            "TLC-AUDIT-" + datetime.now().strftime("%Y%m%d%H%M%S%f"),
            election_id,
            "ELECTION CLOSED",
            "Voting period closed and final results calculated.",
            "SYSTEM",
            datetime.now()
        ])

        wb.save(ELECTIONS_FILE)

    finally:
        wb.close()

    return {
        "success": True,
        "message": "Election closed and final results calculated.",
        "results": results
    }
