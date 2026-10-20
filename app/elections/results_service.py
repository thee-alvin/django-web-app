from pathlib import Path
from datetime import datetime
from collections import Counter
from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR

BASE_DIR = Path(__file__).resolve().parents[2]
ELECTIONS_FILE = DATA_DIR / "spreadsheets" / "06_ELECTIONS.xlsx"
CANDIDATES_FILE = DATA_DIR / "spreadsheets" / "07_CANDIDATES.xlsx"


def clean(value):
    return str(value or "").strip()


def load_records(path, sheet_name):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        records = []
        for row in rows:
            if not any(x is not None for x in row):
                continue
            records.append({
                h: row[i] if i < len(row) else None
                for i, h in enumerate(headers)
                if h
            })
        return records
    finally:
        wb.close()


def get_active_election():
    elections = load_records(ELECTIONS_FILE, "Elections")
    return elections[0] if elections else None


def calculate_results(election_id=None):
    election = get_active_election()
    if not election:
        raise ValueError("No election is configured.")

    election_id = election_id or clean(election.get("Election ID"))
    if not election_id:
        raise ValueError("Election ID is missing.")

    positions = [
        r for r in load_records(ELECTIONS_FILE, "Election Positions")
        if clean(r.get("Election ID")) == election_id
    ]

    candidates = [
        r for r in load_records(CANDIDATES_FILE, "Candidates")
        if clean(r.get("Election ID")) == election_id
        and clean(r.get("Candidate Status")).upper() == "APPROVED"
    ]

    votes = [
        r for r in load_records(ELECTIONS_FILE, "Voting Records")
        if clean(r.get("Election ID")) == election_id
        and clean(r.get("Validation Status")).upper() == "VALID"
    ]

    counts = Counter(
        (
            clean(v.get("Position ID")),
            clean(v.get("Candidate ID"))
        )
        for v in votes
    )

    now = datetime.now()
    results = []

    for position in sorted(
        positions,
        key=lambda r: int(r.get("Position Order") or 999999)
    ):
        position_id = clean(position.get("Position ID"))
        position_name = clean(position.get("Position Name"))

        position_candidates = [
            c for c in candidates
            if clean(c.get("Position ID")) == position_id
        ]

        total_votes = sum(
            counts[(position_id, clean(c.get("Candidate ID")))]
            for c in position_candidates
        )

        vote_values = [
            counts[(position_id, clean(c.get("Candidate ID")))]
            for c in position_candidates
        ]

        highest = max(vote_values, default=0)

        winners = [
            clean(c.get("Candidate ID"))
            for c in position_candidates
            if counts[(position_id, clean(c.get("Candidate ID")))] == highest
            and highest > 0
        ]

        tie = len(winners) > 1

        for candidate in position_candidates:
            candidate_id = clean(candidate.get("Candidate ID"))
            vote_count = counts[(position_id, candidate_id)]
            percentage = (
                round((vote_count / total_votes) * 100, 2)
                if total_votes else 0.0
            )

            results.append({
                "Election ID": election_id,
                "Position ID": position_id,
                "Position Name": position_name,
                "Candidate ID": candidate_id,
                "Candidate Name": clean(candidate.get("Full Name")) or "Unknown Candidate",
                "Vote Count": vote_count,
                "Total Votes": total_votes,
                "Percentage": percentage,
                "Result Status": "TIE" if tie else "CALCULATED",
                "Winner": "NO" if tie else ("YES" if candidate_id in winners else "NO"),
                "Calculated At": now,
            })

    return results


def calculate_and_store_results(election_id=None):
    election = get_active_election()
    election_id = election_id or clean(election.get("Election ID"))

    results = calculate_results(election_id)

    wb = load_workbook(ELECTIONS_FILE)
    try:
        ws = wb["Election Results"]
        headers = [clean(c.value) for c in ws[1]]

        required = [
            "Election ID", "Position ID", "Candidate ID",
            "Candidate Name", "Vote Count", "Percentage",
            "Result Status", "Winner", "Calculated At"
        ]

        missing = [h for h in required if h not in headers]
        if missing:
            raise ValueError(
                "Missing Election Results columns: " + ", ".join(missing)
            )

        election_col = headers.index("Election ID") + 1

        for row in range(ws.max_row, 1, -1):
            if clean(ws.cell(row, election_col).value) == election_id:
                ws.delete_rows(row)

        for result in results:
            ws.append([
                result["Election ID"],
                result["Position ID"],
                result["Candidate ID"],
                result["Candidate Name"],
                result["Vote Count"],
                result["Percentage"],
                result["Result Status"],
                result["Winner"],
                result["Calculated At"],
            ])

        wb.save(ELECTIONS_FILE)
        return results
    finally:
        wb.close()
