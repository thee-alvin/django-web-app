from pathlib import Path
from openpyxl import load_workbook

CANDIDATES_FILE = (
    Path(__file__).resolve().parents[2]
    / "data"
    / "spreadsheets"
    / "07_CANDIDATES.xlsx"
)

def clean(value):
    return str(value or "").strip()

def get_candidate_status(election_id, member_id, position_id):
    wb = load_workbook(
        CANDIDATES_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Candidates"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        for row in rows:
            if not any(x is not None for x in row):
                continue

            record = {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

            if (
                clean(record.get("Election ID")) == election_id
                and clean(record.get("Member ID")).upper()
                    == member_id.upper()
                and clean(record.get("Position ID"))
                    == position_id
            ):
                return record

        return None

    finally:
        wb.close()
