from pathlib import Path
from datetime import datetime
import uuid

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

ELECTIONS_FILE = BASE_DIR / "data" / "spreadsheets" / "06_ELECTIONS.xlsx"
CANDIDATES_FILE = BASE_DIR / "data" / "spreadsheets" / "07_CANDIDATES.xlsx"


def clean(value):
    return str(value or "").strip()


def get_active_election():
    wb = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Elections"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        for row in rows:
            if not any(x is not None for x in row):
                continue

            return {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

        return None

    finally:
        wb.close()


def get_positions(election_id):
    wb = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Election Positions"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        positions = []

        for row in rows:
            if not any(x is not None for x in row):
                continue

            record = {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

            if (
                clean(record.get("Election ID")) == election_id
                and clean(record.get("Status")).upper() == "ACTIVE"
            ):
                positions.append(record)

        return positions

    finally:
        wb.close()


def member_already_applied(member_id, election_id, position_id):
    wb = load_workbook(
        CANDIDATES_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Candidates"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        for row in rows:
            if not any(x is not None for x in row):
                continue

            record = {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

            if (
                clean(record.get("Election ID")) == election_id
                and clean(record.get("Member ID")).upper()
                == member_id.upper()
                and clean(record.get("Position ID"))
                == position_id
                and clean(record.get("Candidate Status")).upper()
                in ("PENDING", "APPROVED")
            ):
                return True

        return False

    finally:
        wb.close()


def submit_candidate_application(
    member_id,
    full_name,
    position_id,
    manifesto
):
    if not member_id:
        return False, "Authentication required.", None

    if not full_name:
        return False, "Member name is required.", None

    if not position_id:
        return False, "Please select a position.", None

    if not manifesto:
        return False, "Please enter your manifesto.", None

    election = get_active_election()

    if not election:
        return False, "No election is configured.", None

    election_id = clean(
        election.get("Election ID")
    )

    nomination_start = election.get(
        "Nomination Start"
    )

    nomination_end = election.get(
        "Nomination End"
    )

    now = datetime.now()

    if nomination_start and now < nomination_start:
        return False, "Candidate applications are not yet open.", None

    if nomination_end and now > nomination_end:
        return False, "Candidate applications are closed.", None

    positions = get_positions(election_id)

    position = next(
        (
            p for p in positions
            if clean(p.get("Position ID")) == position_id
        ),
        None
    )

    if not position:
        return False, "The selected position is not available.", None

    if member_already_applied(
        member_id,
        election_id,
        position_id
    ):
        return (
            False,
            "You have already submitted an application for this position.",
            None
        )

    candidate_id = (
        "TLC-CAN-"
        + uuid.uuid4().hex.upper()
    )

    submitted_at = datetime.now()

    candidates_wb = load_workbook(
        CANDIDATES_FILE
    )

    elections_wb = load_workbook(
        ELECTIONS_FILE
    )

    try:
        candidates_ws = candidates_wb["Candidates"]

        candidates_ws.append([
            candidate_id,
            election_id,
            member_id,
            full_name,
            position_id,
            clean(position.get("Position Name")),
            manifesto,
            "PENDING",
            submitted_at,
            "",
            submitted_at
        ])

        candidates_wb.save(
            CANDIDATES_FILE
        )

        audit_ws = elections_wb["Election Audit"]

        audit_ws.append([
            "TLC-AUDIT-" + uuid.uuid4().hex.upper(),
            election_id,
            "CANDIDATE APPLICATION SUBMITTED",
            (
                f"Candidate application submitted for "
                f"{clean(position.get('Position Name'))}."
            ),
            member_id,
            submitted_at
        ])

        elections_wb.save(
            ELECTIONS_FILE
        )

    finally:
        candidates_wb.close()
        elections_wb.close()

    return (
        True,
        "Candidate application submitted successfully.",
        candidate_id
    )
