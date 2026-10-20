from pathlib import Path
from flask import Blueprint, jsonify, session
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parents[2]
CANDIDATES_FILE = BASE_DIR / "data" / "spreadsheets" / "07_CANDIDATES.xlsx"

candidate_status_bp = Blueprint(
    "candidate_status",
    __name__
)


def clean(value):
    return str(value or "").strip()


@candidate_status_bp.route(
    "/api/member/elections/my-applications",
    methods=["GET"]
)
def my_applications():

    member_id = clean(
        session.get("member_id")
    )

    if not member_id:
        return jsonify({
            "success": False,
            "message": "Authentication required."
        }), 401

    wb = load_workbook(
        CANDIDATES_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Candidates"]
        rows = ws.iter_rows(values_only=True)

        headers = [
            clean(value)
            for value in next(rows, ())
        ]

        applications = []

        for row in rows:

            if not any(value is not None for value in row):
                continue

            record = {
                headers[index]: row[index]
                for index in range(
                    min(len(headers), len(row))
                )
                if headers[index]
            }

            if clean(
                record.get("Member ID")
            ).upper() != member_id.upper():
                continue

            applications.append({
                "Candidate ID": clean(
                    record.get("Candidate ID")
                ),
                "Member ID": clean(
                    record.get("Member ID")
                ),
                "Election ID": clean(
                    record.get("Election ID")
                ),
                "Position ID": clean(
                    record.get("Position ID")
                ),
                "Position": clean(
                    record.get("Position")
                ),
                "Full Name": clean(
                    record.get("Full Name")
                ),
                "Manifesto": clean(
                    record.get("Manifesto")
                ),
                "Candidate Status": clean(
                    record.get("Candidate Status")
                ),
                "Nomination Date": (
                    record.get("Nomination Date").strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                    if hasattr(
                        record.get("Nomination Date"),
                        "strftime"
                    )
                    else clean(
                        record.get("Nomination Date")
                    )
                )
            })

        return jsonify({
            "success": True,
            "applications": applications
        })

    finally:
        wb.close()
