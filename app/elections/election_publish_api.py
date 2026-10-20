from datetime import datetime
from pathlib import Path
import uuid

from flask import Blueprint, jsonify, session
from openpyxl import load_workbook

from app.elections.results_service import calculate_and_store_results


BASE_DIR = Path(__file__).resolve().parents[2]
ELECTIONS_FILE = BASE_DIR / "data" / "spreadsheets" / "06_ELECTIONS.xlsx"

election_publish_bp = Blueprint(
    "election_publish",
    __name__
)


def clean(value):
    return str(value or "").strip()


def get_election():
    wb = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        ws = wb["Elections"]
        rows = ws.iter_rows(values_only=True)
        headers = [clean(x) for x in next(rows, ())]

        for row in rows:
            if any(x is not None for x in row):
                return {
                    headers[i]: row[i]
                    for i in range(min(len(headers), len(row)))
                    if headers[i]
                }

        return None
    finally:
        wb.close()


@election_publish_bp.route(
    "/api/admin/elections/publish-results",
    methods=["POST"]
)
def publish_results():

    if not session.get("member_id"):
        return jsonify({
            "success": False,
            "message": "Authentication required."
        }), 401

    election = get_election()

    if not election:
        return jsonify({
            "success": False,
            "message": "No election is configured."
        }), 404

    election_id = clean(
        election.get("Election ID")
    )

    status = clean(
        election.get("Status")
    ).upper()

    results_status = clean(
        election.get("Results Status")
    ).upper()

    if status != "ELECTION CLOSED":
        return jsonify({
            "success": False,
            "message": "Election must be closed before results can be published."
        }), 400

    if results_status != "READY":
        return jsonify({
            "success": False,
            "message": "Final results are not ready for publication."
        }), 400

    results = calculate_and_store_results(
        election_id
    )

    wb = load_workbook(ELECTIONS_FILE)

    try:
        ws = wb["Elections"]

        headers = [
            clean(cell.value)
            for cell in ws[1]
        ]

        election_col = headers.index(
            "Election ID"
        ) + 1

        results_status_col = headers.index(
            "Results Status"
        ) + 1

        target_row = None

        for row in range(2, ws.max_row + 1):
            if clean(
                ws.cell(
                    row=row,
                    column=election_col
                ).value
            ) == election_id:
                target_row = row
                break

        if target_row is None:
            raise ValueError(
                "Election could not be found."
            )

        ws.cell(
            target_row,
            results_status_col
        ).value = "PUBLISHED"

        audit_ws = wb["Election Audit"]

        audit_ws.append([
            "TLC-AUDIT-" + uuid.uuid4().hex.upper(),
            election_id,
            "RESULTS PUBLISHED",
            "Final election results were published.",
            "SYSTEM",
            datetime.now()
        ])

        wb.save(ELECTIONS_FILE)

    finally:
        wb.close()

    return jsonify({
        "success": True,
        "message": "Final election results published successfully.",
        "results": results
    })
