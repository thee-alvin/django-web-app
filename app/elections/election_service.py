
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

ELECTIONS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "06_ELECTIONS.xlsx"
)

CANDIDATES_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "07_CANDIDATES.xlsx"
)


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def _read_elections():
    workbook = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Elections"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [_clean(value) for value in headers]

        records = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            record = {}

            for index, header in enumerate(headers):
                if not header:
                    continue

                value = row[index] if index < len(row) else None

                if hasattr(value, "strftime"):
                    if header in (
                        "Nomination Start",
                        "Nomination End",
                        "Voting Start",
                        "Voting End"
                    ):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    elif header == "Created At":
                        value = value.strftime("%Y-%m-%d %H:%M:%S")

                record[header] = value if value is not None else ""

            records.append(record)

        return records

    finally:
        workbook.close()


def _read_positions():
    workbook = load_workbook(
        ELECTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Election Positions"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [_clean(value) for value in headers]

        records = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            record = {}

            for index, header in enumerate(headers):
                if not header:
                    continue

                value = row[index] if index < len(row) else None
                record[header] = value if value is not None else ""

            records.append(record)

        return records

    finally:
        workbook.close()


def _read_candidates():
    workbook = load_workbook(
        CANDIDATES_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Candidates"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [_clean(value) for value in headers]

        records = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            record = {}

            for index, header in enumerate(headers):
                if not header:
                    continue

                value = row[index] if index < len(row) else None

                if hasattr(value, "strftime"):
                    value = value.strftime("%Y-%m-%d %H:%M:%S")

                record[header] = value if value is not None else ""

            records.append(record)

        return records

    finally:
        workbook.close()


def get_election_overview():
    elections = _read_elections()
    positions = _read_positions()
    candidates = _read_candidates()

    return {
        "elections": elections,
        "positions": positions,
        "candidates": candidates
    }
