from pathlib import Path
from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]
DOCUMENTS_WORKBOOK = DATA_DIR / "spreadsheets" / "08_DOCUMENTS.xlsx"


def get_documents():
    """
    Read active documents from the Documents workbook.
    """

    if not DOCUMENTS_WORKBOOK.exists():
        return []

    try:
        wb = load_workbook(DOCUMENTS_WORKBOOK, data_only=True)

        if "Documents" not in wb.sheetnames:
            return []

        ws = wb["Documents"]

        documents = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue

            document = {
                "document_id": row[0] or "",
                "document_name": row[1] or "",
                "category": row[2] or "",
                "description": row[3] or "",
                "file_name": row[4] or "",
                "date_added": row[5] or "",
                "status": row[6] or "",
            }

            if str(document["status"]).strip().upper() == "ACTIVE":
                documents.append(document)

        return documents

    except Exception:
        return []


def get_document(document_id):
    """
    Find one active document by its Document ID.
    """

    document_id = str(document_id).strip()

    for document in get_documents():
        if str(document["document_id"]).strip() == document_id:
            return document

    return None
