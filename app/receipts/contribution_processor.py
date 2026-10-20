from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook

from app.receipts.receipt_number import format_receipt_number

from app.receipts.receipt_storage import (
    get_member_details,
    load_receipts,
    add_receipt,
)

from app.receipts.receipt_generator import create_receipt

from app.receipts.receipt_pdf_generator import (
    generate_receipt_pdf
)

from app.services.email_service import (
    send_receipt_email
)

from app.security.audit_log import record_activity

from app.receipts.contribution_process_lock import (
    process_lock
)
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

CONTRIBUTIONS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "02_CONTRIBUTIONS.xlsx"
)


def get_next_receipt_number():

    receipts = load_receipts()

    current_year = datetime.now().year

    highest_sequence = 0

    for receipt in receipts:

        number = str(
            receipt.get("receipt_number", "")
        ).strip()

        parts = number.split("/")

        if len(parts) != 4:
            continue

        if parts[0] != "TLC":
            continue

        if parts[1] != "RCPT":
            continue

        if parts[2] != str(current_year):
            continue

        try:
            sequence = int(parts[3])
        except ValueError:
            continue

        highest_sequence = max(
            highest_sequence,
            sequence
        )

    return format_receipt_number(
        highest_sequence + 1,
        current_year
    )


def get_next_contribution_id(ws):

    highest = 0

    for row in range(2, ws.max_row + 1):

        value = ws.cell(
            row=row,
            column=1
        ).value

        if not value:
            continue

        text = str(value).strip().upper()

        if not text.startswith("TLC-CON-"):
            continue

        try:
            number = int(
                text.split("-")[-1]
            )
        except ValueError:
            continue

        highest = max(
            highest,
            number
        )

    return f"TLC-CON-{highest + 1:06d}"


def find_existing_receipt(
    receipt_number="",
    member_id="",
    transaction_reference="",
    amount=None,
    payment_type=""
):

    receipts = load_receipts()

    receipt_number = str(
        receipt_number or ""
    ).strip()

    member_id = str(
        member_id or ""
    ).strip().upper()

    transaction_reference = str(
        transaction_reference or ""
    ).strip()

    payment_type = str(
        payment_type or ""
    ).strip()

    try:
        amount_value = float(amount)
    except (TypeError, ValueError):
        amount_value = None

    # ------------------------------------------------------
    # FIRST: Exact receipt-number match
    # ------------------------------------------------------

    if receipt_number:

        for receipt in receipts:

            if str(
                receipt.get("receipt_number", "")
            ).strip() == receipt_number:

                return receipt

    # ------------------------------------------------------
    # SECOND: Match the original payment
    # ------------------------------------------------------

    for receipt in receipts:

        stored_member_id = str(
            receipt.get("member_id", "")
        ).strip().upper()

        stored_reference = str(
            receipt.get(
                "transaction_reference",
                receipt.get(
                    "payment_reference",
                    ""
                )
            ) or ""
        ).strip()

        stored_payment_type = str(
            receipt.get(
                "payment_type",
                ""
            ) or ""
        ).strip()

        if stored_member_id != member_id:
            continue

        if (
            transaction_reference
            and stored_reference
            != transaction_reference
        ):
            continue

        if (
            transaction_reference == ""
            and stored_reference != ""
        ):
            continue

        if (
            payment_type
            and stored_payment_type
            != payment_type
        ):
            continue

        if amount_value is not None:

            try:
                stored_amount = float(
                    receipt.get("amount")
                )
            except (TypeError, ValueError):
                continue

            if stored_amount != amount_value:
                continue

        return receipt

    return None


def process_pending_contributions():

    """
    Process spreadsheet rows marked PENDING.

    A system-wide lock prevents simultaneous processing.

    Any stale PROCESSING rows are recovered to PENDING
    because this function only proceeds after acquiring
    the system-wide lock.

    Existing saved receipts are reused to prevent duplicate
    receipt numbers after a failed PDF/email attempt.
    """

    lock = process_lock()

    if not lock.acquire():
        return []

    workbook = None

    try:

        if not CONTRIBUTIONS_FILE.exists():

            raise FileNotFoundError(
                f"Missing contributions workbook: "
                f"{CONTRIBUTIONS_FILE}"
            )

        workbook = load_workbook(
            CONTRIBUTIONS_FILE
        )

        if "Contributions" not in workbook.sheetnames:

            raise ValueError(
                "Contributions sheet was not found."
            )

        ws = workbook["Contributions"]

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in ws[1]
        ]

        required_headers = [
            "Contribution ID",
            "Receipt Number",
            "Member ID",
            "Full Name",
            "Phone Number",
            "Email",
            "Payment Date",
            "Payment Time",
            "Payment Type",
            "Description",
            "Amount",
            "Payment Method",
            "Transaction Reference",
            "Status",
            "Recorded By",
            "Created At",
        ]

        missing = [
            header
            for header in required_headers
            if header not in headers
        ]

        if missing:

            raise ValueError(
                "Missing spreadsheet columns: "
                + ", ".join(missing)
            )

        columns = {
            header: headers.index(header) + 1
            for header in required_headers
        }

        # ------------------------------------------------------
        # RECOVER STALE PROCESSING ROWS
        # ------------------------------------------------------

        recovered_rows = False

        for row_number in range(
            2,
            ws.max_row + 1
        ):

            status = str(
                ws.cell(
                    row=row_number,
                    column=columns["Status"]
                ).value or ""
            ).strip().upper()

            if status == "PROCESSING":

                ws.cell(
                    row=row_number,
                    column=columns["Status"]
                ).value = "PENDING"

                recovered_rows = True

                print(
                    "[TLC PROCESSOR] Recovered stale "
                    f"PROCESSING row {row_number}.",
                    flush=True
                )

        if recovered_rows:

            workbook.save(
                CONTRIBUTIONS_FILE
            )

        processed = []

        for row_number in range(
            2,
            ws.max_row + 1
        ):

            status = str(
                ws.cell(
                    row=row_number,
                    column=columns["Status"]
                ).value or ""
            ).strip().upper()

            if status != "PENDING":
                continue

            # --------------------------------------------------
            # CLAIM ROW
            # --------------------------------------------------

            ws.cell(
                row=row_number,
                column=columns["Status"]
            ).value = "PROCESSING"

            workbook.save(
                CONTRIBUTIONS_FILE
            )

            try:

                # ----------------------------------------------
                # READ MEMBER
                # ----------------------------------------------

                member_id = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Member ID"]
                    ).value or ""
                ).strip().upper()

                if not member_id:

                    ws.cell(
                        row=row_number,
                        column=columns["Status"]
                    ).value = (
                        "ERROR: Member ID required"
                    )

                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )

                    continue

                member = get_member_details(
                    member_id
                )

                if not member:

                    ws.cell(
                        row=row_number,
                        column=columns["Status"]
                    ).value = (
                        "ERROR: Member not found"
                    )

                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )

                    continue

                # ----------------------------------------------
                # PAYMENT DATA
                # ----------------------------------------------

                amount = ws.cell(
                    row=row_number,
                    column=columns["Amount"]
                ).value

                payment_type = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Payment Type"]
                    ).value or ""
                ).strip()

                description = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Description"]
                    ).value or ""
                ).strip()

                payment_method = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Payment Method"]
                    ).value or "M-PESA"
                ).strip()

                transaction_reference = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Transaction Reference"]
                    ).value or ""
                ).strip()

                recorded_by = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Recorded By"]
                    ).value or ""
                ).strip()

                payment_date = ws.cell(
                    row=row_number,
                    column=columns["Payment Date"]
                ).value

                payment_time = ws.cell(
                    row=row_number,
                    column=columns["Payment Time"]
                ).value

                # ----------------------------------------------
                # VALIDATION
                # ----------------------------------------------

                if not payment_type:

                    ws.cell(
                        row=row_number,
                        column=columns["Status"]
                    ).value = (
                        "ERROR: Payment type required"
                    )

                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )

                    continue

                if amount is None:

                    ws.cell(
                        row=row_number,
                        column=columns["Status"]
                    ).value = (
                        "ERROR: Amount required"
                    )

                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )

                    continue

                try:

                    amount = float(amount)

                    if amount <= 0:
                        raise ValueError

                except (TypeError, ValueError):

                    ws.cell(
                        row=row_number,
                        column=columns["Status"]
                    ).value = (
                        "ERROR: Invalid amount"
                    )

                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )

                    continue

                # ----------------------------------------------
                # DEFAULT DATE/TIME
                # ----------------------------------------------

                now = datetime.now()

                if not payment_date:

                    payment_date = now.strftime(
                        "%Y-%m-%d"
                    )

                if not payment_time:

                    payment_time = now.strftime(
                        "%H:%M:%S"
                    )

                # ----------------------------------------------
                # REUSE EXISTING RECEIPT WHEN AVAILABLE
                # ----------------------------------------------

                existing_receipt_number = str(
                    ws.cell(
                        row=row_number,
                        column=columns["Receipt Number"]
                    ).value or ""
                ).strip()

                saved_receipt = find_existing_receipt(
                    receipt_number=existing_receipt_number,
                    member_id=member["member_id"],
                    transaction_reference=(
                        transaction_reference
                    ),
                    amount=amount,
                    payment_type=payment_type,
                )

                # ----------------------------------------------
                # CREATE OR REUSE RECEIPT
                # ----------------------------------------------

                if saved_receipt:

                    receipt_number = str(
                        saved_receipt.get(
                            "receipt_number",
                            existing_receipt_number
                        )
                    ).strip()

                    contribution_id = str(
                        ws.cell(
                            row=row_number,
                            column=columns["Contribution ID"]
                        ).value or ""
                    ).strip()

                    if not contribution_id:

                        contribution_id = (
                            get_next_contribution_id(ws)
                        )

                    print(
                        "[TLC PROCESSOR] Reusing receipt:",
                        receipt_number,
                        flush=True
                    )

                else:

                    receipt_number = (
                        get_next_receipt_number()
                    )

                    contribution_id = (
                        str(
                            ws.cell(
                                row=row_number,
                                column=columns[
                                    "Contribution ID"
                                ]
                            ).value or ""
                        ).strip()
                    )

                    if not contribution_id:

                        contribution_id = (
                            get_next_contribution_id(ws)
                        )

                    receipt = create_receipt(
                        member_id=member["member_id"],
                        member_name=member["full_name"],
                        amount=amount,
                        payment_type=payment_type,
                        payment_reference=(
                            transaction_reference
                        ),
                        description=description,
                        receipt_number=receipt_number,
                        payment_date=str(
                            payment_date
                        ),
                        phone_number=(
                            member["phone_number"]
                        ),
                        email=member["email"],
                        payment_method=payment_method,
                        recorded_by=recorded_by,
                    )

                    receipt["payment_time"] = str(
                        payment_time
                    )

                    receipt[
                        "transaction_reference"
                    ] = transaction_reference

                    saved_receipt = add_receipt(
                        receipt
                    )

                    print(
                        "[TLC PROCESSOR] Created receipt:",
                        receipt_number,
                        flush=True
                    )

                # ----------------------------------------------
                # SAVE IDENTIFIERS IMMEDIATELY
                # ----------------------------------------------

                ws.cell(
                    row=row_number,
                    column=columns["Contribution ID"]
                ).value = contribution_id

                ws.cell(
                    row=row_number,
                    column=columns["Receipt Number"]
                ).value = receipt_number

                ws.cell(
                    row=row_number,
                    column=columns["Full Name"]
                ).value = member["full_name"]

                ws.cell(
                    row=row_number,
                    column=columns["Phone Number"]
                ).value = member["phone_number"]

                ws.cell(
                    row=row_number,
                    column=columns["Email"]
                ).value = member["email"]

                ws.cell(
                    row=row_number,
                    column=columns["Payment Date"]
                ).value = payment_date

                ws.cell(
                    row=row_number,
                    column=columns["Payment Time"]
                ).value = payment_time

                ws.cell(
                    row=row_number,
                    column=columns["Created At"]
                ).value = (
                    now.strftime(
                        "%Y-%m-%d %H:%M:%S"
                    )
                )

                workbook.save(
                    CONTRIBUTIONS_FILE
                )

                # ----------------------------------------------
                # GENERATE EXACT PORTAL PDF
                # ----------------------------------------------

                print(
                    "[TLC PROCESSOR] Generating PDF:",
                    receipt_number,
                    flush=True
                )

                pdf_bytes = (
                    generate_receipt_pdf(
                        saved_receipt
                    )
                )

                # ----------------------------------------------
                # EMAIL PDF
                # ----------------------------------------------

                member_email = str(
                    member.get("email", "")
                ).strip()

                if not member_email:

                    raise ValueError(
                        "Member has no email address."
                    )

                print(
                    "[TLC PROCESSOR] Sending email:",
                    receipt_number,
                    "->",
                    member_email,
                    flush=True
                )

                email_sent, email_message = (
                    send_receipt_email(
                        recipient_email=member_email,
                        receipt_number=receipt_number,
                        member_name=member["full_name"],
                        amount=amount,
                        pdf_bytes=pdf_bytes,
                    )
                )

                if not email_sent:

                    raise RuntimeError(
                        email_message
                    )

                # ----------------------------------------------
                # COMPLETE PAYMENT
                # ----------------------------------------------

                ws.cell(
                    row=row_number,
                    column=columns["Status"]
                ).value = (
                    "RECORDED - EMAILED"
                )

                workbook.save(
                    CONTRIBUTIONS_FILE
                )

                print(
                    "[TLC PROCESSOR] SUCCESS:",
                    receipt_number,
                    flush=True
                )

                processed.append({
                    "row": row_number,
                    "contribution_id": (
                        contribution_id
                    ),
                    "receipt_number": (
                        receipt_number
                    ),
                    "email_sent": True,
                    "email_message": (
                        "Receipt PDF emailed successfully."
                    ),
                })

            except Exception as error:

                # ----------------------------------------------
                # SAFE AUTOMATIC RETRY
                # ----------------------------------------------

                ws.cell(
                    row=row_number,
                    column=columns["Status"]
                ).value = "PENDING"

                try:
                    workbook.save(
                        CONTRIBUTIONS_FILE
                    )
                except Exception as save_error:

                    print(
                        "[TLC PROCESSOR] FAILED TO SAVE "
                        "RETRY STATUS:",
                        repr(save_error),
                        flush=True
                    )

                print(
                    "[TLC PROCESSOR] ERROR on row "
                    f"{row_number}:",
                    repr(error),
                    flush=True
                )

        return processed

    finally:

        if workbook is not None:

            try:
                workbook.close()
            except Exception:
                pass

        lock.release()
