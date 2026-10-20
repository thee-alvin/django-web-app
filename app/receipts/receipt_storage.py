from pathlib import Path
import json

from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

RECEIPTS_DIR = DATA_DIR / "receipts"
RECEIPTS_FILE = RECEIPTS_DIR / "receipts.json"

MEMBERS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

CONTRIBUTIONS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "02_CONTRIBUTIONS.xlsx"
)


def ensure_receipt_storage():
    """Create the receipt storage location when it does not exist."""
    RECEIPTS_DIR.mkdir(parents=True, exist_ok=True)

    if not RECEIPTS_FILE.exists():
        RECEIPTS_FILE.write_text(
            "[]",
            encoding="utf-8"
        )


def load_receipts():
    """Return all stored TLC receipt records."""
    ensure_receipt_storage()

    try:
        data = json.loads(
            RECEIPTS_FILE.read_text(
                encoding="utf-8"
            )
        )

        if not isinstance(data, list):
            return []

        return data

    except (json.JSONDecodeError, OSError):
        return []


def save_receipts(receipts):
    """Save the complete TLC receipt collection."""
    ensure_receipt_storage()

    RECEIPTS_FILE.write_text(
        json.dumps(
            receipts,
            indent=4,
            ensure_ascii=False
        ),
        encoding="utf-8"
    )


def add_receipt(receipt):
    """Add one receipt record to storage."""
    receipts = load_receipts()

    receipts.append(receipt)

    save_receipts(receipts)

    return receipt


def get_member_details(member_id):
    """
    Look up member information from 01_MEMBERS.xlsx.
    """

    member_id = str(member_id).strip().upper()

    if not member_id:
        return None

    if not MEMBERS_FILE.exists():
        return None

    try:
        workbook = load_workbook(
            MEMBERS_FILE,
            read_only=True,
            data_only=True
        )

        if "Members" not in workbook.sheetnames:
            return None

        sheet = workbook["Members"]

        rows = sheet.iter_rows(values_only=True)

        headers = next(rows, None)

        if not headers:
            return None

        headers = [
            str(header).strip()
            if header is not None
            else ""
            for header in headers
        ]

        for row in rows:

            member = dict(zip(headers, row))

            current_member_id = str(
                member.get("Member ID") or ""
            ).strip().upper()

            if current_member_id != member_id:
                continue

            return {
                "member_id": current_member_id,
                "full_name": str(
                    member.get("Full Name") or ""
                ).strip(),
                "phone_number": str(
                    member.get("Phone Number") or ""
                ).strip(),
                "email": str(
                    member.get("Email Address") or ""
                ).strip(),
                "membership_status": str(
                    member.get("Membership Status") or ""
                ).strip(),
                "registration_date": str(
                    member.get("Registration Date") or ""
                ).strip(),
                "account_status": str(
                    member.get("Account Status") or ""
                ).strip()
            }

        return None

    except Exception:
        return None


def enrich_receipt_with_member(receipt):
    """
    Add the current member information to a receipt.
    """

    if not isinstance(receipt, dict):
        return receipt

    member_id = receipt.get("member_id")

    member = get_member_details(member_id)

    if not member:
        return receipt

    enriched = dict(receipt)

    enriched["member_id"] = member["member_id"]
    enriched["full_name"] = member["full_name"]
    enriched["phone_number"] = member["phone_number"]
    enriched["email"] = member["email"]

    return enriched


def get_receipts_for_member(member_id):
    """
    Return receipts belonging only to the specified member.

    Member identity information is refreshed from
    01_MEMBERS.xlsx.
    """

    member_id = str(member_id).strip().upper()

    if not member_id:
        return []

    receipts = []

    for receipt in load_receipts():

        if not isinstance(receipt, dict):
            continue

        receipt_member_id = str(
            receipt.get("member_id", "")
        ).strip().upper()

        if receipt_member_id != member_id:
            continue

        receipts.append(
            enrich_receipt_with_member(receipt)
        )

    return receipts
