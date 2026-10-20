from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

WELFARE_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "05_WELFARE.xlsx"
)


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _read_sheet(sheet_name, member_id):
    if not member_id:
        return []

    if not WELFARE_FILE.exists():
        raise FileNotFoundError(
            f"Welfare workbook not found: {WELFARE_FILE}"
        )

    workbook = load_workbook(
        WELFARE_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if sheet_name not in workbook.sheetnames:
            return []

        sheet = workbook[sheet_name]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        records = []

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            record = dict(
                zip(headers, row)
            )

            if (
                _clean(
                    record.get("Member ID")
                )
                != _clean(member_id)
            ):
                continue

            records.append({
                key: _clean(value)
                if key not in {
                    "Amount Requested",
                    "Amount",
                    "Amount / Value",
                }
                else _amount(value)
                for key, value
                in record.items()
            })

        return records

    finally:
        workbook.close()


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def get_member_welfare(member_id):

    return {
        "requests": _read_sheet(
            "Welfare Requests",
            member_id
        ),

        "contributions": _read_sheet(
            "Welfare Contributions",
            member_id
        ),

        "support": _read_sheet(
            "Welfare Support",
            member_id
        ),

        "cases": _read_sheet(
            "Welfare Cases",
            member_id
        ),
    }


def get_member_welfare_summary(member_id):

    welfare = get_member_welfare(
        member_id
    )

    requests = welfare["requests"]
    contributions = welfare["contributions"]
    support = welfare["support"]
    cases = welfare["cases"]

    return {
        "request_count": len(requests),
        "contribution_count": len(contributions),
        "support_count": len(support),
        "case_count": len(cases),
        "latest_request": (
            requests[-1]
            if requests
            else None
        ),
        "latest_support": (
            support[-1]
            if support
            else None
        ),
    }
