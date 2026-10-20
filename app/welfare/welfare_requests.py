from datetime import datetime
from pathlib import Path

from flask import Blueprint, jsonify, request, session
from openpyxl import load_workbook

from app.security.audit_log import record_activity
from app.config.data_paths import DATA_DIR


welfare_requests_bp = Blueprint(
    "welfare_requests",
    __name__,
)

BASE_DIR = Path(__file__).resolve().parents[2]

MEMBERS_FILE = DATA_DIR / "spreadsheets" / "01_MEMBERS.xlsx"
WELFARE_FILE = DATA_DIR / "spreadsheets" / "05_WELFARE.xlsx"


REQUEST_HEADERS = [
    "Request ID",
    "Member ID",
    "Full Name",
    "Request Date",
    "Welfare Category",
    "Description",
    "Amount Requested",
    "Status",
    "Reviewed By",
    "Decision Date",
    "Officer Remarks",
    "Created At",
]


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def _get_member(member_id):
    workbook = load_workbook(
        MEMBERS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)

        headers = next(rows, None)

        if not headers:
            return None

        header_map = {
            _clean(header): index
            for index, header in enumerate(headers)
        }

        member_id_index = header_map.get("Member ID")
        full_name_index = header_map.get("Full Name")

        if member_id_index is None or full_name_index is None:
            return None

        for row in rows:
            if (
                len(row) > member_id_index
                and _clean(row[member_id_index]) == member_id
            ):
                return {
                    "member_id": member_id,
                    "full_name": _clean(row[full_name_index]),
                }

        return None

    finally:
        workbook.close()


def _next_request_id(sheet):
    highest = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):
        request_id = _clean(row[0])

        if not request_id:
            continue

        if request_id.startswith("TLC-WEL-"):
            try:
                number = int(request_id.split("-")[-1])

                if number > highest:
                    highest = number

            except (ValueError, IndexError):
                continue

    return f"TLC-WEL-{highest + 1:06d}"


def create_welfare_request(
    member_id,
    welfare_category,
    description,
    amount_requested,
):
    member = _get_member(member_id)

    if not member:
        raise ValueError("Member account could not be found.")

    workbook = load_workbook(WELFARE_FILE)

    try:
        if "Welfare Requests" not in workbook.sheetnames:
            raise ValueError(
                "Welfare Requests sheet is missing from the welfare workbook."
            )

        sheet = workbook["Welfare Requests"]

        existing_headers = [
            _clean(cell.value)
            for cell in sheet[1]
        ]

        if existing_headers != REQUEST_HEADERS:
            raise ValueError(
                "Welfare Requests sheet structure does not match the expected format."
            )

        request_id = _next_request_id(sheet)

        now = datetime.now()

        sheet.append(
            [
                request_id,
                member["member_id"],
                member["full_name"],
                now.strftime("%Y-%m-%d"),
                welfare_category,
                description,
                amount_requested,
                "PENDING",
                "",
                "",
                "",
                now.strftime("%Y-%m-%d %H:%M:%S"),
            ]
        )

        workbook.save(WELFARE_FILE)

        record_activity(
            member["member_id"],
            action="Welfare Request Submitted",
            description=(
                f"Welfare request {request_id} was submitted "
                f"for KES {amount_requested:,.2f}."
            ),
            category="welfare",
            details={
                "request_id": request_id,
                "welfare_category": welfare_category,
                "amount_requested": amount_requested,
            }
        )

        return {
            "request_id": request_id,
            "member_id": member["member_id"],
            "full_name": member["full_name"],
            "request_date": now.strftime("%Y-%m-%d"),
            "welfare_category": welfare_category,
            "description": description,
            "amount_requested": amount_requested,
            "status": "PENDING",
        }

    finally:
        workbook.close()


@welfare_requests_bp.route(
    "/api/member/welfare/request",
    methods=["POST"],
)
def submit_welfare_request():
    member_id = session.get("member_id")

    if not member_id:
        return jsonify(
            {
                "success": False,
                "message": "Authentication required.",
            }
        ), 401

    data = request.get_json(silent=True) or {}

    welfare_category = _clean(
        data.get("welfare_category")
    )

    description = _clean(
        data.get("description")
    )

    amount_raw = data.get("amount_requested")

    if not welfare_category:
        return jsonify(
            {
                "success": False,
                "message": "Welfare category is required.",
            }
        ), 400

    if not description:
        return jsonify(
            {
                "success": False,
                "message": "Description is required.",
            }
        ), 400

    try:
        amount_requested = float(amount_raw)

    except (TypeError, ValueError):
        return jsonify(
            {
                "success": False,
                "message": "Amount requested must be a valid number.",
            }
        ), 400

    if amount_requested <= 0:
        return jsonify(
            {
                "success": False,
                "message": "Amount requested must be greater than zero.",
            }
        ), 400

    if amount_requested > 100000000:
        return jsonify(
            {
                "success": False,
                "message": "Amount requested is too large.",
            }
        ), 400

    description = description[:2000]
    welfare_category = welfare_category[:100]

    try:
        result = create_welfare_request(
            member_id=member_id,
            welfare_category=welfare_category,
            description=description,
            amount_requested=amount_requested,
        )

        return jsonify(
            {
                "success": True,
                "message": "Welfare request submitted successfully.",
                "request": result,
            }
        ), 201

    except ValueError as exc:
        return jsonify(
            {
                "success": False,
                "message": str(exc),
            }
        ), 400

    except Exception:
        return jsonify(
            {
                "success": False,
                "message": "Unable to save the welfare request.",
            }
        ), 500


