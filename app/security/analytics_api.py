from pathlib import Path
from datetime import datetime, timezone, timedelta

from flask import Blueprint, jsonify, request, send_from_directory
from openpyxl import load_workbook

from app.auth.session import get_current_member_id
from app.security.audit_log import _read_records
from app.config.data_paths import DATA_DIR


analytics_api = Blueprint(
    "analytics_api",
    __name__
)

BASE_DIR = Path(__file__).resolve().parents[2]

FRONTEND_DIR = BASE_DIR / "frontend"

MEMBERS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

ANALYTICS_RESET_FILE = (
    DATA_DIR
    / "activity"
    / "analytics_reset.json"
)

ADMIN_MEMBER_ID = "TLC/001/2026"


def _is_admin():
    return (
        str(
            get_current_member_id() or ""
        ).strip().upper()
        == ADMIN_MEMBER_ID
    )


def _clean(value):
    return str(value or "").strip()


def _parse_timestamp(value):
    text = _clean(value)

    if not text:
        return None

    try:
        return datetime.fromisoformat(
            text.replace("Z", "+00:00")
        )
    except ValueError:
        return None


def _get_reset_time():
    if not ANALYTICS_RESET_FILE.exists():
        return None

    try:
        text = ANALYTICS_RESET_FILE.read_text(
            encoding="utf-8"
        ).strip()

        if not text:
            return None

        return datetime.fromisoformat(text)

    except Exception:
        return None


def _set_reset_time():
    ANALYTICS_RESET_FILE.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    timestamp = datetime.now(timezone.utc)

    ANALYTICS_RESET_FILE.write_text(
        timestamp.isoformat(),
        encoding="utf-8"
    )

    return timestamp


def _load_members():
    if not MEMBERS_FILE.exists():
        return []

    workbook = load_workbook(
        MEMBERS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if "Members" not in workbook.sheetnames:
            return []

        sheet = workbook["Members"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [_clean(value) for value in headers]

        if (
            "Member ID" not in headers
            or "Full Name" not in headers
        ):
            return []

        member_id_index = headers.index("Member ID")
        name_index = headers.index("Full Name")

        members = []

        for row in rows:
            member_id = _clean(
                row[member_id_index]
                if member_id_index < len(row)
                else ""
            ).upper()

            name = _clean(
                row[name_index]
                if name_index < len(row)
                else ""
            )

            if not member_id:
                continue

            members.append({
                "member_id": member_id,
                "name": name or member_id
            })

        return members

    finally:
        workbook.close()


def _active_records():
    records = _read_records()

    reset_time = _get_reset_time()

    if reset_time is None:
        return records

    active = []

    for record in records:
        timestamp = _parse_timestamp(
            record.get("timestamp")
        )

        if timestamp is None:
            continue

        if timestamp >= reset_time:
            active.append(record)

    return active


def _matches(
    record,
    member_id,
    action,
    category,
    date_from,
    date_to
):
    record_member = _clean(
        record.get("member_id")
    ).upper()

    if member_id and record_member != member_id:
        return False

    record_action = _clean(
        record.get("action")
    )

    if (
        action
        and record_action.lower()
        != action.lower()
    ):
        return False

    record_category = _clean(
        record.get("category")
    )

    if (
        category
        and record_category.lower()
        != category.lower()
    ):
        return False

    timestamp = _parse_timestamp(
        record.get("timestamp")
    )

    if timestamp is None:
        return False

    if date_from and timestamp.date().isoformat() < date_from:
        return False

    if date_to and timestamp.date().isoformat() > date_to:
        return False

    return True


def _build_daily_series(records):
    end_date = datetime.now(
        timezone.utc
    ).date()

    start_date = (
        end_date - timedelta(days=29)
    )

    activity_counts = {}
    login_counts = {}

    for record in records:
        timestamp = _parse_timestamp(
            record.get("timestamp")
        )

        if timestamp is None:
            continue

        day = timestamp.date().isoformat()

        activity_counts[day] = (
            activity_counts.get(day, 0) + 1
        )

        action = _clean(
            record.get("action")
        ).lower()

        if action == "member login":
            login_counts[day] = (
                login_counts.get(day, 0) + 1
            )

    series = []

    current = start_date

    while current <= end_date:
        day = current.isoformat()

        series.append({
            "date": day,
            "activities": activity_counts.get(day, 0),
            "logins": login_counts.get(day, 0)
        })

        current += timedelta(days=1)

    return series


@analytics_api.get("/api/tlc/analytics")
def analytics():
    if not _is_admin():
        return jsonify({
            "success": False,
            "message": "Access denied."
        }), 403

    records = _active_records()

    member_id = _clean(
        request.args.get("member_id")
    ).upper()

    action = _clean(
        request.args.get("action")
    )

    category = _clean(
        request.args.get("category")
    )

    date_from = _clean(
        request.args.get("date_from")
    )

    date_to = _clean(
        request.args.get("date_to")
    )

    filtered_records = [
        record
        for record in records
        if _matches(
            record,
            member_id,
            action,
            category,
            date_from,
            date_to
        )
    ]

    filtered_records.sort(
        key=lambda item: str(
            item.get("timestamp", "")
        ),
        reverse=True
    )

    categories = {}
    actions = {}

    for record in filtered_records:
        category_name = (
            _clean(record.get("category"))
            or "uncategorized"
        )

        action_name = (
            _clean(record.get("action"))
            or "activity"
        )

        categories[category_name] = (
            categories.get(category_name, 0) + 1
        )

        actions[action_name] = (
            actions.get(action_name, 0) + 1
        )

    login_records = [
        record
        for record in filtered_records
        if _clean(
            record.get("action")
        ).lower() == "member login"
    ]

    today = datetime.now(
        timezone.utc
    ).date().isoformat()

    today_logins = [
        record
        for record in login_records
        if _clean(
            record.get("timestamp")
        ).startswith(today)
    ]

    unique_members = len({
        _clean(
            record.get("member_id")
        ).upper()
        for record in filtered_records
        if _clean(record.get("member_id"))
    })

    selected_member = None

    if member_id:
        for member in _load_members():
            if member["member_id"] == member_id:
                selected_member = member
                break

    member_last_login = None

    if member_id:
        member_logins = [
            record
            for record in records
            if (
                _clean(
                    record.get("member_id")
                ).upper() == member_id
                and
                _clean(
                    record.get("action")
                ).lower() == "member login"
            )
        ]

        member_logins.sort(
            key=lambda item: str(
                item.get("timestamp", "")
            ),
            reverse=True
        )

        if member_logins:
            member_last_login = member_logins[0]

    # ------------------------------------------------------
    # EXECUTIVE ANALYTICS
    # ------------------------------------------------------

    member_names = {
        member["member_id"]: member["name"]
        for member in _load_members()
    }

    member_activity_counts = {}

    for record in filtered_records:

        current_id = _clean(
            record.get("member_id")
        ).upper()

        if not current_id:
            continue

        member_activity_counts[current_id] = (
            member_activity_counts.get(
                current_id,
                0
            ) + 1
        )

    most_active_member = None

    if member_activity_counts:

        top_member_id = max(
            member_activity_counts,
            key=member_activity_counts.get
        )

        most_active_member = {
            "member_id": top_member_id,
            "name": member_names.get(
                top_member_id,
                top_member_id
            ),
            "activities": member_activity_counts[
                top_member_id
            ]
        }

    daily_series = _build_daily_series(
        filtered_records
    )

    peak_activity_day = None

    if daily_series:

        peak = max(
            daily_series,
            key=lambda item: item["activities"]
        )

        if peak["activities"] > 0:

            peak_activity_day = peak

    latest_activity = (
        filtered_records[0]
        if filtered_records
        else None
    )

    return jsonify({
        "success": True,

        "summary": {
            "total_members": len(_load_members()),
            "total_activities": len(filtered_records),
            "total_logins": len(login_records),
            "today_logins": len(today_logins),
            "unique_members": unique_members,

            "member_last_login": (
                member_last_login.get("timestamp")
                if member_last_login
                else None
            ),

            "most_active_member":
                most_active_member,

            "peak_activity_day":
                peak_activity_day,

            "latest_activity":
                latest_activity,

            "login_events":
                len(login_records)
        },
        "member": selected_member,
        "members": _load_members(),
        "categories": categories,
        "actions": actions,
        "daily_series": _build_daily_series(
            filtered_records
        ),
        "activities": filtered_records[:500],
        "filters": {
            "member_id": member_id,
            "action": action,
            "category": category,
            "date_from": date_from,
            "date_to": date_to
        }
    })


@analytics_api.post("/api/tlc/analytics/reset")
def reset_analytics():
    if not _is_admin():
        return jsonify({
            "success": False,
            "message": "Access denied."
        }), 403

    timestamp = _set_reset_time()

    return jsonify({
        "success": True,
        "message": "Analytics session reset successfully.",
        "reset_time": timestamp.isoformat()
    })


@analytics_api.get("/tlc-analytics")
def analytics_page():
    if not _is_admin():
        return jsonify({
            "success": False,
            "message": "Access denied."
        }), 403

    return send_from_directory(
        FRONTEND_DIR / "pages",
        "tlc_analytics.html"
    )
