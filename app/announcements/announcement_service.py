from pathlib import Path
from datetime import date, datetime

from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

ANNOUNCEMENTS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "03_ANNOUNCEMENTS.xlsx"
)


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _expiry_date(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%d/%m/%Y",
    ):
        try:
            return datetime.strptime(
                text,
                fmt
            ).date()
        except ValueError:
            continue

    return None


def get_announcements():

    if not ANNOUNCEMENTS_FILE.exists():
        raise FileNotFoundError(
            f"Announcements workbook not found: {ANNOUNCEMENTS_FILE}"
        )

    workbook = load_workbook(
        ANNOUNCEMENTS_FILE,
        read_only=True,
        data_only=True
    )

    try:

        if "Announcements" not in workbook.sheetnames:
            raise ValueError(
                "Announcements sheet was not found."
            )

        sheet = workbook["Announcements"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(header)
            for header in headers
        ]

        records = []

        today = date.today()

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            record = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                if hasattr(value, "strftime"):

                    if header in (
                        "Announcement Date",
                        "Expiry Date"
                    ):
                        value = value.strftime(
                            "%Y-%m-%d"
                        )

                    elif header == "Created At":
                        value = value.strftime(
                            "%Y-%m-%d %H:%M:%S"
                        )

                record[header] = (
                    value
                    if value is not None
                    else ""
                )

            # Only published announcements are public.
            status = _clean(
                record.get("Status")
            ).upper()

            if status != "PUBLISHED":
                continue

            # Automatically hide expired announcements.
            expiry = _expiry_date(
                record.get("Expiry Date")
            )

            if expiry is not None and expiry < today:
                continue

            records.append(record)

        records.sort(
            key=lambda item: str(
                item.get("Announcement Date")
                or ""
            ),
            reverse=True
        )

        return records

    finally:
        workbook.close()
