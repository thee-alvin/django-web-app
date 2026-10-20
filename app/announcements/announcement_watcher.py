from pathlib import Path
import json
import hashlib
import threading

from openpyxl import load_workbook

from app.announcements.announcement_email_service import (
    send_announcement_email
)
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

ANNOUNCEMENTS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "03_ANNOUNCEMENTS.xlsx"
)

MEMBERS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

STATE_DIR = (
    DATA_DIR
    / "announcements"
)

STATE_FILE = (
    STATE_DIR
    / "announcement_notification_state.json"
)

CHECK_INTERVAL = 2


class AnnouncementWatcher:

    def __init__(self):
        self._stop_event = threading.Event()
        self._thread = None

    def start(self):

        if (
            self._thread
            and self._thread.is_alive()
        ):
            return

        self._stop_event.clear()

        self._thread = threading.Thread(
            target=self._run,
            name="TLC-Announcement-Watcher",
            daemon=True,
        )

        self._thread.start()

        print(
            "[TLC ANNOUNCEMENT WATCHER] Started.",
            flush=True,
        )

    def stop(self):

        self._stop_event.set()

        if self._thread:
            self._thread.join(timeout=3)

    def _clean(self, value):

        if value is None:
            return ""

        if hasattr(value, "strftime"):

            try:
                return value.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            except Exception:
                pass

        return str(value).strip()

    def _load_state(self):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        if not STATE_FILE.exists():
            STATE_FILE.write_text(
                "{}",
                encoding="utf-8"
            )

        try:
            data = json.loads(
                STATE_FILE.read_text(
                    encoding="utf-8"
                )
            )

            return (
                data
                if isinstance(data, dict)
                else {}
            )

        except Exception:
            return {}

    def _save_state(self, state):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        STATE_FILE.write_text(
            json.dumps(
                state,
                indent=4,
                ensure_ascii=False
            ),
            encoding="utf-8"
        )

    def _fingerprint(self, record):

        fields = [
            "Announcement ID",
            "Announcement Date",
            "Title",
            "Message",
            "Announcement Type",
            "Priority",
            "Status",
            "Posted By",
            "Expiry Date",
            "Created At",
        ]

        payload = {
            field: self._clean(
                record.get(field)
            )
            for field in fields
        }

        raw = json.dumps(
            payload,
            sort_keys=True,
            ensure_ascii=False
        )

        return hashlib.sha256(
            raw.encode("utf-8")
        ).hexdigest()

    def _load_announcements(self):

        if not ANNOUNCEMENTS_FILE.exists():
            return []

        workbook = load_workbook(
            ANNOUNCEMENTS_FILE,
            read_only=True,
            data_only=True
        )

        try:

            if "Announcements" not in workbook.sheetnames:
                return []

            sheet = workbook["Announcements"]

            rows = sheet.iter_rows(
                values_only=True
            )

            headers = next(rows, None)

            if not headers:
                return []

            headers = [
                self._clean(header)
                for header in headers
            ]

            records = []

            for row in rows:

                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                record = {}

                for index, header in enumerate(headers):

                    if not header:
                        continue

                    value = (
                        row[index]
                        if index < len(row)
                        else None
                    )

                    record[header] = self._clean(
                        value
                    )

                if (
                    record.get("Status", "")
                    .upper()
                    != "PUBLISHED"
                ):
                    continue

                announcement_id = (
                    record.get(
                        "Announcement ID",
                        ""
                    ).strip()
                )

                if not announcement_id:
                    continue

                records.append(record)

            return records

        finally:
            workbook.close()

    def _load_member_contacts(self):

        if not MEMBERS_FILE.exists():
            return []

        workbook = load_workbook(
            MEMBERS_FILE,
            read_only=True,
            data_only=True
        )

        try:

            if "Members" not in workbook.sheetnames:
                return []

            sheet = workbook["Members"]

            rows = sheet.iter_rows(
                values_only=True
            )

            headers = next(rows, None)

            if not headers:
                return []

            headers = [
                self._clean(header)
                for header in headers
            ]

            try:
                name_index = headers.index(
                    "Full Name"
                )

                email_index = headers.index(
                    "Email Address"
                )

            except ValueError:
                return []

            contacts = []
            seen = set()

            for row in rows:

                if (
                    name_index >= len(row)
                    or email_index >= len(row)
                ):
                    continue

                name = self._clean(
                    row[name_index]
                )

                email = self._clean(
                    row[email_index]
                )

                if not email:
                    continue

                email_key = email.lower()

                if email_key in seen:
                    continue

                seen.add(email_key)

                contacts.append({
                    "name": name,
                    "email": email,
                })

            return contacts

        finally:
            workbook.close()

    def _notify_members(
        self,
        announcement,
        is_update
    ):

        contacts = self._load_member_contacts()

        if not contacts:

            print(
                "[TLC ANNOUNCEMENT WATCHER] "
                "No member contacts found.",
                flush=True,
            )

            return

        title = announcement.get(
            "Title",
            ""
        )

        sent = 0
        failed = 0

        print(
            "[TLC ANNOUNCEMENT WATCHER] "
            f"Sending {'UPDATE' if is_update else 'NEW'}: "
            f"{title}",
            flush=True,
        )

        for contact in contacts:

            success, response = (
                send_announcement_email(
                    recipient_email=contact["email"],
                    recipient_name=contact["name"],
                    title=title,
                    message=announcement.get(
                        "Message",
                        ""
                    ),
                    announcement_type=announcement.get(
                        "Announcement Type",
                        "GENERAL"
                    ),
                    priority=announcement.get(
                        "Priority",
                        "NORMAL"
                    ),
                    announcement_date=announcement.get(
                        "Announcement Date",
                        ""
                    ),
                    expiry_date=announcement.get(
                        "Expiry Date",
                        ""
                    ),
                    is_update=is_update,
                )
            )

            if success:
                sent += 1

                print(
                    "[TLC ANNOUNCEMENT WATCHER] "
                    f"SENT -> {contact['name']} "
                    f"<{contact['email']}>",
                    flush=True,
                )

            else:

                failed += 1

                print(
                    "[TLC ANNOUNCEMENT WATCHER] "
                    f"FAILED -> {contact['name']} "
                    f"<{contact['email']}>: {response}",
                    flush=True,
                )

        print(
            "[TLC ANNOUNCEMENT WATCHER] "
            f"Email result: {sent} sent, {failed} failed.",
            flush=True,
        )

    def _run(self):

        first_scan = True

        while not self._stop_event.is_set():

            try:

                announcements = (
                    self._load_announcements()
                )

                state = self._load_state()

                current = {}

                # If there is no notification history yet, bootstrap it
                # from the announcements already present in the workbook.
                # Existing announcements are NOT emailed during bootstrap.
                if not state:

                    for announcement in announcements:

                        announcement_id = (
                            announcement.get(
                                "Announcement ID",
                                ""
                            ).strip()
                        )

                        if not announcement_id:
                            continue

                        current[announcement_id] = (
                            self._fingerprint(
                                announcement
                            )
                        )

                    self._save_state(current)

                    print(
                        "[TLC ANNOUNCEMENT WATCHER] "
                        "Existing announcements registered. "
                        "No old announcements emailed.",
                        flush=True,
                    )

                else:

                    for announcement in announcements:

                        announcement_id = (
                            announcement.get(
                                "Announcement ID",
                                ""
                            ).strip()
                        )

                        if not announcement_id:
                            continue

                        fingerprint = (
                            self._fingerprint(
                                announcement
                            )
                        )

                        current[announcement_id] = fingerprint

                        if announcement_id not in state:

                            print(
                                "[TLC ANNOUNCEMENT WATCHER] "
                                f"NEW ANNOUNCEMENT: "
                                f"{announcement_id}",
                                flush=True,
                            )

                            self._notify_members(
                                announcement,
                                is_update=False,
                            )

                        elif (
                            state[announcement_id]
                            != fingerprint
                        ):

                            print(
                                "[TLC ANNOUNCEMENT WATCHER] "
                                f"UPDATED ANNOUNCEMENT: "
                                f"{announcement_id}",
                                flush=True,
                            )

                            self._notify_members(
                                announcement,
                                is_update=True,
                            )

                    self._save_state(current)

                first_scan = False

            except Exception as error:

                print(
                    "[TLC ANNOUNCEMENT WATCHER] ERROR:",
                    repr(error),
                    flush=True,
                )

            self._stop_event.wait(
                CHECK_INTERVAL
            )


announcement_watcher = AnnouncementWatcher()

