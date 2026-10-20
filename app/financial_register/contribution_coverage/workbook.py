from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side


BASE_DIR = Path(__file__).resolve().parents[3]

COVERAGE_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "03_CONTRIBUTION_COVERAGE.xlsx"
)


DEFAULT_SETTINGS = {
    "Weekly Contribution Amount": "100",
    "Contribution Weekday": "THURSDAY",
    "Coverage Allocation Rule": (
        "Current contribution week, then consecutive future weeks"
    ),
}


COVERAGE_HEADERS = [
    "Coverage ID",
    "Contribution ID",
    "Receipt Number",
    "Member ID",
    "Full Name",
    "Payment Date",
    "Coverage Week #",
    "Coverage Week Start",
    "Coverage Week End",
    "Coverage Month",
    "Allocated Amount",
    "Source Payment Amount",
    "Allocation Status",
]


EXCEPTION_HEADERS = [
    "Contribution ID",
    "Member ID",
    "Full Name",
    "Payment Date",
    "Payment Amount",
    "Weekly Rate",
    "Issue",
    "Created At",
]


NAVY = "172033"
GOLD = "C9A227"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"
RED = "FDECEC"


def load_settings():
    settings = dict(DEFAULT_SETTINGS)

    if not COVERAGE_FILE.exists():
        return settings

    workbook = load_workbook(
        COVERAGE_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        if "Settings" not in workbook.sheetnames:
            return settings

        sheet = workbook["Settings"]

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            if not row:
                continue

            key = str(row[0] or "").strip()

            if not key:
                continue

            value = (
                str(row[1]).strip()
                if len(row) > 1 and row[1] is not None
                else ""
            )

            if value:
                settings[key] = value

        return settings

    finally:
        workbook.close()


def _style_header(sheet, row_number=1):
    fill = PatternFill(
        "solid",
        fgColor=NAVY,
    )

    font = Font(
        bold=True,
        color=WHITE,
    )

    thin = Side(
        style="thin",
        color="D9DEE7",
    )

    for cell in sheet[row_number]:
        if cell.value is None:
            continue

        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(
            bottom=thin,
        )


def _format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_cells in sheet.columns:
        values = [
            str(cell.value or "")
            for cell in column_cells
        ]

        width = min(
            max(
                max(
                    len(value)
                    for value in values
                ) + 2,
                12,
            ),
            32,
        )

        column_letter = column_cells[0].column_letter

        sheet.column_dimensions[
            column_letter
        ].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )


def _clear_or_create_sheet(workbook, name):
    if name in workbook.sheetnames:
        sheet = workbook[name]

        if sheet.max_row:
            sheet.delete_rows(
                1,
                sheet.max_row,
            )

        return sheet

    return workbook.create_sheet(name)


def write_coverage_workbook(
    coverage_rows,
    exceptions,
    summary,
    settings,
    monthly_rows,
):
    COVERAGE_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if COVERAGE_FILE.exists():
        workbook = load_workbook(
            COVERAGE_FILE
        )
    else:
        workbook = Workbook()

    try:
        default_sheet = workbook.active

        summary_sheet = _clear_or_create_sheet(
            workbook,
            "Summary",
        )

        coverage_sheet = _clear_or_create_sheet(
            workbook,
            "Coverage Ledger",
        )

        monthly_sheet = _clear_or_create_sheet(
            workbook,
            "Monthly Coverage",
        )

        exceptions_sheet = _clear_or_create_sheet(
            workbook,
            "Exceptions",
        )

        settings_sheet = _clear_or_create_sheet(
            workbook,
            "Settings",
        )

        if (
            default_sheet.title
            not in {
                "Summary",
                "Coverage Ledger",
                "Monthly Coverage",
                "Exceptions",
                "Settings",
            }
        ):
            workbook.remove(default_sheet)

        # ------------------------------------------------------
        # SUMMARY
        # ------------------------------------------------------

        summary_sheet.append(
            ["TLC CONTRIBUTION COVERAGE SUMMARY", ""]
        )

        summary_items = [
            (
                "Payment Records Processed",
                summary["payment_records_processed"],
            ),
            (
                "Coverage Weeks Generated",
                summary["coverage_weeks_generated"],
            ),
            (
                "Cash Amount Represented",
                summary["cash_amount"],
            ),
            (
                "Allocated Coverage",
                summary["allocated_amount"],
            ),
            (
                "Exceptions",
                summary["exceptions"],
            ),
        ]

        for label, value in summary_items:
            summary_sheet.append(
                [label, value]
            )

        summary_sheet["A1"].fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )

        summary_sheet["A1"].font = Font(
            bold=True,
            color=WHITE,
            size=14,
        )

        summary_sheet["B1"].fill = PatternFill(
            "solid",
            fgColor=NAVY,
        )

        summary_sheet.column_dimensions["A"].width = 32
        summary_sheet.column_dimensions["B"].width = 22

        # ------------------------------------------------------
        # COVERAGE LEDGER
        # ------------------------------------------------------

        coverage_sheet.append(
            COVERAGE_HEADERS
        )

        for row in coverage_rows:
            coverage_sheet.append(
                [
                    row.get(header, "")
                    for header in COVERAGE_HEADERS
                ]
            )

        _style_header(
            coverage_sheet
        )

        for row in coverage_sheet.iter_rows(
            min_row=2
        ):
            for cell in row:
                if cell.column in {
                    11,
                    12,
                }:
                    cell.number_format = (
                        '#,##0.00'
                    )

        _format_sheet(
            coverage_sheet
        )

        # ------------------------------------------------------
        # MONTHLY COVERAGE
        # ------------------------------------------------------

        monthly_sheet.append(
            [
                "Member ID",
                "Full Name",
                "Month",
                "Weeks Covered",
                "Amount Covered",
            ]
        )

        for row in monthly_rows:
            monthly_sheet.append(
                [
                    row["member_id"],
                    row["full_name"],
                    row["month"],
                    row["weeks"],
                    row["amount"],
                ]
            )

        _style_header(
            monthly_sheet
        )

        _format_sheet(
            monthly_sheet
        )

        for row in monthly_sheet.iter_rows(
            min_row=2
        ):
            row[4].number_format = (
                '#,##0.00'
            )

        # ------------------------------------------------------
        # EXCEPTIONS
        # ------------------------------------------------------

        exceptions_sheet.append(
            EXCEPTION_HEADERS
        )

        for row in exceptions:
            exceptions_sheet.append(
                [
                    row.get(header, "")
                    for header in EXCEPTION_HEADERS
                ]
            )

        _style_header(
            exceptions_sheet
        )

        for row in exceptions_sheet.iter_rows(
            min_row=2
        ):
            for cell in row:
                cell.fill = PatternFill(
                    "solid",
                    fgColor=RED,
                )

        _format_sheet(
            exceptions_sheet
        )

        # ------------------------------------------------------
        # SETTINGS
        # ------------------------------------------------------

        settings_sheet.append(
            [
                "Setting",
                "Value",
            ]
        )

        for key, value in settings.items():
            settings_sheet.append(
                [key, value]
            )

        _style_header(
            settings_sheet
        )

        _format_sheet(
            settings_sheet
        )

        workbook.save(
            COVERAGE_FILE
        )

    finally:
        workbook.close()
