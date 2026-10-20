import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation

from dateutil.rrule import WEEKLY, rrule
from openpyxl import load_workbook

from .workbook import (
    COVERAGE_FILE,
    load_settings,
    write_coverage_workbook,
)


CONTRIBUTIONS_FILE = (
    COVERAGE_FILE.parent
    / "02_CONTRIBUTIONS.xlsx"
)

EXCLUDED_STATUSES = {
    "CANCELLED",
    "REJECTED",
}

WEEKLY_PAYMENT_TYPE = "WEEKLY CONTRIBUTION"

WEEKDAY_NUMBERS = {
    "MONDAY": 0,
    "TUESDAY": 1,
    "WEDNESDAY": 2,
    "THURSDAY": 3,
    "FRIDAY": 4,
    "SATURDAY": 5,
    "SUNDAY": 6,
}


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _parse_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = _clean(value)

    if not text:
        return None

    for parser in (
        lambda x: datetime.fromisoformat(
            x.replace("Z", "+00:00")
        ).date(),
        lambda x: datetime.strptime(
            x,
            "%Y-%m-%d",
        ).date(),
    ):
        try:
            return parser(text)
        except ValueError:
            continue

    return None


def _parse_time(value):
    if isinstance(value, datetime):
        return value.time()

    if isinstance(value, time):
        return value

    text = _clean(value)

    if not text:
        return time.min

    for fmt in (
        "%H:%M:%S",
        "%H:%M",
    ):
        try:
            return datetime.strptime(
                text,
                fmt,
            ).time()
        except ValueError:
            continue

    return time.min


def _amount(value):
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _format_date(value):
    if not value:
        return ""

    return value.strftime("%Y-%m-%d")


def _weekday_number(name):
    normalized = (
        str(name or "THURSDAY")
        .strip()
        .upper()
    )

    if normalized not in WEEKDAY_NUMBERS:
        raise ValueError(
            "Invalid contribution weekday: "
            + normalized
        )

    return WEEKDAY_NUMBERS[normalized]


def _next_or_same_weekday(
    payment_date,
    weekday_number,
):
    offset = (
        weekday_number
        - payment_date.weekday()
    ) % 7

    return payment_date + timedelta(
        days=offset
    )


def _normalize_description(description):
    text = _clean(description).upper()

    text = (
        text
        .replace("-", "-")
        .replace("-", "-")
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    ).strip()

    return text


def _parse_coverage_description(description):
    """
    Convert TLC descriptions into explicit
    coverage instructions.

    Supported examples:

        WEEK 1
        WEEK3
        WEEK 3-4
        2-WEEKS
        2- WEEKS
        10-WEEKS
    """

    text = _normalize_description(
        description
    )

    if not text:
        return {
            "kind": "count",
            "count": None,
        }

    match = re.fullmatch(
        r"WEEK\s*(\d+)\s*-\s*(\d+)",
        text,
    )

    if match:
        start = int(match.group(1))
        end = int(match.group(2))

        if end < start:
            raise ValueError(
                "Invalid week range: "
                f"{text}"
            )

        return {
            "kind": "range",
            "start": start,
            "end": end,
        }

    match = re.fullmatch(
        r"WEEK\s*(\d+)",
        text,
    )

    if match:
        week = int(match.group(1))

        return {
            "kind": "range",
            "start": week,
            "end": week,
        }

    match = re.fullmatch(
        r"(\d+)\s*-\s*WEEKS?",
        text,
    )

    if match:
        return {
            "kind": "count",
            "count": int(
                match.group(1)
            ),
        }

    match = re.fullmatch(
        r"(\d+)\s*WEEKS?",
        text,
    )

    if match:
        return {
            "kind": "count",
            "count": int(
                match.group(1)
            ),
        }

    return {
        "kind": "unknown",
        "text": text,
    }


def _load_contribution_records():
    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            "Contribution workbook not found: "
            f"{CONTRIBUTIONS_FILE}"
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(
            rows,
            None,
        )

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        records = []

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            records.append(
                dict(
                    zip(
                        headers,
                        row,
                    )
                )
            )

        return records

    finally:
        workbook.close()


def _sort_key(record):
    payment_date = (
        _parse_date(
            record.get(
                "Payment Date"
            )
        )
        or date.min
    )

    payment_time = _parse_time(
        record.get(
            "Payment Time"
        )
    )

    contribution_id = _clean(
        record.get(
            "Contribution ID"
        )
    )

    return (
        payment_date,
        payment_time,
        contribution_id,
    )


def _discover_week_one_start(
    records,
    weekday_number,
):
    candidates = []

    for record in records:

        payment_type = _clean(
            record.get(
                "Payment Type"
            )
        ).upper()

        if payment_type != WEEKLY_PAYMENT_TYPE:
            continue

        parsed = _parse_coverage_description(
            record.get(
                "Description"
            )
        )

        if parsed.get("kind") != "range":
            continue

        if parsed.get("start") != 1:
            continue

        payment_date = _parse_date(
            record.get(
                "Payment Date"
            )
        )

        if payment_date:
            candidates.append(
                payment_date
            )

    if not candidates:
        return None

    earliest_payment = min(
        candidates
    )

    return _next_or_same_weekday(
        earliest_payment,
        weekday_number,
    )


def _generate_week_dates(
    first_week,
    count,
):
    if count <= 0:
        return []

    start_datetime = datetime.combine(
        first_week,
        time.min,
    )

    generated = rrule(
        WEEKLY,
        count=count,
        dtstart=start_datetime,
    )

    return [
        item.date()
        for item in generated
    ]


def _week_date(
    week_one_start,
    week_number,
):
    return week_one_start + timedelta(
        days=7 * (week_number - 1)
    )


def _coverage_week_numbers_for_count(
    member_id,
    count,
    covered_weeks_by_member,
):
    if count <= 0:
        return []

    existing = covered_weeks_by_member[
        member_id
    ]

    if not existing:
        start_week = 1
    else:
        start_week = max(existing) + 1

    return list(
        range(
            start_week,
            start_week + count,
        )
    )


def build_coverage_ledger():
    settings = load_settings()

    weekly_rate = _amount(
        settings.get(
            "Weekly Contribution Amount",
            "100",
        )
    )

    if weekly_rate <= 0:
        raise ValueError(
            "Weekly Contribution Amount "
            "must be greater than zero."
        )

    weekday_number = _weekday_number(
        settings.get(
            "Contribution Weekday",
            "THURSDAY",
        )
    )

    records = sorted(
        _load_contribution_records(),
        key=_sort_key,
    )

    week_one_start = (
        _discover_week_one_start(
            records,
            weekday_number,
        )
    )

    coverage_rows = []
    exceptions = []

    covered_weeks_by_member = {}

    coverage_id_number = 1

    payment_records_processed = 0
    total_cash_amount = Decimal("0")
    total_allocated_amount = Decimal("0")

    for record in records:

        status = _clean(
            record.get("Status")
        ).upper()

        if status in EXCLUDED_STATUSES:
            continue

        if not status.startswith("RECORDED"):
            continue

        payment_type = _clean(
            record.get(
                "Payment Type"
            )
        ).upper()

        if payment_type != WEEKLY_PAYMENT_TYPE:
            continue

        member_id = _clean(
            record.get(
                "Member ID"
            )
        ).upper()

        if not member_id:
            continue

        if member_id.startswith(
            "TLC/TEST/"
        ):
            continue

        contribution_id = _clean(
            record.get(
                "Contribution ID"
            )
        )

        payment_date = _parse_date(
            record.get(
                "Payment Date"
            )
        )

        amount = _amount(
            record.get("Amount")
        )

        description = _clean(
            record.get("Description")
        )

        payment_records_processed += 1
        total_cash_amount += amount

        if not contribution_id:
            exceptions.append({
                "Contribution ID": "",
                "Member ID": member_id,
                "Full Name": _clean(
                    record.get(
                        "Full Name"
                    )
                ),
                "Payment Date": _format_date(
                    payment_date
                ),
                "Payment Amount": float(
                    amount
                ),
                "Weekly Rate": float(
                    weekly_rate
                ),
                "Issue": (
                    "Missing Contribution ID."
                ),
                "Created At": datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            })
            continue

        if payment_date is None:
            exceptions.append({
                "Contribution ID": contribution_id,
                "Member ID": member_id,
                "Full Name": _clean(
                    record.get(
                        "Full Name"
                    )
                ),
                "Payment Date": "",
                "Payment Amount": float(
                    amount
                ),
                "Weekly Rate": float(
                    weekly_rate
                ),
                "Issue": (
                    "Invalid payment date."
                ),
                "Created At": datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            })
            continue

        if amount <= 0:
            continue

        if member_id not in covered_weeks_by_member:
            covered_weeks_by_member[
                member_id
            ] = set()

        parsed = _parse_coverage_description(
            description
        )

        # ------------------------------------------------------
        # DETERMINE INTENDED WEEK NUMBERS
        # ------------------------------------------------------

        if parsed["kind"] == "range":

            start_week = parsed[
                "start"
            ]

            end_week = parsed[
                "end"
            ]

            week_numbers = list(
                range(
                    start_week,
                    end_week + 1,
                )
            )

        elif parsed["kind"] == "count":

            count = parsed.get(
                "count"
            )

            if count is None:

                count = int(
                    amount / weekly_rate
                )

            week_numbers = (
                _coverage_week_numbers_for_count(
                    member_id,
                    count,
                    covered_weeks_by_member,
                )
            )

        else:

            count = int(
                amount / weekly_rate
            )

            if count <= 0:
                exceptions.append({
                    "Contribution ID": contribution_id,
                    "Member ID": member_id,
                    "Full Name": _clean(
                        record.get(
                            "Full Name"
                        )
                    ),
                    "Payment Date": _format_date(
                        payment_date
                    ),
                    "Payment Amount": float(
                        amount
                    ),
                    "Weekly Rate": float(
                        weekly_rate
                    ),
                    "Issue": (
                        "Could not determine "
                        "number of covered weeks."
                    ),
                    "Created At": datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),
                })
                continue

            week_numbers = (
                _coverage_week_numbers_for_count(
                    member_id,
                    count,
                    covered_weeks_by_member,
                )
            )

            exceptions.append({
                "Contribution ID": contribution_id,
                "Member ID": member_id,
                "Full Name": _clean(
                    record.get(
                        "Full Name"
                    )
                ),
                "Payment Date": _format_date(
                    payment_date
                ),
                "Payment Amount": float(
                    amount
                ),
                "Weekly Rate": float(
                    weekly_rate
                ),
                "Issue": (
                    "Description was not "
                    "recognized; coverage was "
                    "inferred from amount."
                ),
                "Created At": datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            })

        if not week_numbers:
            continue

        expected_amount = (
            weekly_rate
            * len(week_numbers)
        )

        if expected_amount != amount:
            exceptions.append({
                "Contribution ID": contribution_id,
                "Member ID": member_id,
                "Full Name": _clean(
                    record.get(
                        "Full Name"
                    )
                ),
                "Payment Date": _format_date(
                    payment_date
                ),
                "Payment Amount": float(
                    amount
                ),
                "Weekly Rate": float(
                    weekly_rate
                ),
                "Issue": (
                    "Description indicates "
                    f"{len(week_numbers)} weeks "
                    f"but payment is KSh {amount}."
                ),
                "Created At": datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
            })

        # ------------------------------------------------------
        # CREATE COVERAGE ROWS
        # ------------------------------------------------------

        valid_week_numbers = []

        for week_number in week_numbers:

            if week_number <= 0:
                continue

            if (
                week_number
                in covered_weeks_by_member[
                    member_id
                ]
            ):

                exceptions.append({
                    "Contribution ID": contribution_id,
                    "Member ID": member_id,
                    "Full Name": _clean(
                        record.get(
                            "Full Name"
                        )
                    ),
                    "Payment Date": _format_date(
                        payment_date
                    ),
                    "Payment Amount": float(
                        amount
                    ),
                    "Weekly Rate": float(
                        weekly_rate
                    ),
                    "Issue": (
                        f"Week {week_number} "
                        "is already covered."
                    ),
                    "Created At": datetime.now().strftime(
                        "%Y-%m-%d %H:%M:%S"
                    ),
                })

                continue

            valid_week_numbers.append(
                week_number
            )

        for sequence, week_number in enumerate(
            valid_week_numbers,
            start=1,
        ):

            if week_one_start is not None:

                week_start = _week_date(
                    week_one_start,
                    week_number,
                )

            else:

                fallback_start = (
                    _next_or_same_weekday(
                        payment_date,
                        weekday_number,
                    )
                )

                week_dates = (
                    _generate_week_dates(
                        fallback_start,
                        len(valid_week_numbers),
                    )
                )

                week_start = week_dates[
                    sequence - 1
                ]

            week_end = (
                week_start
                + timedelta(days=6)
            )

            coverage_rows.append({
                "Coverage ID": (
                    f"COV-{coverage_id_number:06d}"
                ),
                "Contribution ID": contribution_id,
                "Receipt Number": _clean(
                    record.get(
                        "Receipt Number"
                    )
                ),
                "Member ID": member_id,
                "Full Name": _clean(
                    record.get(
                        "Full Name"
                    )
                ),
                "Payment Date": _format_date(
                    payment_date
                ),
                "Coverage Week #": week_number,
                "Coverage Week Start": _format_date(
                    week_start
                ),
                "Coverage Week End": _format_date(
                    week_end
                ),
                "Coverage Month": (
                    week_start.strftime(
                        "%Y-%m"
                    )
                ),
                "Allocated Amount": float(
                    weekly_rate
                ),
                "Source Payment Amount": float(
                    amount
                ),
                "Allocation Status": "COVERED",
            })

            covered_weeks_by_member[
                member_id
            ].add(
                week_number
            )

            total_allocated_amount += (
                weekly_rate
            )

            coverage_id_number += 1

    # ----------------------------------------------------------
    # MONTHLY SUMMARY
    # ----------------------------------------------------------

    monthly = {}

    for row in coverage_rows:

        key = (
            row["Member ID"],
            row["Coverage Month"],
        )

        if key not in monthly:
            monthly[key] = {
                "member_id": row[
                    "Member ID"
                ],
                "full_name": row[
                    "Full Name"
                ],
                "month": row[
                    "Coverage Month"
                ],
                "weeks": 0,
                "amount": 0.0,
            }

        monthly[key]["weeks"] += 1

        monthly[key]["amount"] += float(
            row["Allocated Amount"]
        )

    monthly_rows = sorted(
        monthly.values(),
        key=lambda item: (
            item["month"],
            item["full_name"].upper(),
            item["member_id"],
        ),
    )

    return {
        "settings": settings,
        "coverage_rows": coverage_rows,
        "exceptions": exceptions,
        "monthly_rows": monthly_rows,
        "summary": {
            "payment_records_processed": (
                payment_records_processed
            ),
            "coverage_weeks_generated": (
                len(coverage_rows)
            ),
            "cash_amount": float(
                total_cash_amount
            ),
            "allocated_amount": float(
                total_allocated_amount
            ),
            "exceptions": len(
                exceptions
            ),
        },
    }


def sync_coverage_workbook():
    result = build_coverage_ledger()

    write_coverage_workbook(
        coverage_rows=result[
            "coverage_rows"
        ],
        exceptions=result[
            "exceptions"
        ],
        summary=result[
            "summary"
        ],
        settings=result[
            "settings"
        ],
        monthly_rows=result[
            "monthly_rows"
        ],
    )

    return result


def get_monthly_coverage(
    year,
    month,
    member_id=None,
):
    # Statement reads must never rewrite the
    # coverage workbook. Build the current
    # in-memory ledger without saving it.
    result = build_coverage_ledger()

    target_month = (
        f"{int(year):04d}-"
        f"{int(month):02d}"
    )

    normalized_member_id = (
        str(member_id or "")
        .strip()
        .upper()
    )

    return [
        row
        for row in result[
            "coverage_rows"
        ]
        if row["Coverage Month"]
        == target_month
        and (
            not normalized_member_id
            or row["Member ID"]
            == normalized_member_id
        )
    ]


def get_member_monthly_coverage(
    member_id,
    year,
    month,
):
    rows = get_monthly_coverage(
        year=year,
        month=month,
        member_id=member_id,
    )

    total = sum(
        float(
            row["Allocated Amount"]
        )
        for row in rows
    )

    return {
        "member_id": (
            str(
                member_id
            ).strip().upper()
        ),
        "month": (
            f"{int(year):04d}-"
            f"{int(month):02d}"
        ),
        "weeks": rows,
        "week_count": len(rows),
        "amount": total,
    }
