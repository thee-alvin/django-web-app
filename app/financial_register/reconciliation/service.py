from datetime import datetime

from openpyxl import load_workbook

from ..ledger_service import (
    BASE_DIR,
    MEMBERS_FILE,
    CONTRIBUTIONS_FILE,
    EXCLUDED_STATUSES,
    _clean,
    _amount,
    _format_date,
)


def _load_members():
    workbook = load_workbook(
        MEMBERS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Members"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return {}

        headers = [_clean(value) for value in headers]

        member_id_index = headers.index("Member ID")
        name_index = headers.index("Full Name")

        members = {}

        for row in rows:
            member_id = _clean(
                row[member_id_index]
                if member_id_index < len(row)
                else ""
            ).upper()

            if not member_id:
                continue

            if member_id.startswith("TLC/TEST/"):
                continue

            members[member_id] = _clean(
                row[name_index]
                if name_index < len(row)
                else ""
            ) or member_id

        return members

    finally:
        workbook.close()


def _load_records():
    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [_clean(value) for value in headers]

        records = []

        for row in rows:
            if not any(
                value is not None
                for value in row
            ):
                continue

            records.append(
                dict(zip(headers, row))
            )

        return records

    finally:
        workbook.close()


def build_reconciliation():
    members = _load_members()
    records = _load_records()

    valid_records = []
    excluded_records = []

    missing_receipts = []
    duplicate_receipts = []
    missing_contribution_ids = []
    unknown_members = []

    receipt_seen = {}

    raw_total = 0.0

    registration_fees = 0.0
    weekly_contributions = 0.0
    fines = 0.0

    for record in records:

        status = _clean(
            record.get("Status")
        ).upper()

        member_id = _clean(
            record.get("Member ID")
        ).upper()

        receipt_number = _clean(
            record.get("Receipt Number")
        )

        contribution_id = _clean(
            record.get("Contribution ID")
        )

        amount = _amount(
            record.get("Amount")
        )

        payment_type = _clean(
            record.get("Payment Type")
        ).upper()

        payment_date = _format_date(
            record.get("Payment Date")
        )

        # Excluded records are tracked separately.
        if status in EXCLUDED_STATUSES:
            excluded_records.append({
                "receipt_number": receipt_number,
                "member_id": member_id,
                "amount": amount,
                "status": status,
            })
            continue

        # Test member never belongs in the official register.
        if member_id.startswith("TLC/TEST/"):
            continue

        valid_records.append(record)

        raw_total += amount

        if payment_type == "REGISTRATION FEE":
            registration_fees += amount

        elif payment_type == "WEEKLY CONTRIBUTION":
            weekly_contributions += amount

        elif payment_type == "FINE":
            fines += amount

        if not receipt_number:
            missing_receipts.append({
                "member_id": member_id,
                "full_name": _clean(
                    record.get("Full Name")
                ),
                "date": payment_date,
                "amount": amount,
            })
        else:
            receipt_seen.setdefault(
                receipt_number.upper(),
                []
            ).append(record)

        if not contribution_id:
            missing_contribution_ids.append({
                "member_id": member_id,
                "full_name": _clean(
                    record.get("Full Name")
                ),
                "receipt_number": receipt_number,
                "amount": amount,
            })

        if member_id and member_id not in members:
            unknown_members.append({
                "member_id": member_id,
                "full_name": _clean(
                    record.get("Full Name")
                ),
                "receipt_number": receipt_number,
            })

    for receipt_number, receipt_records in receipt_seen.items():
        if len(receipt_records) > 1:
            duplicate_receipts.append({
                "receipt_number": receipt_number,
                "count": len(receipt_records),
                "records": [
                    {
                        "member_id": _clean(
                            item.get("Member ID")
                        ),
                        "full_name": _clean(
                            item.get("Full Name")
                        ),
                        "amount": _amount(
                            item.get("Amount")
                        ),
                        "date": _format_date(
                            item.get("Payment Date")
                        ),
                    }
                    for item in receipt_records
                ],
            })

    ledger_registration = registration_fees
    ledger_weekly = weekly_contributions
    ledger_fines = fines

    ledger_contributions = (
        ledger_registration
        + ledger_weekly
    )

    ledger_total_funds = (
        ledger_contributions
        + ledger_fines
    )

    category_amount_match = (
        abs(registration_fees - ledger_registration) < 0.005
        and abs(weekly_contributions - ledger_weekly) < 0.005
        and abs(fines - ledger_fines) < 0.005
    )

    total_amount_match = (
        abs(raw_total - ledger_total_funds) < 0.005
    )

    issues = []

    if missing_receipts:
        issues.append(
            f"{len(missing_receipts)} transaction(s) missing receipt numbers."
        )

    if duplicate_receipts:
        issues.append(
            f"{len(duplicate_receipts)} duplicate receipt reference(s)."
        )

    if missing_contribution_ids:
        issues.append(
            f"{len(missing_contribution_ids)} transaction(s) missing contribution IDs."
        )

    if unknown_members:
        issues.append(
            f"{len(unknown_members)} transaction(s) belong to unknown member IDs."
        )

    if not category_amount_match:
        issues.append(
            "Payment-type totals do not match the official ledger."
        )

    if not total_amount_match:
        issues.append(
            "Total funds do not match the official ledger."
        )

    status = (
        "MATCHED"
        if not issues
        else "REVIEW REQUIRED"
    )

    return {
        "status": status,

        "checked_at": datetime.now().isoformat(
            timespec="seconds"
        ),

        "register": {
            "total_records": len(valid_records),
            "excluded_records": len(excluded_records),
            "total_amount": raw_total,
            "registration_fees": registration_fees,
            "weekly_contributions": weekly_contributions,
            "total_contributions": ledger_contributions,
            "fines": fines,
            "total_funds": ledger_total_funds,
        },

        "integrity": {
            "category_amount_match": category_amount_match,
            "total_amount_match": total_amount_match,
            "receipts_complete": not missing_receipts,
            "receipts_unique": not duplicate_receipts,
            "contribution_ids_complete": not missing_contribution_ids,
            "members_valid": not unknown_members,
        },

        "issues": issues,

        "details": {
            "missing_receipts": missing_receipts,
            "duplicate_receipts": duplicate_receipts,
            "missing_contribution_ids": missing_contribution_ids,
            "unknown_members": unknown_members,
            "excluded_records": excluded_records,
        },
    }

