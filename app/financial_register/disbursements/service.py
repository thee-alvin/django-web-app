from pathlib import Path
from datetime import datetime, date
from openpyxl import Workbook, load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

DISBURSEMENTS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "11_TREASURY_DISBURSEMENTS.xlsx"
)

SHEET_NAME = "Disbursements"

EXCLUDED_STATUSES = {
    "CANCELLED",
    "REJECTED",
    "VOID",
}


HEADERS = [
    "Disbursement ID",
    "Date",
    "Time",
    "Category",
    "Purpose",
    "Paid To",
    "Amount",
    "Payment Method",
    "Reference",
    "Authorized By",
    "Recorded By",
    "Status",
    "Notes",
]


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _format_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return _clean(value)


def _format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")

    if hasattr(value, "strftime"):
        try:
            return value.strftime("%H:%M:%S")
        except Exception:
            pass

    return _clean(value)


def _ensure_workbook():
    DISBURSEMENTS_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if DISBURSEMENTS_FILE.exists():
        return

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = SHEET_NAME

    sheet.append(HEADERS)

    for cell in sheet[1]:
        cell.font = cell.font.copy(
            bold=True
        )

    widths = {
        1: 24,
        2: 14,
        3: 12,
        4: 22,
        5: 35,
        6: 28,
        7: 16,
        8: 20,
        9: 25,
        10: 25,
        11: 25,
        12: 18,
        13: 40,
    }

    for column, width in widths.items():
        sheet.column_dimensions[
            chr(64 + column)
        ].width = width

    sheet.freeze_panes = "A2"

    workbook.save(
        DISBURSEMENTS_FILE
    )


def _load_records():
    _ensure_workbook()

    workbook = load_workbook(
        DISBURSEMENTS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        if SHEET_NAME not in workbook.sheetnames:
            return []

        sheet = workbook[SHEET_NAME]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        records = []

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            records.append(
                dict(
                    zip(headers, row)
                )
            )

        return records

    finally:
        workbook.close()


def _next_disbursement_id(
    records,
    year,
):
    prefix = (
        f"TLC/DISB/{year}/"
    )

    largest = 0

    for record in records:

        value = _clean(
            record.get(
                "Disbursement ID"
            )
        )

        if not value.startswith(prefix):
            continue

        number_text = value[len(prefix):]

        try:
            largest = max(
                largest,
                int(number_text)
            )
        except ValueError:
            continue

    return (
        f"{prefix}{largest + 1:06d}"
    )


def load_disbursements(
    year=None,
    month=None,
):
    records = _load_records()

    result = []

    for record in records:

        status = _clean(
            record.get("Status")
        ).upper()

        if status in EXCLUDED_STATUSES:
            continue

        transaction_date = _format_date(
            record.get("Date")
        )

        if year:
            if not transaction_date.startswith(
                f"{year}-"
            ):
                continue

        if month:

            month_text = (
                f"{int(month):02d}"
            )

            if not transaction_date.startswith(
                f"{year}-{month_text}"
            ):
                continue

        result.append({
            "disbursement_id":
                _clean(
                    record.get(
                        "Disbursement ID"
                    )
                ),

            "date":
                transaction_date,

            "time":
                _format_time(
                    record.get("Time")
                ),

            "category":
                _clean(
                    record.get("Category")
                ),

            "purpose":
                _clean(
                    record.get("Purpose")
                ),

            "paid_to":
                _clean(
                    record.get("Paid To")
                ),

            "amount":
                _amount(
                    record.get("Amount")
                ),

            "payment_method":
                _clean(
                    record.get(
                        "Payment Method"
                    )
                ),

            "reference":
                _clean(
                    record.get("Reference")
                ),

            "authorized_by":
                _clean(
                    record.get(
                        "Authorized By"
                    )
                ),

            "recorded_by":
                _clean(
                    record.get(
                        "Recorded By"
                    )
                ),

            "status":
                status,

            "notes":
                _clean(
                    record.get("Notes")
                ),
        })

    result.sort(
        key=lambda item: (
            item["date"],
            item["time"],
        ),
        reverse=True,
    )

    return result


def build_disbursement_register(
    year=None,
    month=None,
):
    records = load_disbursements(
        year=year,
        month=month,
    )

    total_withdrawals = sum(
        record["amount"]
        for record in records
    )

    return {
        "generated_at":
            datetime.now().isoformat(
                timespec="seconds"
            ),

        "summary": {
            "transactions":
                len(records),

            "total_withdrawals":
                total_withdrawals,
        },

        "transactions":
            records,
    }


def record_disbursement(
    amount,
    category,
    purpose,
    paid_to,
    payment_method,
    reference,
    authorized_by,
    recorded_by,
    notes="",
    transaction_date=None,
    transaction_time=None,
):
    amount = _amount(amount)

    if amount <= 0:
        raise ValueError(
            "Disbursement amount must be greater than zero."
        )

    category = _clean(category)
    purpose = _clean(purpose)
    paid_to = _clean(paid_to)
    payment_method = _clean(payment_method)
    reference = _clean(reference)
    authorized_by = _clean(authorized_by)
    recorded_by = _clean(recorded_by)
    notes = _clean(notes)

    if not category:
        raise ValueError(
            "Disbursement category is required."
        )

    if not purpose:
        raise ValueError(
            "Disbursement purpose is required."
        )

    if not paid_to:
        raise ValueError(
            "Paid To is required."
        )

    if not authorized_by:
        raise ValueError(
            "Authorized By is required."
        )

    if not recorded_by:
        raise ValueError(
            "Recorded By is required."
        )

    now = datetime.now()

    transaction_date = (
        transaction_date
        or now.strftime("%Y-%m-%d")
    )

    transaction_time = (
        transaction_time
        or now.strftime("%H:%M:%S")
    )

    transaction_date = _format_date(
        transaction_date
    )

    if len(transaction_date) != 10:
        raise ValueError(
            "Transaction date must use YYYY-MM-DD."
        )

    try:
        year = int(
            transaction_date[:4]
        )
    except ValueError as exc:
        raise ValueError(
            "Invalid transaction date."
        ) from exc

    _ensure_workbook()

    records = _load_records()

    disbursement_id = (
        _next_disbursement_id(
            records,
            year,
        )
    )

    workbook = load_workbook(
        DISBURSEMENTS_FILE
    )

    try:
        sheet = workbook[SHEET_NAME]

        sheet.append([
            disbursement_id,
            transaction_date,
            transaction_time,
            category,
            purpose,
            paid_to,
            amount,
            payment_method,
            reference,
            authorized_by,
            recorded_by,
            "PAID",
            notes,
        ])

        workbook.save(
            DISBURSEMENTS_FILE
        )

    finally:
        workbook.close()

    return {
        "disbursement_id":
            disbursement_id,

        "date":
            transaction_date,

        "time":
            transaction_time,

        "category":
            category,

        "purpose":
            purpose,

        "paid_to":
            paid_to,

        "amount":
            amount,

        "payment_method":
            payment_method,

        "reference":
            reference,

        "authorized_by":
            authorized_by,

        "recorded_by":
            recorded_by,

        "status":
            "PAID",

        "notes":
            notes,
    }


_ensure_workbook()
