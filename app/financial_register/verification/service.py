from pathlib import Path
from datetime import datetime, date

from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[3]

CONTRIBUTIONS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "02_CONTRIBUTIONS.xlsx"
)

EXCLUDED_STATUSES = {
    "CANCELLED",
    "REJECTED",
}


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _format_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    text = _clean(value)

    if not text:
        return ""

    try:
        return datetime.fromisoformat(
            text.replace("Z", "+00:00")
        ).strftime("%Y-%m-%d")
    except ValueError:
        pass

    return text.split(" ", 1)[0]


def _format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")

    if hasattr(value, "strftime"):
        try:
            return value.strftime("%H:%M:%S")
        except Exception:
            pass

    return _clean(value)


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def search_receipts(query, limit=10):
    query = _clean(query)

    if not query:
        return []

    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            "TLC contribution workbook was not found."
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        query_upper = query.upper()
        results = []

        for row in rows:
            record = dict(zip(headers, row))

            receipt_number = _clean(
                record.get("Receipt Number")
            )

            if not receipt_number:
                continue

            status = _clean(
                record.get("Status")
            ).upper()

            if status in EXCLUDED_STATUSES:
                continue

            receipt_upper = receipt_number.upper()

            receipt_digits = "".join(
                character
                for character in receipt_number
                if character.isdigit()
            )

            if not (
                query_upper in receipt_upper
                or (
                    query.isdigit()
                    and query in receipt_digits
                )
            ):
                continue

            results.append({
                "receipt_number": receipt_number,
                "member_id": _clean(
                    record.get("Member ID")
                ),
                "full_name": _clean(
                    record.get("Full Name")
                ),
                "payment_date": _format_date(
                    record.get("Payment Date")
                ),
                "payment_type": _clean(
                    record.get("Payment Type")
                ),
                "amount": _amount(
                    record.get("Amount")
                ),
                "status": status,
            })

        return results[:max(1, int(limit))]

    finally:
        workbook.close()

def find_receipt(receipt_number):
    receipt_number = _clean(
        receipt_number
    )

    if not receipt_number:
        return None

    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            "TLC contribution workbook was not found."
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True,
    )

    try:

        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return None

        headers = [
            _clean(value)
            for value in headers
        ]

        for row in rows:

            record = dict(
                zip(headers, row)
            )

            stored_receipt = _clean(
                record.get("Receipt Number")
            )

            if (
                stored_receipt.upper()
                != receipt_number.upper()
            ):
                continue

            status = _clean(
                record.get("Status")
            ).upper()

            if status in EXCLUDED_STATUSES:
                return None

            return {
                "verified": True,

                "receipt_number":
                    stored_receipt,

                "member_id":
                    _clean(
                        record.get("Member ID")
                    ),

                "full_name":
                    _clean(
                        record.get("Full Name")
                    ),

                "payment_date":
                    _format_date(
                        record.get("Payment Date")
                    ),

                "payment_time":
                    _format_time(
                        record.get("Payment Time")
                    ),

                "payment_type":
                    _clean(
                        record.get("Payment Type")
                    ),

                "description":
                    _clean(
                        record.get("Description")
                    ),

                "amount":
                    _amount(
                        record.get("Amount")
                    ),

                "payment_method":
                    _clean(
                        record.get("Payment Method")
                    ),

                "transaction_reference":
                    _clean(
                        record.get(
                            "Transaction Reference"
                        )
                    ),

                "contribution_id":
                    _clean(
                        record.get(
                            "Contribution ID"
                        )
                    ),

                "status":
                    status,

                "recorded_by":
                    _clean(
                        record.get(
                            "Recorded By"
                        )
                    ),
            }

        return None

    finally:
        workbook.close()



def list_receipt_options():
    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            "TLC contribution workbook was not found."
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        results = []

        for row in rows:
            record = dict(zip(headers, row))

            receipt_number = _clean(
                record.get("Receipt Number")
            )

            if not receipt_number:
                continue

            status = _clean(
                record.get("Status")
            ).upper()

            if status in EXCLUDED_STATUSES:
                continue

            results.append({
                "receipt_number": receipt_number,
                "full_name": _clean(
                    record.get("Full Name")
                ),
            })

        return results

    finally:
        workbook.close()


