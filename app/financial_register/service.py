from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

MEMBERS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

CONTRIBUTIONS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "02_CONTRIBUTIONS.xlsx"
)

EXCLUDED_STATUSES = {
    "CANCELLED",
    "REJECTED",
}


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _format_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return _clean(value)


def _format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")

    if hasattr(value, "strftime"):
        try:
            return value.strftime("%H:%M:%S")
        except Exception:
            pass

    return _clean(value)


def _load_members():
    if not MEMBERS_FILE.exists():
        raise FileNotFoundError(
            f"Members workbook not found: {MEMBERS_FILE}"
        )

    workbook = load_workbook(
        MEMBERS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if "Members" not in workbook.sheetnames:
            raise ValueError(
                "Members sheet was not found."
            )

        sheet = workbook["Members"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return {}

        headers = [
            _clean(value)
            for value in headers
        ]

        if "Member ID" not in headers:
            raise ValueError(
                "Member ID column is missing."
            )

        member_id_index = headers.index(
            "Member ID"
        )

        name_index = (
            headers.index("Full Name")
            if "Full Name" in headers
            else None
        )

        members = {}

        for row in rows:
            member_id = _clean(
                row[member_id_index]
                if member_id_index < len(row)
                else ""
            ).upper()

            if not member_id:
                continue

            name = _clean(
                row[name_index]
                if name_index is not None
                and name_index < len(row)
                else ""
            )

            members[member_id] = {
                "member_id": member_id,
                "full_name": name or member_id,
                "registration_fees": 0.0,
                "weekly_contributions": 0.0,
                "total_contributions": 0.0,
                "fines": 0.0,
                "transaction_count": 0,
                "transactions": [],
            }

        return members

    finally:
        workbook.close()


def _load_contributions():
    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            f"Contributions workbook not found: "
            f"{CONTRIBUTIONS_FILE}"
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if "Contributions" not in workbook.sheetnames:
            raise ValueError(
                "Contributions sheet was not found."
            )

        sheet = workbook["Contributions"]

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(value)
            for value in headers
        ]

        records = []

        for row in rows:
            if not any(
                value is not None
                for value in row
            ):
                continue

            records.append(
                dict(
                    zip(headers, row)
                )
            )

        return records

    finally:
        workbook.close()


def build_financial_register():
    members = _load_members()
    contribution_records = _load_contributions()

    for record in contribution_records:

        status = _clean(
            record.get("Status")
        ).upper()

        if status in EXCLUDED_STATUSES:
            continue

        member_id = _clean(
            record.get("Member ID")
        ).upper()

        if not member_id:
            continue

        if member_id not in members:
            members[member_id] = {
                "member_id": member_id,
                "full_name": _clean(
                    record.get("Full Name")
                ) or member_id,
                "registration_fees": 0.0,
                "weekly_contributions": 0.0,
                "total_contributions": 0.0,
                "fines": 0.0,
                "transaction_count": 0,
                "transactions": [],
            }

        member = members[member_id]

        amount = _amount(
            record.get("Amount")
        )

        payment_type = _clean(
            record.get("Payment Type")
        ).upper()

        transaction = {
            "contribution_id": _clean(
                record.get("Contribution ID")
            ),
            "receipt_number": _clean(
                record.get("Receipt Number")
            ),
            "date": _format_date(
                record.get("Payment Date")
            ),
            "time": _format_time(
                record.get("Payment Time")
            ),
            "amount": amount,
            "payment_type": payment_type,
            "status": status,
        }

        member["transaction_count"] += 1

        member["transactions"].append(
            transaction
        )

        if payment_type == "REGISTRATION FEE":
            member["registration_fees"] += amount

        elif payment_type == "WEEKLY CONTRIBUTION":
            member["weekly_contributions"] += amount

        elif payment_type == "FINE":
            member["fines"] += amount

    member_rows = []

    for member in members.values():

        member["total_contributions"] = (
            member["registration_fees"]
            + member["weekly_contributions"]
        )

        member["transactions"].sort(
            key=lambda item: (
                item["date"],
                item["time"]
            ),
            reverse=True
        )

        member_rows.append(member)

    member_rows.sort(
        key=lambda item: (
            item["full_name"].upper(),
            item["member_id"]
        )
    )

    total_registration_fees = sum(
        item["registration_fees"]
        for item in member_rows
    )

    total_weekly_contributions = sum(
        item["weekly_contributions"]
        for item in member_rows
    )

    total_contributions = sum(
        item["total_contributions"]
        for item in member_rows
    )

    total_fines = sum(
        item["fines"]
        for item in member_rows
    )

    total_transactions = sum(
        item["transaction_count"]
        for item in member_rows
    )

    total_funds = (
        total_contributions
        + total_fines
    )

    return {
        "generated_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "summary": {
            "members": len(member_rows),
            "registration_fees": total_registration_fees,
            "weekly_contributions": total_weekly_contributions,
            "total_contributions": total_contributions,
            "fines": total_fines,
            "transactions": total_transactions,
            "total_funds": total_funds,
        },
        "members": member_rows,
    }
