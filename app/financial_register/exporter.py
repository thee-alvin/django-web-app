from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    PageBreak,
)

from .service import build_financial_register


BASE_DIR = Path(__file__).resolve().parents[2]
LOGO_FILE = BASE_DIR / "frontend" / "assets" / "tlc_logo.jpeg"


def _money(value):
    return f"KSh {float(value or 0):,.2f}"


def create_excel_report():
    report = build_financial_register()

    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "Financial Register"

    headers = [
        "Member",
        "Registration Fees",
        "Weekly Contributions",
        "Total Contributions",
        "Fines",
        "Transactions",
    ]

    summary_sheet.append([
        "THE LEGACY CIRCLE",
        "",
        "",
        "",
        "",
        "",
    ])

    summary_sheet.append([
        "PRIVATE FINANCIAL CONTRIBUTION REGISTER",
        "",
        "",
        "",
        "",
        "",
    ])

    summary_sheet.append([])
    summary_sheet.append(headers)

    for member in report["members"]:
        summary_sheet.append([
            member["full_name"],
            member["registration_fees"],
            member["weekly_contributions"],
            member["total_contributions"],
            member["fines"],
            member["transaction_count"],
        ])

    summary = report["summary"]

    summary_sheet.append([
        "TOTAL",
        summary["registration_fees"],
        summary["weekly_contributions"],
        summary["total_contributions"],
        summary["fines"],
        summary["transactions"],
    ])

    summary_sheet.append([])
    summary_sheet.append([
        "TOTAL TLC FUNDS",
        summary["total_funds"],
    ])

    title_fill = PatternFill(
        "solid",
        fgColor="0B1F33"
    )

    header_fill = PatternFill(
        "solid",
        fgColor="123A5A"
    )

    gold_fill = PatternFill(
        "solid",
        fgColor="D4AF37"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    title_font = Font(
        size=18,
        bold=True,
        color="FFFFFF"
    )

    subtitle_font = Font(
        size=11,
        bold=True,
        color="D4AF37"
    )

    thin = Side(
        style="thin",
        color="D9E2EC"
    )

    for cell in summary_sheet[1]:
        cell.fill = title_fill
        cell.font = title_font

    for cell in summary_sheet[2]:
        cell.fill = title_fill
        cell.font = subtitle_font

    for cell in summary_sheet[4]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    total_row = summary_sheet.max_row - 1

    for cell in summary_sheet[total_row]:
        cell.fill = gold_fill
        cell.font = Font(bold=True)

    for row in summary_sheet.iter_rows(
        min_row=5,
        max_row=total_row
    ):
        for cell in row:
            cell.border = Border(
                bottom=thin
            )

    for row in range(5, total_row + 1):
        for column in range(2, 6):
            summary_sheet.cell(
                row=row,
                column=column
            ).number_format = '#,##0.00'

    summary_sheet.freeze_panes = "A5"

    widths = {
        1: 28,
        2: 20,
        3: 24,
        4: 22,
        5: 14,
        6: 16,
    }

    for column, width in widths.items():
        summary_sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    transaction_sheet = workbook.create_sheet(
        "All Transactions"
    )

    transaction_headers = [
        "Member",
        "Member ID",
        "Date",
        "Time",
        "Payment Type",
        "Amount",
        "Receipt Number",
        "Status",
    ]

    transaction_sheet.append(transaction_headers)

    for member in report["members"]:
        for transaction in member["transactions"]:
            transaction_sheet.append([
                member["full_name"],
                member["member_id"],
                transaction["date"],
                transaction["time"],
                transaction["payment_type"],
                transaction["amount"],
                transaction["receipt_number"],
                transaction["status"],
            ])

    for cell in transaction_sheet[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in transaction_sheet.iter_rows(
        min_row=2
    ):
        row[5].number_format = '#,##0.00'

    transaction_sheet.freeze_panes = "A2"

    transaction_widths = {
        1: 26,
        2: 20,
        3: 14,
        4: 12,
        5: 24,
        6: 16,
        7: 24,
        8: 22,
    }

    for column, width in transaction_widths.items():
        transaction_sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return output


def create_pdf_report():
    report = build_financial_register()
    summary = report["summary"]

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=28,
        leftMargin=28,
        topMargin=28,
        bottomMargin=28,
        title="TLC Private Financial Contribution Register",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TLC_TITLE",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=22,
        leading=26,
        textColor=colors.HexColor("#0B1F33"),
    )

    subtitle_style = ParagraphStyle(
        "TLC_SUBTITLE",
        parent=styles["Normal"],
        alignment=TA_CENTER,
        fontSize=10,
        textColor=colors.HexColor("#64748B"),
    )

    right_style = ParagraphStyle(
        "TLC_RIGHT",
        parent=styles["Normal"],
        alignment=TA_RIGHT,
        fontSize=10,
    )

    story = []

    if LOGO_FILE.exists():
        try:
            from reportlab.platypus import Image

            logo = Image(
                str(LOGO_FILE),
                width=70,
                height=70,
                kind="proportional",
            )

            logo.hAlign = "CENTER"
            story.append(logo)
            story.append(Spacer(1, 8))

        except Exception:
            pass

    story.append(
        Paragraph(
            "THE LEGACY CIRCLE",
            title_style
        )
    )

    story.append(
        Paragraph(
            "PRIVATE FINANCIAL CONTRIBUTION REGISTER",
            subtitle_style
        )
    )

    story.append(Spacer(1, 14))

    summary_data = [
        [
            "Registration Fees",
            "Weekly Contributions",
            "Total Contributions",
            "Fines",
            "Transactions",
            "TOTAL TLC FUNDS",
        ],
        [
            _money(summary["registration_fees"]),
            _money(summary["weekly_contributions"]),
            _money(summary["total_contributions"]),
            _money(summary["fines"]),
            str(summary["transactions"]),
            _money(summary["total_funds"]),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            100,
            115,
            115,
            80,
            80,
            110,
        ],
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#123A5A"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER",
            ),
            (
                "BACKGROUND",
                (-1, 1),
                (-1, 1),
                colors.HexColor("#D4AF37"),
            ),
            (
                "FONTNAME",
                (-1, 1),
                (-1, 1),
                "Helvetica-Bold",
            ),
            (
                "BOX",
                (0, 0),
                (-1, -1),
                0.7,
                colors.HexColor("#D9E2EC"),
            ),
            (
                "INNERGRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor("#D9E2EC"),
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                9,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                9,
            ),
        ])
    )

    story.append(summary_table)
    story.append(Spacer(1, 18))

    member_data = [[
        "Member",
        "Registration Fees",
        "Weekly Contributions",
        "Total Contributions",
        "Fines",
        "Transactions",
    ]]

    for member in report["members"]:
        member_data.append([
            member["full_name"],
            _money(member["registration_fees"]),
            _money(member["weekly_contributions"]),
            _money(member["total_contributions"]),
            _money(member["fines"]),
            str(member["transaction_count"]),
        ])

    member_data.append([
        "TOTAL",
        _money(summary["registration_fees"]),
        _money(summary["weekly_contributions"]),
        _money(summary["total_contributions"]),
        _money(summary["fines"]),
        str(summary["transactions"]),
    ])

    member_table = Table(
        member_data,
        repeatRows=1,
        colWidths=[
            155,
            115,
            125,
            120,
            85,
            80,
        ],
    )

    member_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#123A5A"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (1, 1),
                (-1, -1),
                "RIGHT",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#D9E2EC"),
            ),
            (
                "BACKGROUND",
                (0, -1),
                (-1, -1),
                colors.HexColor("#D4AF37"),
            ),
            (
                "FONTNAME",
                (0, -1),
                (-1, -1),
                "Helvetica-Bold",
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                7,
            ),
        ])
    )

    story.append(member_table)
    story.append(Spacer(1, 14))

    story.append(
        Paragraph(
            "Generated from the official TLC contribution workbook.",
            subtitle_style
        )
    )

    story.append(
        Paragraph(
            "Private Treasurer Report",
            subtitle_style
        )
    )

    document.build(story)

    output.seek(0)

    return output
