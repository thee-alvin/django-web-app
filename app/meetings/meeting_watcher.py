from pathlib import Path
import hashlib
import json
import threading

from openpyxl import load_workbook

from app.meetings.meeting_email_service import (
    send_meeting_email
)
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

MEETINGS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "04_MEETINGS.xlsx"
)

MEMBERS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

STATE_DIR = (
    DATA_DIR
    / "meetings"
)

STATE_FILE = (
    STATE_DIR
    / "meeting_notification_state.json"
)

CHECK_INTERVAL = 2


class MeetingWatcher:

    def __init__(self):
        self._stop_event = threading.Event()
        self._thread = None

    def start(self):

        if (
            self._thread
            and self._thread.is_alive()
        ):
            return

        self._stop_event.clear()

        self._thread = threading.Thread(
            target=self._run,
            name="TLC-Meeting-Watcher",
            daemon=True,
        )

        self._thread.start()

        print(
            "[TLC MEETING WATCHER] Started.",
            flush=True,
        )

    def stop(self):

        self._stop_event.set()

        if self._thread:
            self._thread.join(timeout=3)

    def _clean(self, value):

        if value is None:
            return ""

        if hasattr(value, "strftime"):

            try:
                return value.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            except Exception:
                pass

        return str(value).strip()

    def _load_state(self):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        if not STATE_FILE.exists():

            STATE_FILE.write_text(
                "{}",
                encoding="utf-8"
            )

        try:

            data = json.loads(
                STATE_FILE.read_text(
                    encoding="utf-8"
                )
            )

            return (
                data
                if isinstance(data, dict)
                else {}
            )

        except Exception:

            return {}

    def _save_state(self, state):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        STATE_FILE.write_text(
            json.dumps(
                state,
                indent=4,
                ensure_ascii=False
            ),
            encoding="utf-8"
        )

    def _fingerprint(self, meeting):

        fields = [
            "Meeting ID",
            "Meeting Date",
            "Meeting Time",
            "Meeting Type",
            "Venue",
            "Agenda",
            "Chairperson",
            "Secretary",
            "Status",
            "Minutes",
            "Created At",
        ]

        payload = {
            field: self._clean(
                meeting.get(field)
            )
            for field in fields
        }

        raw = json.dumps(
            payload,
            sort_keys=True,
            ensure_ascii=False
        )

        return hashlib.sha256(
            raw.encode("utf-8")
        ).hexdigest()

    def _load_meetings(self):

        if not MEETINGS_FILE.exists():
            return []

        workbook = load_workbook(
            MEETINGS_FILE,
            read_only=True,
            data_only=True
        )

        try:

            if "Meetings" not in workbook.sheetnames:
                return []

            sheet = workbook["Meetings"]

            rows = sheet.iter_rows(
                values_only=True
            )

            headers = next(rows, None)

            if not headers:
                return []

            headers = [
                self._clean(header)
                for header in headers
            ]

            meetings = []

            for row in rows:

                if not any(
                    value is not None
                    for value in row
                ):
                    continue

                record = {}

                for index, header in enumerate(headers):

                    if not header:
                        continue

                    value = (
                        row[index]
                        if index < len(row)
                        else None
                    )

                    record[header] = self._clean(
                        value
                    )

                status = (
                    record.get(
                        "Status",
                        ""
                    )
                    .strip()
                    .upper()
                )

                # Only scheduled meetings trigger emails.
                if status != "SCHEDULED":
                    continue

                meeting_id = (
                    record.get(
                        "Meeting ID",
                        ""
                    ).strip()
                )

                if not meeting_id:
                    continue

                meetings.append(record)

            return meetings

        finally:
            workbook.close()

    def _load_member_contacts(self):

        if not MEMBERS_FILE.exists():
            return []

        workbook = load_workbook(
            MEMBERS_FILE,
            read_only=True,
            data_only=True
        )

        try:

            if "Members" not in workbook.sheetnames:
                return []

            sheet = workbook["Members"]

            rows = sheet.iter_rows(
                values_only=True
            )

            headers = next(rows, None)

            if not headers:
                return []

            headers = [
                self._clean(header)
                for header in headers
            ]

            try:

                name_index = headers.index(
                    "Full Name"
                )

                email_index = headers.index(
                    "Email Address"
                )

            except ValueError:

                return []

            contacts = []
            seen = set()

            for row in rows:

                if (
                    name_index >= len(row)
                    or email_index >= len(row)
                ):
                    continue

                name = self._clean(
                    row[name_index]
                )

                email = self._clean(
                    row[email_index]
                )

                if not email:
                    continue

                key = email.lower()

                if key in seen:
                    continue

                seen.add(key)

                contacts.append({
                    "name": name,
                    "email": email,
                })

            return contacts

        finally:
            workbook.close()

    def _notify_members(
        self,
        meeting,
        is_update
    ):

        contacts = self._load_member_contacts()

        if not contacts:

            print(
                "[TLC MEETING WATCHER] "
                "No member contacts found.",
                flush=True,
            )

            return

        sent = 0
        failed = 0

        print(
            "[TLC MEETING WATCHER] "
            f"Sending {'UPDATE' if is_update else 'NEW'}: "
            f"{meeting.get('Meeting ID')}",
            flush=True,
        )

        for contact in contacts:

            success, response = (
                send_meeting_email(
                    recipient_email=contact["email"],
                    recipient_name=contact["name"],
                    meeting_id=meeting.get(
                        "Meeting ID",
                        ""
                    ),
                    meeting_date=meeting.get(
                        "Meeting Date",
                        ""
                    ),
                    meeting_time=meeting.get(
                        "Meeting Time",
                        ""
                    ),
                    meeting_type=meeting.get(
                        "Meeting Type",
                        ""
                    ),
                    venue=meeting.get(
                        "Venue",
                        ""
                    ),
                    agenda=meeting.get(
                        "Agenda",
                        ""
                    ),
                    chairperson=meeting.get(
                        "Chairperson",
                        ""
                    ),
                    secretary=meeting.get(
                        "Secretary",
                        ""
                    ),
                    status=meeting.get(
                        "Status",
                        ""
                    ),
                    is_update=is_update,
                )
            )

            if success:

                sent += 1

                print(
                    "[TLC MEETING WATCHER] "
                    f"SENT -> {contact['name']} "
                    f"<{contact['email']}>",
                    flush=True,
                )

            else:

                failed += 1

                print(
                    "[TLC MEETING WATCHER] "
                    f"FAILED -> {contact['name']} "
                    f"<{contact['email']}>: "
                    f"{response}",
                    flush=True,
                )

        print(
            "[TLC MEETING WATCHER] "
            f"Email result: {sent} sent, {failed} failed.",
            flush=True,
        )

    def _run(self):

        first_scan = True

        while not self._stop_event.is_set():

            try:

                meetings = self._load_meetings()

                state = self._load_state()

                current = {}

                for meeting in meetings:

                    meeting_id = (
                        meeting.get(
                            "Meeting ID",
                            ""
                        ).strip()
                    )

                    if not meeting_id:
                        continue

                    fingerprint = (
                        self._fingerprint(
                            meeting
                        )
                    )

                    current[meeting_id] = fingerprint

                    if meeting_id not in state:

                        if True:

                            print(
                                "[TLC MEETING WATCHER] "
                                f"NEW MEETING: {meeting_id}",
                                flush=True,
                            )

                            self._notify_members(
                                meeting,
                                is_update=False,
                            )

                    elif (
                        state[meeting_id]
                        != fingerprint
                    ):

                        print(
                            "[TLC MEETING WATCHER] "
                            f"UPDATED MEETING: {meeting_id}",
                            flush=True,
                        )

                        self._notify_members(
                            meeting,
                            is_update=True,
                        )

                self._save_state(current)

                first_scan = False

            except Exception as error:

                print(
                    "[TLC MEETING WATCHER] ERROR:",
                    repr(error),
                    flush=True,
                )

            self._stop_event.wait(
                CHECK_INTERVAL
            )


meeting_watcher = MeetingWatcher()

