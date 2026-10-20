
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

MEETINGS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "04_MEETINGS.xlsx"
)


def _clean(value):
    if value is None:
        return ""

    return str(value).strip()


def get_member_meetings():
    """
    Read meeting records directly from 04_MEETINGS.xlsx.

    Meetings are group-level records, so all valid meeting
    records are returned to authenticated members.
    """

    if not MEETINGS_FILE.exists():
        raise FileNotFoundError(
            f"Meetings workbook not found: {MEETINGS_FILE}"
        )

    workbook = load_workbook(
        MEETINGS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        if "Meetings" not in workbook.sheetnames:
            raise ValueError(
                "Meetings sheet was not found."
            )

        sheet = workbook["Meetings"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            _clean(header)
            for header in headers
        ]

        meetings = []

        for row in rows:

            if not any(
                value is not None
                for value in row
            ):
                continue

            record = {}

            for index, header in enumerate(headers):

                if not header:
                    continue

                value = (
                    row[index]
                    if index < len(row)
                    else None
                )

                if hasattr(
                    value,
                    "strftime"
                ):
                    if header == "Meeting Date":
                        value = value.strftime(
                            "%Y-%m-%d"
                        )

                    elif header == "Meeting Time":
                        value = value.strftime(
                            "%H:%M:%S"
                        )

                    elif header == "Created At":
                        value = value.strftime(
                            "%Y-%m-%d %H:%M:%S"
                        )

                record[header] = (
                    value
                    if value is not None
                    else ""
                )

            meetings.append(record)

        meetings.sort(
            key=lambda item: (
                str(
                    item.get("Meeting Date")
                    or ""
                ),
                str(
                    item.get("Meeting Time")
                    or ""
                )
            ),
            reverse=True
        )

        return meetings

    finally:
        workbook.close()
