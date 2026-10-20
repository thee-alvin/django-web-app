from pathlib import Path
from datetime import date, datetime, timedelta
import json
import threading


from openpyxl import load_workbook


from app.payments.weekly_contribution_email import (
    send_weekly_contribution_reminder
)


BASE_DIR = Path(__file__).resolve().parents[2]


MEMBERS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

CONTRIBUTIONS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "02_CONTRIBUTIONS.xlsx"
)


STATE_DIR = (
    BASE_DIR
    / "data"
    / "weekly_contributions"
)

STATE_FILE = (
    STATE_DIR
    / "reminder_state.json"
)


CHECK_INTERVAL = 60
WEEKLY_AMOUNT = 100.0


class WeeklyContributionWatcher:

    def __init__(self):

        self._stop_event = threading.Event()
        self._thread = None

    def start(self):

        if (
            self._thread
            and self._thread.is_alive()
        ):
            return

        self._stop_event.clear()

        self._thread = threading.Thread(
            target=self._run,
            name="TLC-Weekly-Contribution-Watcher",
            daemon=True,
        )

        self._thread.start()

        print(
            "[TLC WEEKLY WATCHER] Started.",
            flush=True,
        )

    def stop(self):

        self._stop_event.set()

        if self._thread:
            self._thread.join(timeout=3)

    def _clean(self, value):

        if value is None:
            return ""

        if hasattr(value, "strftime"):

            try:
                return value.strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            except Exception:
                pass

        return str(value).strip()

    def _load_state(self):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        if not STATE_FILE.exists():

            STATE_FILE.write_text(
                "{}",
                encoding="utf-8"
            )

        try:

            data = json.loads(
                STATE_FILE.read_text(
                    encoding="utf-8"
                )
            )

            return (
                data
                if isinstance(data, dict)
                else {}
            )

        except Exception:

            return {}

    def _save_state(self, state):

        STATE_DIR.mkdir(
            parents=True,
            exist_ok=True
        )

        STATE_FILE.write_text(
            json.dumps(
                state,
                indent=4,
                ensure_ascii=False
            ),
            encoding="utf-8"
        )

    def _load_members(self):

        workbook = load_workbook(
            MEMBERS_FILE,
            read_only=True,
            data_only=True
        )

        try:

            sheet = workbook["Members"]

            rows = sheet.iter_rows(
                values_only=True
            )

            headers = next(rows, None)

            if not headers:
                return []

            headers = [
                self._clean(value)
                for value in headers
            ]

            index = {
                header: headers.index(header) + 1
                for header in headers
                if header
            }

            members = []

            for row in rows:

                member_id = self._clean(
                    row[index["Member ID"] - 1]
                    if len(row) >= index["Member ID"]
                    else ""
                )

                name = self._clean(
                    row[index["Full Name"] - 1]
                    if len(row) >= index["Full Name"]
                    else ""
                )

                email = self._clean(
                    row[index["Email Address"] - 1]
                    if len(row) >= index["Email Address"]
                    else ""
                )

                account_status = self._clean(
                    row[index["Account Status"] - 1]
                    if len(row) >= index["Account Status"]
                    else ""
                ).lower()

                if not member_id or not email:
                    continue

                # Only active accounts are reminded.
                if account_status and account_status != "active":
                    continue

                members.append({
                    "member_id": member_id.upper(),
                    "name": name,
                    "email": email,
                })

            return members

        finally:

            workbook.close()

    def _parse_weeks(self, description):

        text = str(
            description or ""
        ).strip().upper()

        digits = ""

        for character in text:

            if character.isdigit():
                digits += character
            elif digits:
                break

        if not digits:
            return 0

        try:
            return int(digits)
        except ValueError:
            return 0

    def _load_weekly_contributions(self):
        workbook = load_workbook(
            CONTRIBUTIONS_FILE,
            read_only=True,
            data_only=True
        )

        try:
            sheet = workbook["Contributions"]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                return {}

            headers = [
                self._clean(value)
                for value in headers
            ]

            columns = {
                header: headers.index(header)
                for header in headers
                if header
            }

            paid_today = set()
            today = date.today()

            for row in rows:
                if not row:
                    continue

                payment_type = self._clean(
                    row[columns["Payment Type"]]
                ).upper()

                status = self._clean(
                    row[columns["Status"]]
                ).upper()

                if payment_type != "WEEKLY CONTRIBUTION":
                    continue

                if status not in (
                    "RECORDED - EMAILED",
                    "RECORDED - EMAIL QUEUED",
                ):
                    continue

                member_id = self._clean(
                    row[columns["Member ID"]]
                ).upper()

                payment_date = row[
                    columns["Payment Date"]
                ]

                if isinstance(payment_date, datetime):
                    payment_date = payment_date.date()

                elif not isinstance(payment_date, date):
                    try:
                        payment_date = datetime.strptime(
                            str(payment_date).strip(),
                            "%Y-%m-%d"
                        ).date()
                    except ValueError:
                        continue

                if (
                    member_id
                    and payment_date == today
                ):
                    paid_today.add(member_id)

            return paid_today

        finally:
            workbook.close()

    def _run_today(self):

        today = date.today()

        # Weekly reminders are issued on Thursdays only.
        if today.weekday() != 3:
            return

        members = self._load_members()

        paid_today = self._load_weekly_contributions()

        state = self._load_state()

        date_key = today.isoformat()

        already_sent = state.get(
            date_key,
            {}
        )

        if not isinstance(
            already_sent,
            dict
        ):
            already_sent = {}

        for member in members:

            member_id = member["member_id"]

            if member_id in paid_today:
                continue

            if already_sent.get(member_id):
                continue

            success, response = (
                send_weekly_contribution_reminder(
                    recipient_email=member["email"],
                    recipient_name=member["name"],
                    week_date=today.strftime(
                        "%d %B %Y"
                    ),
                    amount=WEEKLY_AMOUNT,
                )
            )

            if success:

                already_sent[member_id] = True

                print(
                    "[TLC WEEKLY WATCHER] "
                    f"REMINDER SENT -> "
                    f"{member['name']} "
                    f"<{member['email']}>",
                    flush=True,
                )

            else:

                print(
                    "[TLC WEEKLY WATCHER] "
                    f"REMINDER FAILED -> "
                    f"{member['name']}: "
                    f"{response}",
                    flush=True,
                )

        state[date_key] = already_sent

        # Keep recent state only.
        cutoff = (
            today
            - timedelta(days=30)
        )

        state = {
            key: value
            for key, value in state.items()
            if key >= cutoff.isoformat()
        }

        self._save_state(state)

    def _run(self):

        while not self._stop_event.is_set():

            try:

                self._run_today()

            except Exception as error:

                print(
                    "[TLC WEEKLY WATCHER] ERROR:",
                    repr(error),
                    flush=True,
                )

            self._stop_event.wait(
                CHECK_INTERVAL
            )


weekly_contribution_watcher = (
    WeeklyContributionWatcher()
)

