from pathlib import Path
from datetime import datetime, date

from openpyxl import load_workbook
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]
CONTRIBUTIONS_FILE = (
    DATA_DIR / "spreadsheets" / "02_CONTRIBUTIONS.xlsx"
)


EXCLUDED_STATUSES = {
    "CANCELLED",
    "REJECTED",
}


def _clean(value):
    if value is None:
        return ""
    return str(value).strip()


def _amount(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _format_date(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return _clean(value)


def _format_time(value):
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")

    return _clean(value)


def get_member_contributions(member_id):
    if not member_id:
        return []

    if not CONTRIBUTIONS_FILE.exists():
        raise FileNotFoundError(
            f"Contributions workbook not found: {CONTRIBUTIONS_FILE}"
        )

    workbook = load_workbook(
        CONTRIBUTIONS_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Contributions"]

        headers = [
            _clean(cell.value)
            for cell in sheet[1]
        ]

        contributions = []

        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):
            if not any(value is not None for value in row):
                continue

            record = dict(zip(headers, row))

            row_member_id = _clean(
                record.get("Member ID")
            )

            if row_member_id != _clean(member_id):
                continue

            amount = _amount(
                record.get("Amount")
            )

            status = _clean(
                record.get("Status")
            )

            contributions.append({
                "contribution_id": _clean(
                    record.get("Contribution ID")
                ),
                "receipt_number": _clean(
                    record.get("Receipt Number")
                ),
                "member_id": row_member_id,
                "full_name": _clean(
                    record.get("Full Name")
                ),
                "phone_number": _clean(
                    record.get("Phone Number")
                ),
                "email": _clean(
                    record.get("Email")
                ),
                "payment_date": _format_date(
                    record.get("Payment Date")
                ),
                "payment_time": _format_time(
                    record.get("Payment Time")
                ),
                "payment_type": _clean(
                    record.get("Payment Type")
                ),
                "description": _clean(
                    record.get("Description")
                ),
                "amount": amount,
                "payment_method": _clean(
                    record.get("Payment Method")
                ),
                "transaction_reference": _clean(
                    record.get("Transaction Reference")
                ),
                "status": status,
                "recorded_by": _clean(
                    record.get("Recorded By")
                ),
                "created_at": _format_date(
                    record.get("Created At")
                ),
            })

        contributions.sort(
            key=lambda item: (
                item.get("payment_date") or "",
                item.get("payment_time") or "",
            ),
            reverse=True,
        )

        return contributions

    finally:
        workbook.close()


def get_member_contribution_summary(member_id):
    contributions = get_member_contributions(member_id)

    valid_contributions = [
        item
        for item in contributions
        if item["status"].upper() not in EXCLUDED_STATUSES
    ]

    total = sum(
        item["amount"]
        for item in valid_contributions
    )

    current_month = datetime.now().strftime("%Y-%m")

    this_month = sum(
        item["amount"]
        for item in valid_contributions
        if item["payment_date"].startswith(current_month)
    )

    latest = (
        contributions[0]
        if contributions
        else None
    )

    return {
        "total": total,
        "this_month": this_month,
        "count": len(valid_contributions),
        "latest": latest,
    }
