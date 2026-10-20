from pathlib import Path
from datetime import datetime, timedelta
from openpyxl import load_workbook
import os
from dotenv import load_dotenv, dotenv_values
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

DATA_DIR = DATA_DIR
SPREADSHEETS_DIR = DATA_DIR / "spreadsheets"
RECEIPTS_DIR = DATA_DIR / "receipts"
CLOUDFLARE_LOG = DATA_DIR / "cloudflare" / "tunnel.log"

PROCESSOR_LOCK = RECEIPTS_DIR / ".contribution_processor.lock"

SMTP_ENV = BASE_DIR / ".env"


def _result(name, status, message, details=None):
    item = {
        "name": name,
        "status": status,
        "message": message,
    }

    if details is not None:
        item["details"] = details

    return item


def check_excel_files():
    files = [
        "01_MEMBERS.xlsx",
        "02_CONTRIBUTIONS.xlsx",
        "03_ANNOUNCEMENTS.xlsx",
        "04_MEETINGS.xlsx",
    ]

    results = []

    for filename in files:
        path = SPREADSHEETS_DIR / filename

        if not path.exists():
            results.append(
                _result(
                    filename,
                    "ERROR",
                    "File is missing."
                )
            )
            continue

        try:
            wb = load_workbook(
                path,
                read_only=True,
                data_only=True
            )
            sheets = wb.sheetnames
            wb.close()

            results.append(
                _result(
                    filename,
                    "HEALTHY",
                    "Workbook is readable.",
                    {"sheets": sheets}
                )
            )

        except Exception as exc:
            results.append(
                _result(
                    filename,
                    "ERROR",
                    f"{type(exc).__name__}: {exc}"
                )
            )

    return results


def check_smtp():
    config = dotenv_values(SMTP_ENV, encoding="utf-8-sig")

    host = config.get("TLC_SMTP_HOST")
    port = config.get("TLC_SMTP_PORT")
    username = config.get("TLC_SMTP_USERNAME")
    password = config.get("TLC_SMTP_PASSWORD")

    missing = []

    if not host:
        missing.append("TLC_SMTP_HOST")

    if not port:
        missing.append("TLC_SMTP_PORT")

    if not username:
        missing.append("TLC_SMTP_USERNAME")

    if not password:
        missing.append("TLC_SMTP_PASSWORD")

    if missing:
        return _result(
            "Email Service",
            "ERROR",
            "SMTP configuration is incomplete.",
            {"missing": missing}
        )

    return _result(
        "Email Service",
        "HEALTHY",
        "SMTP configuration is available.",
        {
            "host": host,
            "port": port,
            "username": username
        }
    )

def check_receipt_lock():
    if not PROCESSOR_LOCK.exists():
        return _result(
            "Receipt Processor Lock",
            "HEALTHY",
            "No processor lock is present."
        )

    modified = datetime.fromtimestamp(
        PROCESSOR_LOCK.stat().st_mtime
    )

    age = datetime.now() - modified

    if age > timedelta(minutes=5):
        return _result(
            "Receipt Processor Lock",
            "WARNING",
            "Processor lock appears stale.",
            {
                "age_seconds": int(age.total_seconds()),
                "last_modified": modified.isoformat()
            }
        )

    return _result(
        "Receipt Processor Lock",
        "HEALTHY",
        "Processor lock is active.",
        {
            "age_seconds": int(age.total_seconds())
        }
    )


def check_processing_rows():
    path = SPREADSHEETS_DIR / "02_CONTRIBUTIONS.xlsx"

    if not path.exists():
        return _result(
            "Contribution Processing",
            "ERROR",
            "Contribution workbook is missing."
        )

    try:
        wb = load_workbook(
            path,
            read_only=True,
            data_only=True
        )
        ws = wb.active

        processing = 0
        pending = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            status = str(row[13]).strip().upper() if len(row) > 13 and row[13] else ""

            if status == "PROCESSING":
                processing += 1

            elif status == "PENDING":
                pending += 1

        wb.close()

        status = "HEALTHY"

        if processing > 0:
            status = "WARNING"

        return _result(
            "Contribution Processing",
            status,
            "Contribution queue checked.",
            {
                "pending": pending,
                "processing": processing
            }
        )

    except Exception as exc:
        return _result(
            "Contribution Processing",
            "ERROR",
            f"{type(exc).__name__}: {exc}"
        )


def check_cloudflare():
    if not CLOUDFLARE_LOG.exists():
        return _result(
            "Cloudflare",
            "WARNING",
            "Cloudflare tunnel log is missing."
        )

    modified = datetime.fromtimestamp(
        CLOUDFLARE_LOG.stat().st_mtime
    )

    age = datetime.now() - modified

    if age > timedelta(minutes=10):
        return _result(
            "Cloudflare",
            "WARNING",
            "Tunnel log has not updated recently.",
            {
                "age_seconds": int(age.total_seconds())
            }
        )

    return _result(
        "Cloudflare",
        "HEALTHY",
        "Cloudflare tunnel log is active.",
        {
            "age_seconds": int(age.total_seconds())
        }
    )


def get_system_health():
    checks = []

    checks.append(
        _result(
            "Portal",
            "HEALTHY",
            "TLC application process is responding."
        )
    )

    checks.append(
        check_smtp()
    )

    checks.append(
        check_receipt_lock()
    )

    checks.append(
        check_processing_rows()
    )

    checks.append(
        check_cloudflare()
    )

    excel_results = check_excel_files()

    for item in excel_results:
        checks.append(item)

    errors = sum(
        1 for item in checks
        if item["status"] == "ERROR"
    )

    warnings = sum(
        1 for item in checks
        if item["status"] == "WARNING"
    )

    if errors:
        overall = "ERROR"
    elif warnings:
        overall = "WARNING"
    else:
        overall = "HEALTHY"

    return {
        "system": "The Legacy Circle",
        "overall": overall,
        "checked_at": datetime.now().isoformat(),
        "summary": {
            "total": len(checks),
            "healthy": sum(
                1 for item in checks
                if item["status"] == "HEALTHY"
            ),
            "warnings": warnings,
            "errors": errors
        },
        "checks": checks
    }






