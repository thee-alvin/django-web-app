from pathlib import Path
from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parents[2]

MEMBERS_FILE = (
    BASE_DIR
    / "data"
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

PROFILE_PICTURES_DIR = (
    BASE_DIR
    / "frontend"
    / "assets"
    / "profile_pictures"
)


def find_profile_picture(member_id):
    safe_id = (
        str(member_id or "")
        .strip()
        .upper()
        .replace("/", "_")
        .replace("\\", "_")
    )

    if not safe_id:
        return ""

    PROFILE_PICTURES_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    for extension in (
        ".jpg",
        ".jpeg",
        ".png",
        ".webp"
    ):
        filename = safe_id + extension

        if (
            PROFILE_PICTURES_DIR / filename
        ).exists():
            return filename

    return ""


def get_member_by_id(member_id):
    """
    Find one TLC member using their Member ID.
    """

    member_id = str(
        member_id or ""
    ).strip().upper()

    if not member_id:
        return None

    if not MEMBERS_FILE.exists():
        return None

    workbook = load_workbook(
        MEMBERS_FILE,
        read_only=True,
        data_only=True
    )

    try:

        sheet = workbook["Members"]

        rows = sheet.iter_rows(
            values_only=True
        )

        headers = next(rows, None)

        if not headers:
            return None

        headers = [
            str(header).strip()
            if header is not None
            else ""
            for header in headers
        ]

        for row in rows:

            member = dict(
                zip(headers, row)
            )

            stored_id = str(
                member.get("Member ID") or ""
            ).strip().upper()

            if stored_id != member_id:
                continue

            return {
                "member_id": stored_id,

                "full_name": str(
                    member.get("Full Name") or ""
                ).strip(),

                "phone": str(
                    member.get("Phone Number") or ""
                ).strip(),

                "email": str(
                    member.get("Email Address") or ""
                ).strip(),

                "membership_status": str(
                    member.get("Membership Status") or ""
                ).strip(),

                "registration_date": str(
                    member.get("Registration Date") or ""
                ).strip(),

                "account_status": str(
                    member.get("Account Status") or ""
                ).strip(),

                "last_login": str(
                    member.get("Last Login") or ""
                ).strip(),

                "profile_picture": find_profile_picture(
                    stored_id
                )
            }

        return None

    finally:
        workbook.close()
