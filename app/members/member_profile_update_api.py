from pathlib import Path

from flask import (
    Blueprint,
    jsonify,
    request,
    session
)

from werkzeug.utils import secure_filename

from openpyxl import load_workbook

from app.security.audit_log import record_activity

from app.members.member_service import (
    get_member_by_id,
)
from app.config.data_paths import DATA_DIR


BASE_DIR = Path(__file__).resolve().parents[2]

MEMBERS_FILE = (
    DATA_DIR
    / "spreadsheets"
    / "01_MEMBERS.xlsx"
)

PROFILE_PICTURES_DIR = (
    BASE_DIR
    / "frontend"
    / "assets"
    / "profile_pictures"
)

ALLOWED_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".webp"
}


member_profile_update_api = Blueprint(
    "member_profile_update_api",
    __name__,
    url_prefix="/api/member/profile"
)


def profile_filename(member_id, extension):
    return (
        str(member_id)
        .strip()
        .upper()
        .replace("/", "_")
        .replace("\\", "_")
        + extension
    )


@member_profile_update_api.route(
    "/update",
    methods=["POST"]
)
def update_profile():

    member_id = str(
        session.get("member_id") or ""
    ).strip().upper()

    if not member_id:

        return jsonify({
            "success": False,
            "message": "Please log in to update your profile."
        }), 401

    if not MEMBERS_FILE.exists():

        return jsonify({
            "success": False,
            "message": "TLC member records are unavailable."
        }), 500

    workbook = None

    try:

        workbook = load_workbook(
            MEMBERS_FILE
        )

        sheet = workbook["Members"]

        headers = [
            str(cell.value).strip()
            if cell.value is not None
            else ""
            for cell in sheet[1]
        ]

        required = [
            "Member ID",
            "Full Name",
            "Phone Number",
            "Email Address"
        ]

        missing = [
            name
            for name in required
            if name not in headers
        ]

        if missing:

            return jsonify({
                "success": False,
                "message": (
                    "Required member columns are missing: "
                    + ", ".join(missing)
                )
            }), 500

        columns = {
            name: headers.index(name) + 1
            for name in headers
        }

        target_row = None

        for row in range(
            2,
            sheet.max_row + 1
        ):

            stored_id = str(
                sheet.cell(
                    row=row,
                    column=columns["Member ID"]
                ).value or ""
            ).strip().upper()

            if stored_id == member_id:
                target_row = row
                break

        if target_row is None:

            return jsonify({
                "success": False,
                "message": "Your TLC member record was not found."
            }), 404

        full_name = str(
            request.form.get("full_name") or ""
        ).strip()

        phone = str(
            request.form.get("phone") or ""
        ).strip()

        email = str(
            request.form.get("email") or ""
        ).strip().lower()

        if not full_name:

            return jsonify({
                "success": False,
                "message": "Full name cannot be empty."
            }), 400

        if not phone:

            return jsonify({
                "success": False,
                "message": "Phone number cannot be empty."
            }), 400

        if not email or "@" not in email:

            return jsonify({
                "success": False,
                "message": "Please enter a valid email address."
            }), 400

        # Only editable member fields are changed.
        sheet.cell(
            row=target_row,
            column=columns["Full Name"]
        ).value = full_name

        sheet.cell(
            row=target_row,
            column=columns["Phone Number"]
        ).value = phone

        sheet.cell(
            row=target_row,
            column=columns["Email Address"]
        ).value = email

        # --------------------------------------------------
        # PROFILE PICTURE FROM DEVICE
        # --------------------------------------------------

        photo = request.files.get(
            "profile_picture"
        )

        if photo and photo.filename:

            original_name = secure_filename(
                photo.filename
            )

            extension = Path(
                original_name
            ).suffix.lower()

            if (
                not original_name
                or extension not in ALLOWED_EXTENSIONS
            ):

                return jsonify({
                    "success": False,
                    "message": (
                        "Use JPG, JPEG, PNG or WEBP "
                        "for your profile picture."
                    )
                }), 400

            PROFILE_PICTURES_DIR.mkdir(
                parents=True,
                exist_ok=True
            )

            # Remove any previous picture belonging
            # to this member.
            safe_id = (
                member_id
                .replace("/", "_")
                .replace("\\", "_")
            )

            for old_extension in ALLOWED_EXTENSIONS:

                old_file = (
                    PROFILE_PICTURES_DIR
                    / f"{safe_id}{old_extension}"
                )

                if old_file.exists():
                    old_file.unlink()

            destination = (
                PROFILE_PICTURES_DIR
                / profile_filename(
                    member_id,
                    extension
                )
            )

            photo.save(
                destination
            )

        workbook.save(
            MEMBERS_FILE
        )

        record_activity(
            member_id,
            action="Profile Updated",
            description="Your TLC member profile was updated successfully.",
            category="profile"
        )

        return jsonify({
            "success": True,
            "message": (
                "Your TLC profile was updated successfully."
            ),
            "member": get_member_by_id(
                member_id
            )
        })

    except PermissionError:

        return jsonify({
            "success": False,
            "message": (
                "The member spreadsheet is currently open. "
                "Please close it and try again."
            )
        }), 423

    except Exception:

        return jsonify({
            "success": False,
            "message": (
                "Unable to update your TLC profile."
            )
        }), 500

    finally:

        if workbook is not None:
            workbook.close()



